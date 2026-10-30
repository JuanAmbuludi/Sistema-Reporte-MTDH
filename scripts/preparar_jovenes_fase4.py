"""Import the fourth-period JEA payment monitor as a separate, undated snapshot.

Usage: python -I scripts/preparar_jovenes_fase4.py 'Transferencia rechazadas.xlsx'

The workbook contains aggregate executing-entity totals and rejection reasons,
not nominal records or a dated payment series. Its source totals and cached
formula results are reconciled before writing. No dates are inferred from file
names, filesystem timestamps or document properties; phases 1–3 are untouched.
"""
from __future__ import annotations

import argparse
from collections import Counter
from decimal import Decimal, InvalidOperation
import hashlib
import json
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "data/processed/jovenes_fase4.json"
PAY_FIELDS = {2: "seleccionados", 3: "autorizados", 5: "monto_autorizado", 6: "acreditadas",
              8: "monto_acreditado", 9: "no_acreditadas", 11: "monto_no_acreditado"}
RATE_COLUMNS = {4: ("autorizados", "seleccionados"), 7: ("acreditadas", "autorizados"),
                10: ("no_acreditadas", "autorizados")}


def text(value) -> str:
    return " ".join(str(value or "").split())


def number(value, location, *, money=False):
    if value is None or isinstance(value, bool):
        raise ValueError(f"{location}: falta un valor numérico o el resultado guardado de la fórmula.")
    try:
        parsed = Decimal(str(value))
    except InvalidOperation as exc:
        raise ValueError(f"{location}: valor numérico inválido.") from exc
    if not parsed.is_finite() or parsed < 0:
        raise ValueError(f"{location}: el valor debe ser finito y no negativo.")
    quantum = Decimal("0.01") if money else Decimal(1)
    if parsed != parsed.quantize(quantum):
        raise ValueError(f"{location}: precisión inválida para {'USD' if money else 'un conteo'}.")
    return float(parsed) if money else int(parsed)


def _balance(row, label):
    if row["seleccionados"] < row["autorizados"]:
        raise ValueError(f"{label}: autorizados supera seleccionados.")
    if row["autorizados"] != row["acreditadas"] + row["no_acreditadas"]:
        raise ValueError(f"{label}: acreditadas y no acreditadas no concilian con autorizados.")
    amount = Decimal(str(row["monto_acreditado"])) + Decimal(str(row["monto_no_acreditado"]))
    if Decimal(str(row["monto_autorizado"])) != amount:
        raise ValueError(f"{label}: los montos no concilian con el monto autorizado.")


def _rates(sheet, line, row):
    for col, (num, den) in RATE_COLUMNS.items():
        value = sheet.cell(line, col).value
        try:
            cached = float(value)
        except (ValueError, TypeError) as exc:
            raise ValueError(f"Pagos {sheet.cell(line, col).coordinate}: porcentaje no disponible.") from exc
        ratio = row[num] / row[den] if row[den] else 0
        if not 0 <= cached <= 1 or abs(cached - ratio) > 1e-10:
            raise ValueError(f"Pagos fila {line}: porcentaje inconsistente con su denominador.")


def read_payments(sheet):
    expected = {"A1": "CUARTO PERÍODO JÓVENES EN ACCIÓN", "A2": "Entidad Ejecutora",
                "B2": "Jóvenes seleccionados", "C2": "Jóvenes autorizados",
                "F2": "Transferencias Acreditadas", "I2": "Transferencias No Acreditadas"}
    if any(text(sheet[cell].value).casefold() != label.casefold() for cell, label in expected.items()):
        raise ValueError("Pagos: cambió la estructura o el período de la fuente; revisar el contrato antes de importar.")
    rows, totals = [], None
    for line in range(4, sheet.max_row + 1):
        entity = text(sheet.cell(line, 1).value)
        if not entity:
            if any(sheet.cell(line, col).value is not None for col in PAY_FIELDS):
                raise ValueError(f"Pagos fila {line}: cifras sin entidad.")
            continue
        row = {field: number(sheet.cell(line, col).value, f"Pagos {sheet.cell(line, col).coordinate}",
                             money=field.startswith("monto_")) for col, field in PAY_FIELDS.items()}
        _balance(row, f"Pagos {entity}")
        _rates(sheet, line, row)
        if entity.casefold() == "total":
            if totals is not None:
                raise ValueError("Pagos: más de una fila Total.")
            totals = row
        else:
            rows.append({"entidad": entity, **row})
    entities = [r["entidad"] for r in rows]
    if not entities or len(set(entities)) != len(entities):
        raise ValueError("Pagos: entidades vacías o duplicadas.")
    if totals is None:
        raise ValueError("Pagos: falta el total de control.")
    for field in PAY_FIELDS.values():
        if sum(Decimal(str(r[field])) for r in rows) != Decimal(str(totals[field])):
            raise ValueError(f"Pagos: el total de {field} no coincide con las entidades.")
    return rows, totals


def read_rejections(sheet, payments, totals):
    if text(sheet["B2"].value).casefold() != "motivo del rechazo":
        raise ValueError("Rechazos: falta el encabezado de motivos.")
    entities = [text(sheet.cell(2, col).value) for col in range(3, sheet.max_column)]
    expected = {r["entidad"]: r["no_acreditadas"] for r in payments}
    if len(entities) != len(set(entities)) or set(entities) != set(expected):
        raise ValueError("Rechazos: las entidades no coinciden con Pagos.")
    if text(sheet.cell(2, sheet.max_column).value).casefold() != "total":
        raise ValueError("Rechazos: falta la columna Total.")
    rows, codes, controls = [], set(), None
    for line in range(3, sheet.max_row + 1):
        value = sheet.cell(line, 1).value
        if value is None:
            continue
        counts = [number(sheet.cell(line, col).value, f"Rechazos {sheet.cell(line, col).coordinate}")
                  for col in range(3, sheet.max_column)]
        total = number(sheet.cell(line, sheet.max_column).value, f"Rechazos total fila {line}")
        if sum(counts) != total:
            raise ValueError(f"Rechazos fila {line}: las entidades no coinciden con el total del motivo.")
        if text(value).casefold() == "total":
            if controls is not None:
                raise ValueError("Rechazos: más de una fila Total.")
            controls = dict(zip(entities, counts))
            if total != totals["no_acreditadas"]:
                raise ValueError("Rechazos: el total no coincide con transferencias no acreditadas.")
            continue
        code = number(value, f"Rechazos código fila {line}")
        if code in codes:
            raise ValueError("Rechazos: código repetido.")
        codes.add(code)
        reason = text(sheet.cell(line, 2).value)
        if not reason:
            raise ValueError(f"Rechazos fila {line}: motivo vacío.")
        rows.extend({"entidad": entity, "codigo_rechazo": code, "motivo": reason,
                     "transferencias": count} for entity, count in zip(entities, counts))
    actual = Counter()
    for row in rows:
        actual[row["entidad"]] += row["transferencias"]
    if controls is None or dict(actual) != expected or controls != expected:
        raise ValueError("Rechazos: los motivos por entidad no concilian con los pagos no acreditados.")
    return rows


def build(path: Path) -> dict:
    book = openpyxl.load_workbook(path, data_only=True, read_only=False)
    try:
        if set(book.sheetnames) != {"Pagos", "Rechazos"}:
            raise ValueError("La fuente debe contener las hojas Pagos y Rechazos; revisar nuevas hojas antes de importar.")
        payments, totals = read_payments(book["Pagos"])
        rejections = read_rejections(book["Rechazos"], payments, totals)
        # A newly added temporal field requires explicit review. File properties
        # and timestamps never define the monitoring period.
        for sheet in book:
            for row in sheet:
                for cell in row:
                    if getattr(cell, "is_date", False) or (isinstance(cell.value, str) and any(
                            word in cell.value.casefold() for word in ("fecha de corte", "corte al", "al corte"))):
                        raise ValueError("La fuente incluye una referencia temporal nueva; revise su significado antes de importar.")
    finally:
        book.close()
    return {
        "version": 1, "fase": 4, "titulo": "Jóvenes en Acción · Cuarto período", "periodo": "Cuarto período",
        "corte": None, "fecha_disponible": False,
        "nota_temporal": "La fuente no informa fecha de corte ni meses de pago. Se presenta como referencia operativa "
                         "del cuarto período, separada de las series mensuales y del acumulado por fechas.",
        "fuente": {"nombre": path.name, "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                   "hojas": ["Pagos", "Rechazos"], "titulo_original": "CUARTO PERÍODO JÓVENES EN ACCIÓN"},
        "pagos": payments, "rechazos": rejections, "totales": totals,
        "control": {"entidades": len(payments), "motivos_rechazo": len({r["codigo_rechazo"] for r in rejections}),
                    "celdas_entidad_motivo": len(rejections), "pagos_conciliados": True,
                    "rechazos_conciliados_por_entidad": True, "porcentajes_fuente_verificados": True},
        "notas": {
            "unidad": "La fuente cuenta jóvenes seleccionados y autorizados, y transferencias acreditadas y no "
                      "acreditadas. No contiene identificadores para verificar personas únicas entre entidades o fases.",
            "entidades": "Las siglas corresponden a entidades ejecutoras según la fuente; no son entidades financieras.",
            "autorizacion": "Porcentaje de autorización: jóvenes autorizados / jóvenes seleccionados.",
            "acreditacion": "Porcentajes de acreditación y no acreditación: transferencias de cada estado / jóvenes autorizados.",
            "diferencia": "La diferencia entre seleccionados y autorizados se mantiene separada de las transferencias no acreditadas.",
            "montos": "Montos en USD tomados de la fuente; no representan devengado presupuestario.",
            "fases": "Este seguimiento del cuarto período no se suma a las fases 1–3 ni a meses sin fecha documentada.",
        },
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("base", type=Path)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    args = parser.parse_args()
    data = build(args.base)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"archivo": args.output.name, **data["control"], **data["totales"]}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
