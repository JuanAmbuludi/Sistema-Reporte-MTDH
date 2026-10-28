"""Actualización reproducible del corte abierto del 27/10/2026. Autor: Juan Ambuludi Maza.

Uso: python scripts/actualizar_enos_27_octubre.py --alertas archivo.xlsx
     --suspendidas archivo.xlsx --brigadas archivo.xlsx --kits archivo.xlsx
Los identificadores existentes se actualizan; los cortes cerrados no se modifican.
"""
import argparse
from collections import Counter
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import sys
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from src import names

def source(path):
    return {"nombre": Path(path).name, "sha256": hashlib.sha256(Path(path).read_bytes()).hexdigest()}

def update(alertas, suspendidas, brigadas, kits):
    folder = ROOT / 'data/processed'
    target = folder / 'seguimiento_enos_2026-10-27.json'
    closed = json.loads((folder / 'cortes_riesgos.json').read_text()).get('cerrados', {})
    if '2026-10-27' in closed:
        raise ValueError('El día está cerrado; abrir una nueva publicación antes de actualizar.')
    data = json.loads(target.read_text())
    old = {r['codigo']: r for r in data['alertas']}
    rows = list(openpyxl.load_workbook(alertas, data_only=True)['Detalle'].values)
    raw = [dict(zip(rows[3], row)) for row in rows[4:] if row[0]]
    assert len(raw) == len({r['Código'] for r in raw})
    summaries = {
      'ALT-2026-000364': 'Daños en techos y canaletas y presencia de goteras; mantenimiento solicitado.',
      'ALT-2026-000363': 'Daños en techos y canaletas y presencia de goteras; mantenimiento solicitado.',
      'ALT-2026-000362': 'Lluvias e inundación dificultan el acceso a cinco CDI de Salitre; dos reportan suspensión.',
      'ALT-2026-000361': 'Filtraciones de agua por el techo y humedad en paredes del centro gerontológico.',
      'ALT-2026-000360': 'Afectación de adecuaciones provisionales para mitigar filtraciones en áreas de adultos mayores.',
    }
    related = {'ALT-2026-000364':['GER-2026-001'], 'ALT-2026-000361':['GER-2026-005'],
               'ALT-2026-000360':['GER-2026-004'], 'ALT-2026-000362':[f'DII-2026-{n:03}' for n in range(14,19)]}
    for r in raw:
        code = str(r['Código SIIMDH'] or '').strip()
        if code in ('—','-'): code = ''
        old[r['Código']] = dict(codigo=r['Código'], fecha=r['Fecha'], provincia=names.place(r['Provincia']),
            canton=names.place(r['Cantón']), estado=r['Estado'], prioridad=r['Prioridad'], codigo_ua=code,
            unidad=' '.join(str(r['Unidad de atención']).split()), afectados_fuente=r['Afectados'],
            afectacion_servicio_confirmada=True, servicio='Desarrollo Infantil' if r['Código']=='ALT-2026-000362' else 'Centros gerontológicos',
            afectacion_servicio=summaries[r['Código']], origen='aplicativo', fecha_texto=r['Fecha'],
            mes_reporte='2026-10', fuente=Path(alertas).name, reportes_relacionados=related.get(r['Código'], []))
    data['alertas'] = list(old.values())
    data['control'].update(alertas_total=len(old), alertas_actualizacion_aplicativo=len(raw),
        encabezado_alertas_fuente=rows[1][0], alertas_detalle_recibido=len(raw))
    data['notas']['alertas'] = ('Seguimiento desde el 28/09/2026. Se cuentan registros de alerta del aplicativo y reportes sectoriales, '
        'no eventos ni unidades únicas. Las nuevas alertas se incorporan por su código oficial; las coincidencias con reportes '
        'sectoriales se conservan vinculadas, sin eliminar antecedentes. El archivo recibido anuncia 7 alertas, pero contiene 5 filas de detalle.')
    data['cortes']['alertas'] = '2026-10-27'
    # El archivo de kits es una actualización parcial: 26 y 27/10. Nunca sumar dos veces IDs presentes.
    w = openpyxl.load_workbook(kits, data_only=True)
    deliveries = {r['id_kobo']: r for r in data['kits']}
    kitrows = [r for r in list(w['Detalle completo'].values)[1:] if r[8]]
    assert len(kitrows) == len({str(r[8]) for r in kitrows})
    for r in kitrows:
        assert int(r[6]) == r[6] and r[6] >= 0 and int(r[7]) == r[7] and r[7] >= 0
        deliveries[str(r[8])] = dict(id_kobo=str(r[8]), fecha=datetime.strptime(r[0],'%d/%m/%Y').date().isoformat(),
            provincia=names.place(r[1]), canton=names.place(r[2]), kits=int(r[6]), familias=int(r[7]))
    data['kits'] = sorted(deliveries.values(), key=lambda r:(r['fecha'],r['provincia'],r['id_kobo']))
    data['cortes']['kits'] = '2026-10-27'
    data['kits_metadatos'].update(fuente=Path(kits).name, corte='27/10/2026 (sin hora indicada)',
        registros=len(deliveries), kits=sum(r['kits'] for r in deliveries.values()),
        familias_reportadas=sum(r['familias'] for r in deliveries.values()),
        registros_actualizacion=len(kitrows), modo_actualizacion='Actualización por ID de registro; conserva entregas anteriores',
        reemplaza_corte_anterior=False)
    for key in ('variacion_kits','registros_nuevos'):
        data['kits_metadatos'].pop(key, None)
    # Una brigada de emergencia, con su territorio y fecha, sin publicar un desglose adicional.
    brigade_rows = list(openpyxl.load_workbook(brigadas,data_only=True).active.values)[1:]
    data['brigadas'] = [dict(id=f'BRI-{r[0]}', fecha=r[1].date().isoformat(), provincia=names.place(r[6]),
        canton=names.place(r[7]), tipo='Emergencia') for r in brigade_rows
        if isinstance(r[1],datetime) and datetime(2026,9,28)<=r[1]<=datetime(2026,10,27)
        and str(r[11]).strip().upper()=='EMERGENCIA']
    data['cortes']['brigadas'] = '2026-10-22'
    assert len(data['brigadas']) == 1
    # El nuevo corte territorial reemplaza el inventario anterior de 99 unidades.
    records = [r for r in list(openpyxl.load_workbook(suspendidas,data_only=True)['DETALLE'].values)[1:] if r[0]]
    detail = []
    for row in records:
        prov, canton, parroquia, unidad, service, mode, admin, coverage = row
        if 'CNH' in mode: service, program = 'CNH','DII'
        elif 'CDI' in mode: service, program = 'CDI','DII'
        elif 'ACOMPAÑAMIENTO FAMILIAR' in mode: service, program = 'AF','AF'
        elif 'ADULTOS' in service: service, program = 'PAM','PAM_SERV'
        elif 'DISCAPACIDAD' in service: service, program = 'PCD','PCD_SERV'
        else: service, program = 'PE','PE'
        detail.append(dict(provincia=names.place(prov), canton=names.place(canton), parroquia=names.place(parroquia),
            servicio=service, modalidad=mode, unidad=unidad, codigo_ua='', estado='Suspendida', unidades=1,
            programa_id=program, corte_fuente='2026-10-27', cobertura=coverage))
    keys=[(r['provincia'],r['canton'],r['servicio'],r['unidad']) for r in detail]
    assert len(keys)==len(set(keys))==124
    counts=Counter((r['provincia'],r['servicio'],r['programa_id']) for r in detail)
    continuity=dict(corte='2026-10-27 19:00', fuente=Path(suspendidas).name, sha256=source(suspendidas)['sha256'],
        registros=[dict(provincia=p,servicio=s,programa_id=g,estado='Suspendida',unidades=n,corte_fuente='2026-10-27')
                   for (p,s,g),n in sorted(counts.items())], detalle=detail,
        fuentes_por_provincia={r['provincia']:'2026-10-27' for r in detail},
        validacion=dict(registros_detalle=124,suspendidas_total=124,cobertura_total=sum(r['cobertura'] for r in detail),
            criterio='Corte completo de unidades sin atención 27/10 a las 19:00. Sustituye las 99 del corte anterior; no se acumulan.'))
    data['generado_utc']=datetime.now(timezone.utc).isoformat()
    new_sources=[source(p) for p in (alertas,suspendidas,brigadas,kits)]
    replaced={s['nombre'] for s in new_sources}
    data['fuentes']=[s for s in data['fuentes'] if s['nombre'] not in replaced]+new_sources
    target.write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n')
    (folder/'continuidad_servicios_2026-10-27.json').write_text(json.dumps(continuity,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(dict(alertas=len(old),suspendidas=len(detail),brigadas=len(data['brigadas']),
        kits_periodo=sum(r['kits'] for r in deliveries.values() if r['fecha']>='2026-09-28'),
        kits_archivo=sum(r['kits'] for r in deliveries.values())),ensure_ascii=False))

if __name__=='__main__':
    p=argparse.ArgumentParser()
    for name in ('alertas','suspendidas','brigadas','kits'): p.add_argument('--'+name,required=True)
    update(**vars(p.parse_args()))
