#!/usr/bin/env python3
"""Convert the four aggregated legacy workbooks to Parquet, preserving cached values.

Run once on an older checkout. The original workbooks remain in Git history.
No nominal SANC input belongs in data/source.
"""
from pathlib import Path
import hashlib
import json
import re
import sys
import unicodedata

import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
KEEP = {
    "Usuarios_vis_vie": {"servicio", "modalidad", "id_provincia", "provincia", "id_canton", "canton", "id_parroquia", "parroquia", "sumsumusuario", "sumunidad", "cooperante", "fecha"},
    "Presupuesto_vis_vie": {"servicio", "vice", "id_provincia", "provincia", "id_canton", "id_parroquia", "sumusuarios", "total_devengado", "total_codificado", "codigo", "distribucion", "presupuesto_devengado_corriente", "presupuesto_devengado_inversion", "presupuesto_codificado_corriente", "presupuesto_codificado_inversion"},
}


def migrate(root=ROOT):
    source = root / "data/source"
    catalog = []
    for path in sorted(source.glob("*.xlsx")):
        book = load_workbook(path, read_only=True, data_only=True)
        for index, sheet in enumerate(book, 1):
            headers = list(next(sheet.iter_rows(max_row=1, values_only=True)))
            # Ignore formatted but empty Excel columns, including a sheet declared with 16,384 columns.
            width = next((i for i, v in enumerate(headers) if v is None), len(headers))
            headers = headers[:width]
            values = list(sheet.iter_rows(min_row=2, max_col=width, values_only=True))
            values = [r for r in values if any(v is not None for v in r)]
            if len(set(headers)) != len(headers):
                raise ValueError(f"Duplicate columns: {path.name}/{sheet.title}")
            table = pa.table({h: [row[i] for row in values] for i, h in enumerate(headers)})
            if path.stem in KEEP:
                table = table.select([c for c in table.column_names if c in KEEP[path.stem]])
            elif path.stem == "Base_de_datos_fin":
                table = table.drop(["% PROGRAMACIÓN"])
            table = table.replace_schema_metadata({b"source.sheet": sheet.title.encode(),
                                                  b"source.original_file": path.name.encode()})
            if len(book.sheetnames) > 1:
                slug = re.sub(r"[^a-z0-9]+", "_", unicodedata.normalize("NFKD", sheet.title).encode("ascii", "ignore").decode().lower()).strip("_")
                target = source / path.stem / f"{index:02d}_{slug}.parquet"
            else:
                target = path.with_suffix(".parquet")
            target.parent.mkdir(parents=True, exist_ok=True)
            pq.write_table(table, target, compression="zstd", compression_level=9)
            assert pq.read_table(target).equals(table)
            catalog.append({"original": path.name, "hoja": sheet.title,
                            "sha256_original": hashlib.sha256(path.read_bytes()).hexdigest(),
                            "parquet": str(target.relative_to(root)), "filas": len(values), "columnas": len(table.schema)})
        book.close()
    if catalog:
        (source / "catalogo_fuentes.json").write_text(json.dumps(catalog, ensure_ascii=False, indent=2)+"\n")
    return catalog


if __name__ == "__main__":
    print(json.dumps(migrate(), ensure_ascii=False, indent=2))
