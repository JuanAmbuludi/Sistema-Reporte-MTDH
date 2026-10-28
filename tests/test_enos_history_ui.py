"""Historical date ranges remain consistent from the dashboard through PDF and Excel."""
from datetime import date
import hashlib
from io import BytesIO
from pathlib import Path
import subprocess
import sys

import openpyxl
import pandas as pd
import pymupdf
import pytest

from src import enos
from src.reports import generate_pdf, generate_xlsx
from src.views.riesgos import RiskView

ROOT = Path(__file__).resolve().parents[1]
ARCHIVE_SHA256 = "da6fb7c6032c1824a41834e5a34867abe934e94fb201712af2a51737b9ffa905"


def view(*, mode="periodo", **selection):
    result = RiskView()
    result.select(result.catalog()[-1], mode=mode, **selection)
    assert result.prepare()
    return result


def values(result):
    return [card["value"] for card in result.kpi_cards()]


def pdf_text(content):
    with pymupdf.open(stream=content, filetype="pdf") as document:
        return "\n".join(page.get_text() for page in document)


def cell_values(sheet):
    return [cell.value for row in sheet.iter_rows() for cell in row if cell.value is not None]


def test_dashboard_defaults_to_history_and_preserves_both_date_controls():
    from test_interface import _logged_in_app, _markdown

    app = _logged_in_app()
    app.segmented_control(key="module").set_value("riesgos").run()
    assert not app.exception
    assert app.segmented_control(key="risk_mode").value == "Por período"
    assert app.date_input(key="risk_period_since").value == date(2026, 2, 27)
    assert app.date_input(key="risk_period_until").value == date(2026, 10, 27)
    assert "Del 27/02/2026 al 27/10/2026" in _markdown(app)

    app.date_input(key="risk_period_since").set_value(date(2026, 8, 28)).run()
    assert not app.exception
    html = _markdown(app)
    assert "Del 28/08/2026 al 27/10/2026" in html
    assert "14.581" in html and "1.171" in html and "354.987,39" in html
    assert "No hay una publicación guardada" not in html
    assert "Sin desglose temporal disponible" not in html

    app.segmented_control(key="risk_mode").set_value("Reportes publicados").run()
    app.date_input(key="risk_until").set_value(date(2026, 10, 26)).run()
    assert not app.exception
    assert "13.390" in _markdown(app)

    app.segmented_control(key="risk_mode").set_value("Por período").run()
    assert app.date_input(key="risk_period_since").value == date(2026, 8, 28)
    app.button(key="risk_all_history").click().run()
    assert not app.exception
    assert app.date_input(key="risk_period_since").value == date(2026, 2, 27)
    assert app.date_input(key="risk_period_until").value == date(2026, 10, 27)


def test_full_history_and_requested_range_do_not_clip_at_publication_start():
    full = view()
    assert (full.since, full.until) == (pd.Timestamp("2026-02-27"), pd.Timestamp("2026-10-27"))
    assert len(full.response.brigadas) == 71
    selected = view(since="2026-08-28", until="2026-10-27")
    assert values(selected) == [44, 14581, 1171, 354987.39, 124, 3]
    assert selected.response.since == pd.Timestamp("2026-08-28")
    assert selected.response.scope_text() == "Del 28/08/2026 al 27/10/2026"
    assert selected.response.kits.id_kobo.nunique() == len(selected.response.kits) == 117
    assert selected.response.kits.provincia.nunique() == 22
    assert "archived_risk_pdf" not in selected.report_context("Plan ENOS")


def test_single_day_before_first_publication_is_inclusive_and_exportable():
    selected = view(since="2026-08-31", until="2026-08-31")
    response = selected.response
    assert response.kits_total == 403
    assert len(response.kits) == 3
    assert pd.to_datetime(response.kits.fecha).eq(pd.Timestamp("2026-08-31")).all()
    assert response.alerts.empty
    assert response.bonus_period_available
    assert values(selected)[2:4] == [1171, 354987.39]
    context = selected.report_context("Plan ENOS")
    assert context["period_label"] == "Del 31/08/2026 al 31/08/2026"
    assert context["response"]["daily"] == [{"date": "31/08/2026", "short": "31/08", "daily": 0, "total": 0}]
    text = pdf_text(generate_pdf(context))
    assert "Del 31/08/2026 al 31/08/2026" in text and "403" in text
    assert "No hay una publicación guardada" not in text
    assert "Sin desglose temporal disponible" not in text


def test_period_pdf_and_excel_share_totals_rows_dates_and_unchanged_bonus():
    selected = view(since="2026-08-28", until="2026-10-27")
    context = selected.report_context("Plan ENOS")
    content = generate_pdf(context)
    text = pdf_text(content)
    assert "Del 28/08/2026 al 27/10/2026" in text
    for number in ("14.581", "1.171", "354.987,39", "124"):
        assert number in text
    assert "Brigadas de emergencia" in text
    assert all(row["codigo"] in text for row in selected.response.alerts.to_dict("records"))
    with pymupdf.open(stream=content, filetype="pdf") as document:
        assert len(document) == context["pdf_pages"]
        for page in document:
            for block in page.get_text("dict")["blocks"]:
                for line in block.get("lines", []):
                    for span in line["spans"]:
                        x0, y0, x1, y1 = span["bbox"]
                        assert 0 <= x0 <= x1 <= page.rect.width + 1
                        assert 0 <= y0 <= y1 <= page.rect.height + 1

    workbook = openpyxl.load_workbook(BytesIO(generate_xlsx(context)), data_only=True)
    assert workbook.sheetnames == selected.sheets()
    assert "Del 28/08/2026 al 27/10/2026" in cell_values(workbook["Ficha técnica"])
    summary = cell_values(workbook["Resumen"])
    assert 14581 in summary and 1171 in summary and 354987.39 in summary
    kits = cell_values(workbook["Entregas Kobo"])
    dates = [value for value in kits if isinstance(value, date)]
    assert len(dates) == 117
    assert min(dates).date() == date(2026, 8, 31)
    assert max(dates).date() == date(2026, 10, 27)
    kit_ids = {str(value) for value in kits}
    assert set(selected.response.kits.id_kobo.astype(str)) <= kit_ids
    alerts = {str(value) for value in cell_values(workbook["Alertas"])}
    assert set(selected.response.alerts.codigo) <= alerts


def test_full_history_pdf_handles_all_brigades_without_outside_page_text():
    selected = view()
    context = selected.report_context("Plan ENOS")
    assert len(context["response"]["brigadas"]) == 71
    with pymupdf.open(stream=generate_pdf(context), filetype="pdf") as document:
        text = "\n".join(page.get_text() for page in document)
        assert "Del 27/02/2026 al 27/10/2026" in text
        assert "Brigadas de emergencia" in text and "71" in text
        for page in document:
            for block in page.get_text("dict")["blocks"]:
                for line in block.get("lines", []):
                    for span in line["spans"]:
                        x0, y0, x1, y1 = span["bbox"]
                        assert 0 <= x0 <= x1 <= page.rect.width + 1
                        assert 0 <= y0 <= y1 <= page.rect.height + 1


def test_publication_mode_preserves_archive_and_has_separate_download_digest():
    published = view(mode="publicacion", until="2026-10-26")
    context = published.report_context("Plan ENOS")
    assert values(published) == [10, 13390, 1171, 354987.39, 99]
    assert hashlib.sha256(generate_pdf(context)).hexdigest() == ARCHIVE_SHA256
    same_dates = view(since=published.since, until=published.until)
    assert "archived_risk_pdf" not in same_dates.report_context("Plan ENOS")
    assert same_dates.digest() != published.digest()
    latest = view(mode="publicacion", until="2026-10-27")
    assert values(latest) == [44, 11207, 1171, 354987.39, 124, 1]


def test_publication_calendar_gap_keeps_effective_report_date(monkeypatch):
    import streamlit as st

    selected = view(mode="publicacion", until="2026-10-26")
    path = enos.in_force(enos.SNAPSHOT, "2026-10-26")[0]
    captions = []
    monkeypatch.setattr(st, "session_state", {})
    monkeypatch.setattr(st, "segmented_control", lambda *a, **k: "Reportes publicados")
    monkeypatch.setattr(st, "date_input", lambda *a, **k: date(2026, 10, 27))
    monkeypatch.setattr(st, "caption", captions.append)
    monkeypatch.setattr(enos, "publications", lambda *a: [(pd.Timestamp("2026-10-26"), path),
                                                       (pd.Timestamp("2026-10-28"), path)])
    monkeypatch.setattr(enos, "in_force", lambda *a: (path, True))
    period = selected.period_controls()
    assert period["until"] == pd.Timestamp("2026-10-26")
    assert any("Se muestra el reporte publicado el 26/10/2026" in text for text in captions)


@pytest.mark.parametrize("arguments", [
    ["--publicacion", "2026-10-26", "--desde", "2026-08-28"],
    ["--publicacion", "2026-10-26", "--todo-historial"],
    ["--todo-historial", "--hasta", "2026-10-27"],
    ["--desde", "2026-10-27", "--hasta", "2026-08-28"],
    ["--publicacion", "2026-10-28"],
])
def test_pdf_cli_rejects_ambiguous_or_reversed_dates(arguments):
    process = subprocess.run([sys.executable, str(ROOT / "scripts/generar_pdf_enos.py"), *arguments],
                             cwd=ROOT, capture_output=True, text=True)
    assert process.returncode == 2


def test_pdf_cli_exports_exact_publication(tmp_path):
    target = tmp_path / "publicacion.pdf"
    process = subprocess.run([sys.executable, str(ROOT / "scripts/generar_pdf_enos.py"),
                              "--publicacion", "2026-10-26", "--salida", str(target)],
                             cwd=ROOT, capture_output=True, text=True)
    assert process.returncode == 0, process.stderr
    assert hashlib.sha256(target.read_bytes()).hexdigest() == ARCHIVE_SHA256


def test_pdf_cli_defaults_to_all_available_history(tmp_path):
    target = tmp_path / "historial.pdf"
    process = subprocess.run([sys.executable, str(ROOT / "scripts/generar_pdf_enos.py"), "--salida", str(target)],
                             cwd=ROOT, capture_output=True, text=True)
    assert process.returncode == 0, process.stderr
    text = pdf_text(target.read_bytes())
    assert "Del 27/02/2026 al 27/10/2026" in text
    assert "Brigadas de emergencia" in text and "71" in text
