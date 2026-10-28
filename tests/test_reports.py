"""Exports: the summary (one or two pages), the slides and the data workbook, built from the on-screen selection and
the territorial level the reader asks for."""
from io import BytesIO
import zipfile

import pymupdf
import pytest
from openpyxl import load_workbook
from pptx import Presentation

from src.data import filter_data, load_periods, load_territory
from src.gestion_laboral import load_labour
from src.institutional import catalog as budget_catalog, load_budget
from src.reports import build_report_context, generate_pdf, generate_pptx, generate_xlsx
from src.reports.national import budget_context

BANNED = ["N/D", "No aplica", "no disponible", "No disponible", "Excel", "fuente_", "no certifica",
          "no documentado", "registros", "Selección ", "deduplicación", "no sumar", "sin identificadores", "fase presupuestaria"]


@pytest.fixture(scope="module")
def periods():
    data, catalog = load_periods()
    return data, catalog


@pytest.fixture(scope="module")
def parishes():
    return load_territory()


@pytest.fixture(scope="module")
def budget():
    frame, _manifest = load_budget()
    period = budget_catalog(frame)[0]
    scoped = frame.loc[frame.periodo_id.eq(period["id"])]
    return budget_context(scoped, national=scoped, scope=scoped, title="Presupuesto institucional", period=period,
                          territory_label="Total institucional", programs_label="Todos los programas")


@pytest.fixture(scope="module")
def labour():
    from src.views.laboral import LabourView
    view = LabourView(load_labour())
    view.select(view.catalog()[0])
    assert view.prepare()
    return view.report_context("Gestión laboral")


@pytest.fixture(scope="module")
def contingencies():
    from src import contingencias
    from src.views.contingencias import ContingencyView
    view = ContingencyView(contingencias.monthly(), contingencias.territorial(), contingencias.catalog())
    view.select(next(p for p in view.catalog() if p["id"] == "2026-09"))
    assert view.prepare()
    return view.report_context("Contingencias y compensaciones")


@pytest.fixture(scope="module")
def youth():
    from src.views.jovenes import YouthView
    view = YouthView()
    view.select(view.catalog()[0])
    assert view.prepare()
    return view.report_context("Jóvenes en Acción")


@pytest.fixture(scope="module")
def social():
    from src import presupuesto_social
    from src.views.presupuesto import SocialBudgetView
    view = SocialBudgetView(*presupuesto_social.load_social_budget())
    view.select(view.catalog()[0])
    assert view.prepare()
    return view.report_context("Presupuesto social")


def _period(data, period_id):
    return data.loc[data.periodo_id.eq(period_id)]


def _context(data, parishes, section, period_id="2026-09", provinces=None, level="todo"):
    rows = filter_data(_period(data, period_id), secciones=[section])
    parish = parishes.loc[parishes.periodo_id.eq(period_id) & parishes.seccion.eq(section)]
    chosen = filter_data(parish, provincias=provinces) if len(parish) else parish
    return build_report_context(filter_data(rows, provincias=provinces), national=rows, selected_codes=set(provinces or []),
                                parishes=chosen, parishes_national=parish, level=level)


def _strings(context: dict) -> list[str]:
    texts = [context["title"], context["period_label"], context["territory_label"], context["programs_label"]]
    texts += [i["text"] for i in context["insights"]]
    texts += [k["label"] for k in context["kpis"]] + [str(k.get("note", "")) for k in context["kpis"]]
    texts += [v for v in context.get("labels", {}).values() if isinstance(v, str)]
    return texts


def _pdf_text(content: bytes, pages: int = 2) -> str:
    document = pymupdf.open(stream=content, filetype="pdf")
    assert document.page_count == pages, "El resumen se mantiene en una o dos páginas"
    return "\n".join(page.get_text() for page in document)


def _deck_texts(content: bytes):
    deck = Presentation(BytesIO(content))
    texts = ["".join(run.text for p in shape.text_frame.paragraphs for run in p.runs)
             for slide in deck.slides for shape in slide.shapes if shape.has_text_frame]
    return deck, texts


def test_context_follows_module_and_period(periods, parishes):
    data, _ = periods
    current = _context(data, parishes, "bonos_pensiones")
    assert current["module_label"] == "Transferencias monetarias no contributivas" and current["period_label"] == "Septiembre 2026"
    assert current["totals"]["people"] == 1578019
    assert [k["label"] for k in current["kpis"]] == ["Beneficiarios", "Ejecución VIS–VIE", "Presupuesto ejecutado 2026",
                                                     "Presupuesto acumulado 2023–2026"]
    assert round(current["totals"]["execution"], 2) == 71.02
    assert current["map_layer"] == "parroquia" and current["parish"]["count"] > 1000
    assert current["levels"] == ["provincia", "canton", "parroquia"] and current["detail_level"] == "parroquia"
    assert current["reach"]["provincia"] == {"count": 24, "half": 5, "total": current["reach"]["provincia"]["total"]}
    assert current["reach"]["parroquia"]["count"] == current["parish"]["count"]
    cantons = current["territories"]["canton"]
    assert cantons[0]["name"] == "Guayaquil" and sum(c["people"] for c in cantons) == current["reach"]["canton"]["total"]
    annual = build_report_context(filter_data(_period(data, "2025"), secciones=["proteccion_integral"]))
    assert annual["period_label"] == "Año 2025" and annual["labels"]["fiscal"] == "Presupuesto 2025"
    assert annual["totals"]["people"] == 443903 and annual["totals"]["units"] == 10906
    assert annual["map_layer"] == "provincia" and annual["parish"] is None


def test_transfers_and_services_are_never_mixed(periods):
    data, _ = periods
    with pytest.raises(ValueError, match="ámbitos distintos"):
        build_report_context(_period(data, "2026-09"))


def test_selected_province_keeps_its_national_place_and_its_parishes(periods, parishes):
    data, _ = periods
    context = _context(data, parishes, "proteccion_integral", provinces=["09"])
    assert context["territory_label"] == "Guayas"
    assert context["totals"]["people"] == 87600
    assert all(k["share"] is not None for k in context["kpis"])
    assert "primer lugar" in context["insights"][1]["text"]
    guayas = next(p for p in context["provinces"] if p["code"] == "09")
    assert guayas["selected"] and guayas["position"] == 1
    assert context["map_focus"] == ["09"]
    assert {r["province_code"] for r in context["parish"]["rows"]} == {"09"}
    assert sum(r["people"] or 0 for r in context["parish"]["rows"]) == 87600
    assert context["parish"]["top"][0]["parish"] == "Tarqui"


def test_documents_use_formal_wording_without_technical_notes(periods, parishes, budget, labour, contingencies, youth, social):
    data, catalog = periods
    contexts = [budget, labour, contingencies, youth, social]
    for period in catalog:
        for section in ["bonos_pensiones", "proteccion_integral"]:
            contexts.append(_context(data, parishes, section, period["id"]))
    for context in contexts:
        for text in _strings(context):
            for banned in BANNED:
                assert banned not in text, f"{banned!r} en {text!r}"


def test_empty_and_multi_period_selections_are_rejected(periods):
    data, _ = periods
    with pytest.raises(ValueError, match="No hay información"):
        build_report_context(data.iloc[0:0])
    both = filter_data(data.loc[data.periodo_id.isin(["2026-09", "2025"])], secciones=["bonos_pensiones"])
    with pytest.raises(ValueError, match="único periodo"):
        build_report_context(both)


def test_levels_follow_the_reader_choice(periods, parishes):
    data, _ = periods
    canton = _context(data, parishes, "proteccion_integral", provinces=["09"], level="canton")
    assert canton["levels"] == ["canton"] and canton["map_layer"] == "canton" and canton["map_focus"] == ["09"]
    assert canton["detail_level"] == "canton" and {d["province"] for d in canton["detail"]} == {"Guayas"}
    assert canton["reach_levels"] == ["canton"] and canton["reach"]["canton"]["count"] == 25
    province = _context(data, parishes, "bonos_pensiones", level="provincia")
    assert province["levels"] == ["provincia"] and province["map_layer"] == "provincia" and province["detail_level"] == "provincia"
    annual = build_report_context(filter_data(_period(data, "2025"), secciones=["bonos_pensiones"]), level="todo")
    assert annual["levels"] == ["provincia"] and annual["level"] == "provincia"


def test_executive_summary_keeps_to_two_pages(periods, parishes):
    data, _ = periods
    text = _pdf_text(generate_pdf(_context(data, parishes, "bonos_pensiones")))
    for expected in ["Transferencias monetarias no contributivas", "Beneficiarios por programa", "Ejecución VIS–VIE por programa",
                     "Aspectos destacados", "1.578.019", "71,0 %", "Beneficiarios por parroquia", "Detalle territorial",
                     "Provincias con más beneficiarios", "Cantones con más beneficiarios", "Parroquias con más beneficiarios",
                     "reúnen la mitad de los beneficiarios", "Página 2 de 2"]:
        assert expected in text, expected
    for banned in BANNED:
        assert banned not in text
    zoom = _pdf_text(generate_pdf(_context(data, parishes, "proteccion_integral", provinces=["09"], level="parroquia")))
    assert "Parroquias con más usuarios" in zoom and "Tarqui" in zoom and "Distribución por servicio" in zoom
    assert "Cantones con más usuarios" not in zoom, "Un solo nivel elegido: solo ese nivel"


def test_presentation_has_a_slide_per_level_with_icons_and_native_charts(periods, parishes):
    data, _ = periods
    deck, texts = _deck_texts(generate_pptx(_context(data, parishes, "proteccion_integral", provinces=["17"])))
    assert len(deck.slides) == 6
    for title in ["Cifras clave", "Resultados por provincia", "Resultados por cantón", "Resultados por parroquia",
                  "Ejecución VIS–VIE", "Resultados por servicio", "6 / 6"]:
        assert title in texts, title
    pictures = [shape for slide in deck.slides for shape in slide.shapes if shape.shape_type == 13]
    assert len(pictures) >= 14, "Iconos, logo y mapas"
    charts = [shape for shape in deck.slides[5].shapes if shape.has_chart]
    assert len(charts) == 1 and len(charts[0].chart.plots[0].categories) == 4
    for text in texts:
        for banned in BANNED:
            assert banned not in text
    single, texts = _deck_texts(generate_pptx(_context(data, parishes, "bonos_pensiones", level="canton")))
    assert len(single.slides) == 4 and "Resultados por cantón" in texts and "Resultados por parroquia" not in texts


def test_data_workbook_is_institutional_numeric_and_formula_driven(periods, parishes):
    data, _ = periods
    content = generate_xlsx(_context(data, parishes, "bonos_pensiones", provinces=["09"]))
    book = load_workbook(BytesIO(content))
    assert book.sheetnames == ["Resumen", "Programas", "Provincias", "Cantones", "Parroquias", "Detalle", "Ficha técnica"]
    values = [cell.value for sheet in book.worksheets for row in sheet.iter_rows() for cell in row if cell.value is not None]
    strings = [v for v in values if isinstance(v, str)]
    assert any(v.startswith("=SUBTOTAL(109,") for v in strings)
    assert 312265 in values
    for banned in BANNED:
        assert not any(banned in v for v in strings if not v.startswith("=")), banned
    detail = book["Detalle"]
    assert detail.auto_filter.ref and detail.freeze_panes
    provinces = book["Provincias"]
    guayas = [row for row in provinces.iter_rows(values_only=True) if "Guayas" in row]
    assert guayas and 1 in guayas[0]
    parish_rows = [row for row in book["Parroquias"].iter_rows(values_only=True) if row[1] == "Guayas"]
    assert len(parish_rows) > 40 and book["Parroquias"].auto_filter.ref
    canton_rows = [row for row in book["Cantones"].iter_rows(values_only=True) if row[1] == "Guayas"]
    assert len(canton_rows) == 25 and book["Cantones"].auto_filter.ref
    detail_headers = [c.value for row in book["Detalle"].iter_rows(max_row=12) for c in row if isinstance(c.value, str)]
    assert "Parroquia" in detail_headers and "Cantón" in detail_headers
    with zipfile.ZipFile(BytesIO(content)) as archive:
        assert 'name="MTDH Institucional"' in archive.read("xl/theme/theme1.xml").decode()


def test_services_workbook_names_its_sheet_after_services(periods):
    data, _ = periods
    context = build_report_context(filter_data(_period(data, "2024"), secciones=["proteccion_integral"]))
    book = load_workbook(BytesIO(generate_xlsx(context)))
    assert book.sheetnames[1] == "Servicios" and "Parroquias" not in book.sheetnames and "Cantones" not in book.sheetnames
    headers = [cell.value for row in book["Servicios"].iter_rows(max_row=12) for cell in row if isinstance(cell.value, str)]
    assert "Presupuesto 2024" in headers and "Unidades de atención" in headers


def test_reports_accept_rows_loaded_without_period_catalog():
    from src.data import load_data
    frame = filter_data(load_data(), secciones=["bonos_pensiones"], programas=["BDH"])
    context = build_report_context(frame)
    assert context["period_label"] == "Septiembre 2026"
    assert context["insights"][0]["text"].startswith("Bono de Desarrollo Humano registra")
    assert sum(p["people"] for p in context["provinces"]) == 619073
    # The zone in study has no polygon: it counts in the totals, not on the map.
    assert sum(context["map_values"].values()) == sum(p["people"] for p in context["provinces"] if p["code"] != "90")


def test_budget_exports_read_execution_by_program_and_zone(budget):
    assert [k["label"] for k in budget["kpis"]] == ["Presupuesto codificado", "Presupuesto devengado", "Ejecución presupuestaria",
                                                    "Saldo por devengar"]
    assert round(budget["stats"]["ejecucion"], 2) == 45.77 and round(budget["stats"]["devengado"], 2) == 978122861.16
    assert budget["zones"][0]["name"] == "Planta central" and set(budget["map_values"]) <= {f"Z{i}" for i in range(1, 10)}
    text = _pdf_text(generate_pdf(budget))
    for expected in ["Presupuesto institucional", "45,8 %", "Ejecución por zona", "Ejecución por programa presupuestario",
                     "Composición del gasto", "Unidades ejecutoras con mayor presupuesto codificado"]:
        assert expected in text, expected
    deck, texts = _deck_texts(generate_pptx(budget))
    assert len(deck.slides) == 4 and "Ejecución por zona" in texts and "Ejecución por programa" in texts
    assert "Composición del gasto" in texts
    book = load_workbook(BytesIO(generate_xlsx(budget)))
    assert book.sheetnames == ["Resumen", "Programas", "Zonas", "Grupos de gasto", "Detalle", "Ficha técnica"]
    names = [c.value for c in book["Programas"]["C"] if isinstance(c.value, str)]
    assert "Servicios de atención gerontológica" in names
    assert not any("PROTECCION" in (c.value or "") for row in book["Detalle"].iter_rows() for c in row if isinstance(c.value, str))


def test_labour_exports_keep_each_figure_with_its_reference(labour):
    # The cut is daily: the figures of the published cut, never a number frozen in the test.
    from src.ui import number
    stock = labour["stock"]
    assert stock["registrados"] > stock["vigentes"] > 2_000_000 and stock["empresas"] > 300_000
    assert labour["levels"] == ["provincia", "canton"]
    assert labour["update"]["frequency"] == "Actualización diaria"
    text = _pdf_text(generate_pdf(labour), pages=2)
    for expected in ["Gestión laboral", number(stock["vigentes"]), number(stock["empresas"]), "Contratos vigentes por provincia",
                     "Perfil de los contratos vigentes",
                     "Compensación jubilar por año", "Contratos registrados por mes en 2026", "El corazón de Ecuador."]:
        assert expected in text, expected
    deck, texts = _deck_texts(generate_pptx(labour))
    assert len(deck.slides) == 5 and "Compensación jubilar" in texts and "Perfil y actividad económica" in texts
    book = load_workbook(BytesIO(generate_xlsx(labour)))
    assert book.sheetnames == ["Resumen", "Provincias", "Cantones", "Perfil y actividad", "Jubilación", "Contratos por mes", "Ficha técnica"]
    links = [c.hyperlink.target for row in book["Ficha técnica"].iter_rows() for c in row if c.hyperlink]
    assert any("sgde-indicadores.github.io" in (link or "") for link in links)


def test_contingency_exports_read_each_program_with_its_criterion(contingencies):
    assert contingencies["period_label"] == "Septiembre 2026" and contingencies["levels"] == ["provincia", "canton", "parroquia"]
    assert [k["label"] for k in contingencies["kpis"]] == ["Orfandad por femicidio", "Cobertura de contingencias",
                                                           "Eventos naturales y antrópicos", "Presupuesto devengado 2026"]
    assert round(contingencies["budget"]["execution"], 1) == 87.8
    text = _pdf_text(generate_pdf(contingencies), pages=2)
    for expected in ["Contingencias y compensaciones", "464", "3.403", "USD 83.681", "Bonos habilitados por provincia",
                     "Resultados por programa", "Evolución mensual de cada programa", "Pago acreditado", "Corte: septiembre de 2026"]:
        assert expected in text, expected
    deck, texts = _deck_texts(generate_pptx(contingencies))
    assert len(deck.slides) == 5 and "Resultados por programa" in texts and "Presupuesto" in texts
    book = load_workbook(BytesIO(generate_xlsx(contingencies)))
    assert book.sheetnames == ["Resumen", "Programas", "Serie mensual", "Provincias", "Cantones", "Parroquias", "Tipos", "Ficha técnica"]
    assert book["Parroquias"].max_row > 300, "El libro trae cada parroquia con bonos habilitados"


def test_open_months_are_not_published():
    from src import contingencias
    assert contingencias.closed_until() == "2026-09"
    assert all(p["id"] <= "2026-09" for p in contingencias.catalog())
    assert contingencias.monthly().periodo_id.max() == "2026-09" and contingencias.territorial().periodo_id.max() == "2026-09"


def test_youth_summary_fits_one_page(youth):
    assert youth["info"]["phase"] == 3 and [p["participants"] for p in youth["phases"]] == [75467, 80000, 150000]
    text = _pdf_text(generate_pdf(youth), pages=1)
    for expected in ["Jóvenes en Acción", "150.000", "637.808 jóvenes", "Presupuesto, ejecución y pagos por fase", "Pagos acreditados por mes",
                     "84,5 % de los habilitados"]:
        assert expected in text, expected
    deck, texts = _deck_texts(generate_pptx(youth))
    assert len(deck.slides) == 4 and "Fases del programa" in texts
    book = load_workbook(BytesIO(generate_xlsx(youth)))
    assert book.sheetnames == ["Resumen", "Fases", "Pagos por mes", "Presupuesto", "Ficha técnica"]


def test_social_budget_exports_read_the_pace_of_the_year(social):
    assert round(social["stats"]["execution"], 1) == 73.5 and len(social["programs"]) == 13
    assert round(social["closing"]["required"] / 1e6, 1) == 151.3
    text = _pdf_text(generate_pdf(social), pages=2)
    for expected in ["Presupuesto social", "USD 1.715,7 M", "Presupuesto devengado por mes", "Ritmo necesario de octubre a diciembre",
                     "Devengado de cada programa por mes", "Bono de Desarrollo Humano (BDH y BDHV)"]:
        assert expected in text, expected
    deck, texts = _deck_texts(generate_pptx(social))
    assert len(deck.slides) == 4 and "Detalle por programa" in texts
    book = load_workbook(BytesIO(generate_xlsx(social)))
    assert book.sheetnames == ["Resumen", "Programas", "Programa y mes", "Ficha técnica"]
