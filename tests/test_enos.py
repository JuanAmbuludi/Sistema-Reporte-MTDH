from io import BytesIO
import json
import pandas as pd
import pymupdf
from openpyxl import load_workbook
from pptx import Presentation
from src import enos
from src.views.enos import EnosView
from src.reports import generate_pdf, generate_pptx, generate_xlsx


def view(provinces=()):
    v=EnosView();v.data=enos.load(v.paths[-1]);v.period=v.catalog()[-1];v.provinces=provinces;v.prepare()
    return v


def test_alert_window_and_source_units():
    v=view();a=v.alerts
    assert len(a)==10 and a.codigo.nunique()==10
    assert pd.to_datetime(a.fecha).ge('2026-09-14').all()
    assert a.estado.value_counts().to_dict()=={'Pendiente':8,'En seguimiento':2}
    assert a.provincia.nunique()==3
    assert a.loc[a.codigo_ua.ne(''),'codigo_ua'].nunique()==10
    assert [k['value'] for k in v.kpi_cards()]==[10,10,10,802]
    assert set(a.columns)=={'codigo','fecha','provincia','canton','estado','prioridad','codigo_ua','unidad','afectados_fuente','afectacion_servicio_confirmada','servicio','afectacion_servicio'}
    assert v.data['control']['alertas_detalle_original']==31


def test_kit_deliveries_bono_scopes_and_merged_province():
    v=view()
    assert len(v.kits)==7
    assert v.kits.groupby('provincia').kits.sum().to_dict()=={'Azuay':197,'Esmeraldas':150,'Guayas':400,'Manabí':55}
    assert v.kits[['provincia','canton']].drop_duplicates().shape[0]==5
    a,b=v.frames['gestionados'],v.frames['articulacion']
    assert a.personas.sum()==113 and round(a.monto.sum(),2)==33200.80
    assert b.personas.sum()==1169 and round(b.monto.sum(),2)==354379.41
    assert b.loc[b.provincia.eq('Los Ríos'),'personas'].item()==109
    assert set(a.evento)|set(b.evento)=={'Inundación'}
    assert v.data['bonos']['planificacion']=={'personas':145523,'nucleos':44098,'monto':13405351.02}
    g=view(['Guayas']);assert len(g.alerts)==6 and g.kits.kits.sum()==400
    assert g.frames['gestionados'].personas.sum()==67
    assert g.digest()!=v.digest()


def test_three_products_keep_counts_and_explicit_limits():
    v=view();ctx=v.report_context('Seguimiento ENOS')
    pdf=pymupdf.open(stream=generate_pdf(ctx),filetype='pdf');assert len(pdf)==2
    text=' '.join(' '.join(p.get_text().split()) for p in pdf)
    for term in ['802','Los Ríos','145.523','14/09','articulación','no identifica pagos exclusivamente atribuibles a ENOS']:
        assert term in text
    ppt=Presentation(BytesIO(generate_pptx(ctx)));assert len(ppt.slides)==4
    wb=load_workbook(BytesIO(generate_xlsx(ctx)),data_only=True)
    assert wb.sheetnames==v.sheets()
    assert 802 in [c.value for row in wb['Resumen'] for c in row]
    assert 'Los Ríos' in [c.value for row in wb['Bonos'] for c in row]
    assert enos.BONO_NOTE in [c.value for row in wb['Ficha técnica'] for c in row]


def test_only_documented_service_impacts_reach_exports():
    v=view()
    expected={'ALT-2026-000'+n for n in ['325','323','322','321','319','318','310','309','316','315']}
    assert set(v.alerts.codigo)==expected
    assert v.alerts.afectacion_servicio_confirmada.sum()==8
    assert set(v.alerts.loc[~v.alerts.afectacion_servicio_confirmada, "codigo"])=={"ALT-2026-000315","ALT-2026-000316"}
    assert v.alerts.afectacion_servicio.str.len().gt(0).all()
    assert v.data['control']['alertas_excluidas']==21
    # A code alone is not evidence; legacy/unreviewed rows cannot reach any product.
    data=json.loads(json.dumps(v.data))
    data['alertas'].append({'codigo':'NEW','codigo_ua':'123','provincia':'Guayas'})
    assert set(enos.frames(data)['alertas'].codigo)==expected
    data['alertas']=[]
    assert enos.frames(data,['Guayas'])['alertas'].empty
    wb=load_workbook(BytesIO(generate_xlsx(v.report_context('Servicios MTDH'))),data_only=True)
    cells=[c.value for row in wb['Alertas'] for c in row]
    assert expected.issubset(set(cells)) and 'ALT-2026-000346' not in cells
    assert 'Afectación al servicio' in cells


def test_september_alerts_and_evolution_reconcile_with_review():
    from src.views.respuesta import Response, alerts_evolution
    v=view();a=v.alerts
    assert v.data['desde']=='2026-09-14'
    assert v.data['control']['alertas_desde_inicio']==31
    assert v.data['control']['alertas_sin_revision']==0
    september=a.loc[a.fecha.lt('2026-09-29')]
    assert september.empty
    assert a[['provincia','canton']].drop_duplicates().shape[0]==4
    r=Response(v.data);fig=alerts_evolution(r)
    assert pd.Timestamp(fig.data[0].x[0])==pd.Timestamp('2026-09-14')
    assert fig.data[0].y[-1]==10 and sum(fig.data[1].y)==10
    daily=dict(zip(pd.to_datetime(fig.data[1].x).strftime('%Y-%m-%d'),fig.data[1].y))
    assert daily['2026-09-14']==0 and daily['2026-09-24']==0 and daily['2026-09-27']==0
    data=json.loads(json.dumps(v.data))
    for code,day in [('BEFORE','2026-09-13'),('AFTER','2026-10-27')]:
        data['alertas'].append({**data['alertas'][0],'codigo':code,'fecha':day})
    assert len(enos.frames(data)['alertas'])==10
