"""The public risk products contain the response only, at the dashboard's grain."""
from io import BytesIO
import re
from zipfile import ZipFile
import openpyxl
import pymupdf
from pptx import Presentation
from src import riesgos
from src.reports import generate_pdf, generate_pptx, generate_xlsx
from src.views.riesgos import RiskView


def view(provinces=(), programs=(), threat=None, levels=riesgos.LEVELS):
    v = RiskView()
    v.select(v.catalog()[-1], provinces, programs, threat, levels)
    assert v.prepare()
    return v


def test_counts_distinguish_alerts_from_units_and_services():
    v = view(); r = v.response
    assert len(r.alerts) == r.alerts.codigo.nunique() == 21
    assert r.units == 10
    assert r.alerts.codigo_ua.eq("").sum() == 11
    assert r.alerts.groupby("servicio").size().to_dict() == {"Desarrollo Infantil": 8, "Creciendo con Nuestros Hijos (CNH)": 2, "Discapacidades · AHC": 5, "Centros gerontológicos": 6}
    cards = v.kpi_cards()
    assert [c["value"] for c in cards] == [21, 13390, 113, 33200.8, 99]
    assert "Núcleos" in cards[2]["label"]
    assert v.kpi_groups() == []
    # The exposure calculation is retained internally, outside the public context.
    assert riesgos.count_records(v.multihazard) == 6220
    assert not {"units", "exposed", "threats", "multihazard"}.intersection(v.report_context("Riesgos"))


def test_response_exports_are_complete_and_exclude_hidden_exposure():
    v = view(); ctx = v.report_context("Riesgos")
    pdf = pymupdf.open(stream=generate_pdf(ctx), filetype="pdf")
    assert len(pdf) == ctx["pdf_pages"] == 3
    pdf_text = " ".join(page.get_text() for page in pdf)
    slides = Presentation(BytesIO(generate_pptx(ctx))).slides
    ppt_text = " ".join(shape.text if shape.has_text_frame else " ".join(cell.text for row in shape.table.rows for cell in row.cells)
                        if shape.has_table else "" for slide in slides for shape in slide.shapes)
    book = generate_xlsx(ctx)
    wb = openpyxl.load_workbook(BytesIO(book), data_only=True)
    assert wb.sheetnames == v.sheets()
    # Dates are real dates (they sort and filter as dates) shown dd/mm/aaaa.
    xlsx_text = " ".join(c.value.strftime("%d/%m/%Y") if hasattr(c.value, "strftime") else str(c.value)
                         for sheet in wb for row in sheet for c in row if c.value is not None)
    kobo_dates = [c for row in wb["Entregas Kobo"].iter_rows() for c in row if hasattr(c.value, "strftime")]
    assert len(kobo_dates) == 89 and {c.number_format for c in kobo_dates} == {"dd/mm/yyyy"}
    for text in [pdf_text, ppt_text, xlsx_text]:
        for hidden in ["UA expuestas", "Multiamenaza", "Susceptibilidad muy alta", "Exposición de las unidades", "6 pendientes · 7 en seguimiento", "Gestionado no acredita cobro", "El código por sí solo no confirma una afectación"]:
            assert hidden not in text
        for a in ctx["response"]["alerts"]:
            assert a["code"] in text
            # All names and impacts survive wrapping, including the longest entries.
            compact = " ".join(text.split())
            assert " ".join(a["unit"].split()) in compact
            assert " ".join(a["impact"].split()) in compact
        assert "Núcleos" in text and "provincias priorizadas" in text
    assert "33.200,80" in pdf_text and "33.200,80" in ppt_text
    normalized_pdf = " ".join(pdf_text.split())
    for kit in ctx["response"]["kits"]:
        assert kit["province"] in normalized_pdf
        assert kit["date"] in xlsx_text and kit["canton"] in xlsx_text and kit["id"] in xlsx_text
    assert "12.681" in normalized_pdf and "99" in normalized_pdf
    assert "64 CNH" in normalized_pdf and "35 CDI" in normalized_pdf
    for value in ["145.523", "44.098", "13.405.351,02", "13.390", "113"]:
        assert value in normalized_pdf
    assert 33200.8 in [c.value for row in wb["Resumen"] for c in row]
    assert len(slides) == 12
    # One kind of native chart for each question: a doughnut (services), daily columns over the running total (alerts,
    # in the summary and in its sheet), a running total of the kits per delivery and a 100 % bar per stage of the bonus.
    with ZipFile(BytesIO(book)) as archive:
        kinds = sorted(" + ".join(re.findall(r"<c:(doughnutChart|barChart|lineChart|scatterChart)>", archive.read(name).decode()))
                       + (" 100 %" if b'<c:grouping val="percentStacked"/>' in archive.read(name) else "")
                       for name in archive.namelist() if name.startswith("xl/charts/chart"))
    assert kinds == ["barChart + lineChart", "barChart + lineChart", "barChart 100 %", "doughnutChart", "scatterChart"]
    # The suspended units carry no column the source leaves empty.
    headers = [c.value for row in wb["Unidades suspendidas"].iter_rows() for c in row if c.value in ("Código UA", "Unidad de atención")]
    assert headers == ["Unidad de atención"]


def test_filters_apply_consistently_to_response_and_do_not_depend_on_exposure():
    ctx = view(("09",), ("DII",), "susc_inundacion", ("Alta",)).report_context("Riesgos")
    assert ctx["response"]["alert_count"] == 6
    assert ctx["response"]["kits_total"] == 6850
    assert ctx["response"]["managed_people"] == 67
    assert sum(a["value"] for a in ctx["response"]["by_province"]) == 6
    assert sum(a["value"] for a in ctx["response"]["by_service"]) == 6
    empty = view((), ("PAM_SERV",), levels=()).report_context("Riesgos")
    assert empty["response"]["alert_count"] == 6 and empty["response"]["kits_total"] == 13390
    assert generate_pdf(empty)


def test_timeline_preserves_the_full_reporting_window():
    r = view().report_context("Riesgos")["response"]
    assert r["daily"][0] == {"date": "14/09/2026", "short": "14/09", "daily": 0, "total": 0}
    assert sum(a["daily"] for a in r["daily"]) == r["daily"][-1]["total"] == 15


def test_continuity_keeps_suspension_and_virtual_attention_separate():
    r = view().response
    assert r.closed_units == 99
    assert r.continuity_summary()["by_province"] == [{"name": "El Oro", "value": 75}, {"name": "Guayas", "value": 17}, {"name": "Santa Elena", "value": 7}]
    assert r.continuity.groupby("estado").unidades.sum().to_dict() == {"Suspendida": 36, "Suspendida presencial": 63}
    assert view(("07",)).response.closed_units == 75
    assert view(("09",)).response.closed_units == 17
    assert view(programs=("PAM_SERV",)).response.closed_units == 0


def test_kobo_snapshot_reconciles_and_is_not_added_to_the_old_cut():
    r = view().response
    assert len(r.kits) == r.kits.id_kobo.nunique() == 89
    assert r.kits_total == 13390
    assert r.kits.familias.sum() == 12681
    assert r.kits.provincia.nunique() == 20
    assert r.data["kits_metadatos"]["desde"] == "2026-08-31"
    assert r.data["kits_metadatos"]["hasta"] == "2026-10-26"
    assert sum(a["kits"] for a in r.kits_summary()) == r.kits_total
    assert r.continuity_summary()["by_service"] == [{"name": "CNH", "value": 64}, {"name": "CDI", "value": 35}]


def test_updated_units_have_complete_names_and_no_stale_summary_rows():
    v = view(); r = v.response
    assert len(r.continuity_detail) == 99
    assert r.continuity_detail.estado.value_counts().to_dict() == {"Suspendida presencial":63,"Suspendida":36}
    assert sum(a["units"] for a in r.suspended_groups()) == 99
    assert sum(a["units"] for a in r.suspended_groups() if a["service"] == "CNH") == 64
    doc = pymupdf.open(stream=generate_pdf(v.report_context("Riesgos · Respuesta del Ministerio")), filetype="pdf")
    text = " ".join(" ".join(p.get_text().split()) for p in doc)
    for a in r.suspended_groups():
        for name in a["names"].split("; "):
            assert name in text, name


def test_guayas_confirmation_preserves_original_source_status():
    r = view(("09",)).response
    assert r.closed_units == 17
    assert r.continuity_detail.estado.eq("Suspendida").all()
    assert r.continuity_detail.estado_fuente.eq("Afectada; modalidad por confirmar").all()


def test_names_read_as_the_gazetteer_writes_them_and_missing_values_are_a_dash():
    r = view().response
    # Kobo writes some cantons in capitals without accents; the alerts name their units in capitals.
    assert {"Espíndola", "Las Naves", "Pindal", "Quilanga", "Santa Ana"} <= set(r.kits.canton)
    assert not any(c.isupper() for c in r.kits.canton)
    assert set(r.alerts.unidad) >= {"Gotita del Amor", "CDI Los Niños del Cisne", "CNH Estrellitas del Amanecer",
                                    "Centro de Desarrollo Infantil Súper Niño"}
    # The CNH units of El Oro have no canton nor name in the source: empty, never a note about the base.
    groups = r.suspended_groups()
    assert groups[0] == {"province": "El Oro", "canton": "", "service": "CNH", "units": 63, "names": ""}
    ctx = view().report_context("Riesgos")
    pdf = " ".join(page.get_text() for page in pymupdf.open(stream=generate_pdf(ctx), filetype="pdf"))
    for note in ["Sin cantón en la base", "Sin nombre", "No consta en la fuente", "no disponibles en la base"]:
        assert note not in pdf


def test_documents_follow_the_content_of_the_selection():
    ctx = view(("17",)).report_context("Riesgos")  # Pichincha: kits only
    assert ctx["pdf_pages"] == 3
    pdf = pymupdf.open(stream=generate_pdf(ctx), filetype="pdf")
    text = " ".join(page.get_text() for page in pdf)
    assert len(pdf) == 3
    assert "Alertas reportadas" in text
    assert "Unidades cerradas / suspendidas" in text
    assert "64" in text and "Pichincha" in text
    assert ctx["response"]["alert_count"] == 0 and "21 alertas reportadas" not in text
    slides = Presentation(BytesIO(generate_pptx(ctx))).slides
    assert len(slides) == 6
    wb = openpyxl.load_workbook(BytesIO(generate_xlsx(ctx)), data_only=True)
    assert wb.sheetnames == view(("17",)).sheets()


def test_scopes_read_without_subtitles_and_anchor_the_geoportal():
    ctx = view().report_context("Riesgos")
    url = ctx["geoportal"]
    pdf = pymupdf.open(stream=generate_pdf(ctx), filetype="pdf")
    last = pdf[-1]
    text = " ".join(last.get_text().split())
    assert "Respuesta por provincia" in text and "Kits en provincias priorizadas" in text
    assert url in text
    assert [link["uri"] for link in last.get_links()] == [url]
    slides = Presentation(BytesIO(generate_pptx(ctx))).slides
    runs = [run for shape in slides[-1].shapes if shape.has_text_frame for p in shape.text_frame.paragraphs for run in p.runs]
    assert [run.hyperlink.address for run in runs if run.hyperlink.address] == [url]
    assert not any(run.text in ("Alertas", "Kits", "Bono de contingencia", "Filtros", "Continuidad") for run in runs)
    wb = openpyxl.load_workbook(BytesIO(generate_xlsx(ctx)))
    links = [c.hyperlink.target for row in wb["Ficha técnica"].iter_rows() for c in row if c.hyperlink]
    assert links == [url]
