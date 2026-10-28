"""Source controls supplied on 27 October, separate scopes and preserved coverage."""
from decimal import Decimal
from io import BytesIO

import pandas as pd
import pytest
from openpyxl import load_workbook

from src import finanzas
from src.data import load_data, load_periods, load_territory
from src.views.finanzas import FinanceView


def fiscal(**filters):
    frame, meta = finanzas.load_fiscal()
    view = FinanceView(frame, metadata=meta)
    view.select(view.catalog()[0], **filters)
    assert view.prepare()
    return view


def test_fiscal_matches_the_published_sgde_source():
    view = fiscal()
    expected = {"codificado": 993606130.45, "certificado": 46556969.04, "comprometido": 390954520.12,
                "devengado": 342231701.21, "saldo_disponible": 554031624.67, "saldo_por_devengar": 651374429.24}
    for key, value in expected.items():
        assert view.stats[key] == pytest.approx(value, abs=0.00001, rel=0)
    assert round(view.stats["ejecucion"], 2) == 34.44
    assert round(fiscal(tipo_gasto=["Corriente"]).stats["ejecucion"], 2) == 34.92
    assert round(fiscal(tipo_gasto=["Inversión"]).stats["ejecucion"], 2) == 32.12
    for dimension in view.dimensions:
        by = finanzas.by(view.filtered, dimension)
        for key in finanzas.FISCAL_MEASURES:
            assert by[key].sum() == pytest.approx(view.stats[key], abs=0.00001, rel=0)
    # Saldo disponible is a reported source measure, never back-calculated to force a balance.
    assert view.stats["saldo_disponible"] != view.stats["saldo_por_devengar"]


def test_fiscal_filters_and_exports_read_the_same_records():
    from src.reports import generate_pdf, generate_pptx, generate_xlsx
    import fitz
    from pptx import Presentation
    view = fiscal(programa=["Desarrollo Infantil"], tipo_gasto=["Corriente"])
    ctx = view.report_context(view.label)
    assert ctx["stats"] == view.stats
    assert len(ctx["tables"][0]["rows"]) == 1
    assert ctx["tables"][0]["rows"][0]["devengado"] == pytest.approx(view.stats["devengado"])
    pdf = fitz.open(stream=generate_pdf(ctx), filetype="pdf")
    assert len(pdf) == 2
    assert "Año fiscal 2026" in "".join(p.get_text() for p in pdf)
    assert len(Presentation(BytesIO(generate_pptx(ctx))).slides) == 3
    book = load_workbook(BytesIO(generate_xlsx(ctx)), data_only=True)
    assert "Detalle" in book and "Ficha técnica" in book
    figures = [cell.value for row in book["Programa"] for cell in row if isinstance(cell.value, (int, float))]
    assert any(abs(x-view.stats["devengado"]) < 0.001 for x in figures)
    assert any(abs(x-view.stats["ejecucion"]/100) < 0.000001 for x in figures)


def test_government_is_the_primary_workbook_and_includes_historical_programs():
    primary = finanzas.load_government()
    assert len(primary) == 296
    assert "usuarios" not in primary
    assert primary.anterior.sum() == pytest.approx(3332400360.6976, abs=0.00001, rel=0)
    assert primary.fiscal.sum() == pytest.approx(1180002196.4343, abs=0.00001, rel=0)
    assert primary.acumulado.sum() == pytest.approx(4512402557.1319, abs=0.00001, rel=0)
    assert primary.loc[primary.programa_id.isin(["PAM", "PCD"]), "acumulado"].sum() == pytest.approx(284102.54, abs=0.00001, rel=0)
    assert ((primary.anterior+primary.fiscal)-primary.acumulado).abs().max() < 0.000001
    current = load_data().set_index(["programa_id", "provincia_codigo"])
    expected = primary.set_index(["programa_id", "provincia_codigo"])
    pd.testing.assert_series_equal(current.monto_gobierno.sort_index(), expected.acumulado.sort_index(), check_names=False)


def test_coverage_and_social_execution_remain_validated():
    data, _ = load_periods()
    current = data.loc[data.periodo_id.eq("2026-09")]
    assert current.loc[current.seccion.eq("bonos_pensiones"), "usuarios"].sum() == 1578019
    assert current.loc[current.seccion.eq("proteccion_integral"), "usuarios"].sum() == 453994
    assert current.loc[current.seccion.eq("proteccion_integral"), "unidades"].sum() == 11215
    assert load_data().usuarios.sum() == 2034375
    annual = pd.read_parquet(finanzas.ROOT/"data/series/presupuesto_social_asignado.parquet")
    monthly = pd.read_parquet(finanzas.ROOT/"data/series/presupuesto_social_mensual.parquet")
    assert annual.presupuesto_aprobado.sum() == Decimal("1715698112.49")
    assert monthly.devengado_mes.sum() == Decimal("1261791379.87")


def test_government_dashboard_and_all_exports():
    from tests.test_interface import _logged_in_app, _markdown, _generate_every_product
    app = _logged_in_app()
    app.segmented_control(key="module").set_value("presupuesto").run()
    app.segmented_control(key="budget_universe").set_value("gobierno").run()
    assert not app.exception
    html = _markdown(app)
    assert "4.512,4" in html and "3.332,4" in html and "1.180,0" in html
    _generate_every_product(app)
