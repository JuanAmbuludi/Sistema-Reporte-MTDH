"""The fourth-phase snapshot keeps its filters and exports separate from phases 1–3."""
from io import BytesIO

import openpyxl
import pymupdf
import pytest
import streamlit as st
from pptx import Presentation

from src import jovenes_fase4 as source, ui
from src.reports import generate_pdf, generate_pptx, generate_xlsx
from src.views.jovenes import YouthPhasesView, YouthView
from src.views.jovenes_fase4 import PhaseFourView


def test_new_program_view_defaults_to_original_phases_even_with_saved_phase4(monkeypatch):
    monkeypatch.setattr(st, "session_state", {"youth_section": "fase4"})
    wrapper, previous = YouthView(), YouthPhasesView()
    assert isinstance(wrapper.active, YouthPhasesView)
    wrapper.select(wrapper.catalog()[0])
    previous.select(previous.catalog()[0])
    assert wrapper.prepare() and previous.prepare()
    assert wrapper.kpi_cards() == previous.kpi_cards()
    assert wrapper.report_context(wrapper.label)["kind"] == "jovenes"


@pytest.mark.parametrize("entities", [[], ["MIT"], ["MINEDEC", "MSP"]])
def test_snapshot_filters_reconcile_both_panels_and_have_no_calendar_period(entities):
    view = PhaseFourView()
    view.select(entities=entities)
    assert view.prepare()
    expected = source.summary(entities or None)
    assert view.stats == expected
    assert view.period["month"] is None and view.period["year"] is None
    assert view.payments.acreditadas.sum() + view.payments.no_acreditadas.sum() == expected["autorizados"]
    assert view.causes.transferencias.sum() == expected["no_acreditadas"]
    assert view.detail("detalle")[0].transferencias.sum() == expected["no_acreditadas"]
    context = view.report_context(view.label)
    assert context["kind"] == "area_simple"
    assert "sin fecha de corte" in context["tables_note"]
    assert sum(row["acreditadas"] for row in context["tables"][0]["rows"]) == expected["acreditadas"]


def test_snapshot_exports_keep_source_totals_dates_and_percentage_units():
    view = PhaseFourView()
    view.select()
    assert view.prepare()
    context = view.report_context(view.label)
    with pymupdf.open(stream=generate_pdf(context), filetype="pdf") as pdf:
        text = " ".join(page.get_text() for page in pdf)
        assert len(pdf) == 2
        for value in ["700.000", "685.867", "14.021", "Jóvenes en Acción", "Fase 4", "sin fecha de corte"]:
            assert value in text
        assert "Año a la fecha" not in text and "año None" not in text
    slides = Presentation(BytesIO(generate_pptx(context)))
    assert len(slides.slides) == 3
    book = openpyxl.load_workbook(BytesIO(generate_xlsx(context)), data_only=True)
    assert book.sheetnames == view.sheets()
    all_cells = [cell for row in book["Pagos y montos"].iter_rows(values_only=True) for cell in row]
    assert sum(value in all_cells for value in view.payments.entidad) == 7
    mit_row = next(row for row in book["Entidades"].iter_rows(values_only=True) if "MIT" in row)
    expected = float(view.payments.set_index("entidad").loc["MIT", "tasa_acreditacion"]) / 100
    assert any(isinstance(value, float) and value == pytest.approx(expected) for value in mit_row)


def test_program_subsection_filters_exports_and_returns_to_unchanged_prior_phases():
    from test_interface import _logged_in_app, _markdown

    legacy = YouthPhasesView()
    legacy.select(legacy.catalog()[0])
    assert legacy.prepare()
    app = _logged_in_app()
    app.segmented_control(key="module").set_value("jovenes").run()
    assert not app.exception
    assert app.segmented_control(key="youth_section").value == "fases"
    assert ui.number(legacy.info["enabled"]) in _markdown(app)
    app.segmented_control(key="youth_section").set_value("fase4").run()
    assert not app.exception
    assert "685.867" in _markdown(app) and "14.021" in _markdown(app)
    assert not any(widget.key in ("period_year", "period_month") for widget in app.segmented_control)
    assert len(app.get("plotly_chart")) == 6
    app.multiselect(key="youth4_entities").set_value(["MIT"]).run()
    assert not app.exception
    expected = source.summary(["MIT"])
    assert ui.number(expected["acreditadas"]) in _markdown(app)
    assert app.dataframe[0].value.entidad.tolist() == ["MIT"]
    app.segmented_control(key="youth4_level").set_value("Por motivo").run()
    assert not app.exception
    assert app.dataframe[0].value.transferencias.sum() == expected["no_acreditadas"]
    app.segmented_control(key="workspace").set_value("Reportes").run()
    for product in ("Resumen ejecutivo", "Presentación", "Libro de datos"):
        app.segmented_control(key="report_product").set_value(product).run()
        app.button(key="generate_report").click().run()
        assert not app.exception
        assert any(button.key == "download_report" for button in app.download_button)
        content = app.session_state["generated_report"]["content"]
        if product == "Resumen ejecutivo":
            with pymupdf.open(stream=content, filetype="pdf") as pdf:
                text = " ".join(page.get_text() for page in pdf)
                assert ui.number(expected["acreditadas"]) in text
        elif product == "Libro de datos":
            book = openpyxl.load_workbook(BytesIO(content), data_only=True)
            cells = {str(cell) for row in book["Entidades"].iter_rows(values_only=True) for cell in row if cell is not None}
            assert set(source.entities()) & cells == {"MIT"}
    app.segmented_control(key="workspace").set_value("Tablero").run()
    app.segmented_control(key="youth_section").set_value("fases").run()
    assert not app.exception
    assert ui.number(legacy.info["enabled"]) in _markdown(app)
    assert "Jóvenes habilitados" in _markdown(app)
    assert app.segmented_control(key="period_month").value == legacy.period["id"]
