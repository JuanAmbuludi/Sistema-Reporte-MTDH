"""Design details that are easy to break silently: wording, number formats, highlights and table exports."""
from io import BytesIO
import zipfile

from openpyxl import load_workbook

from src import charts, names, theme, ui
from src.data import filter_data, load_periods
from src.indicators import MODULES
from src.reports import table_xlsx


def _transfers():
    data, _ = load_periods()
    return filter_data(data.loc[data.periodo_id.eq("2026-09")], secciones=["bonos_pensiones"])


def test_indicator_names_follow_the_period_and_stay_neutral():
    module = MODULES["transferencias"]
    assert [i.label for i in module.indicators(2026)] == ["Beneficiarios", "Presupuesto 2026", "Presupuesto acumulado 2023–2026"]
    assert [i.label for i in module.indicators(2026, executed=True, parishes=True)] == ["Beneficiarios", "Presupuesto ejecutado 2026"]
    assert module.money(2024)["previous"].label == "Presupuesto acumulado 2023"
    assert [i.label for i in MODULES["servicios"].indicators(2025)][:2] == ["Usuarios", "Unidades de atención"]
    for name, text in MODULES["servicios"].definitions(2026):
        assert "_" not in text and text.endswith("."), name


def test_numbers_follow_ecuadorian_format():
    assert ui.number(1578019) == "1.578.019"
    assert ui.number(55867299.384, 2) == "55.867.299,38"
    assert ui.pct(19.84) == "19,8 %"
    assert ui.number(None) == "—"


def test_place_names_from_upper_case_sources():
    assert ui.place_name("SAN MIGUEL DE LOS BANCOS") == "San Miguel de los Bancos"
    assert ui.place_name("ALFREDO BAQUERIZO MORENO (JUJAN)") == "Alfredo Baquerizo Moreno (Juján)"
    assert ui.place_name("Santo Domingo de los Tsáchilas") == "Santo Domingo de los Tsáchilas"
    assert names.place("SAN JOSE DE MINAS") == "San José de Minas"
    assert names.place("BARTOLOMÉ RUIZ (CÉSAR FRANCO CARRIÓN)") == "Bartolomé Ruiz"
    assert names.place("ANDRéS DE VERA") == "Andrés de Vera"
    assert names.place("JULIO E. MORENO") == "Julio E. Moreno" and names.place("10 DE AGOSTO") == "10 de Agosto"


def test_budget_categories_read_as_formal_spanish():
    assert names.budget_program("58", "SERVICIOS DE ATENCION GERONTOLOGICA") == "Servicios de atención gerontológica"
    assert names.budget_program("57", "X", short=True) == "Protección social a la familia"
    assert names.spending_group("78", "TRANSFERENCIAS O DONACIONES PARA INVERSION") == "Transferencias o donaciones para inversión"
    assert names.funding_source("202", "Prestamos Externos") == "Préstamos externos"
    assert names.zone("ZONA 5") == "Zona 5" and names.zone_code("ZONA 5") == "Z5" and names.zone_code("PLANTA CENTRAL") is None
    assert names.executing_unit("DIRECCION DISTRITAL O4D01-SAN PEDRO DE HUACA-TULCAN-MIES") == "Dirección Distrital 04D01 – San Pedro de Huaca, Tulcán"
    assert names.executing_unit("COORDINACION ZONAL 8 - MIES") == "Coordinación Zonal 8"


def test_ranking_marks_the_chosen_province_in_gold_with_its_position():
    people = MODULES["transferencias"].people_indicator()
    fig = charts.province_ranking(_transfers(), people, highlight={"09"})
    bar = fig.data[0]
    index = list(bar.y).index("<b>Guayas</b>")
    assert bar.marker.color[index] == theme.ORO
    assert "1.º" in bar.text[index]
    assert all(color in (theme.ORO, theme.LILA_200) for color in bar.marker.color)


def test_composition_matrix_leaves_out_programs_without_values():
    people = MODULES["transferencias"].people_indicator()
    fig = charts.composition_matrix(_transfers(), people)
    columns = list(fig.data[0].x)
    assert len(columns) == 6, "Las pensiones sin beneficiarios no ocupan columnas vacías"
    assert all("<br>" in c or len(c) <= 11 for c in columns)


def test_category_charts_use_one_institutional_colour():
    people = MODULES["transferencias"].people_indicator()
    fig = charts.category_bars(_transfers(), people)
    assert fig.data[0].marker.color == charts.BAR
    assert "Beneficiarios" not in (fig.layout.title.text or "")
    assert len(fig.data[0].y) == 6, "Las pensiones sin beneficiarios no aparecen como barras vacías"


def test_execution_chart_reads_against_the_total():
    from src.data import aggregate
    rows = aggregate(_transfers(), ["programa_id", "programa"]).sort_values("presupuesto_codificado", ascending=False)
    fig = charts.execution_bullets(rows, "programa", reference=71.0)
    assert len(fig.data[1].y) == 6 and max(fig.data[1].x) <= 100
    assert any("71,0 %" in (a.text or "") for a in fig.layout.annotations)


def test_table_export_is_an_institutional_sheet_with_live_totals():
    rows = [{"programa": "Bono de Desarrollo Humano", "usuarios": 619073, "participacion": 0.6, "unidades": None, "monto_fiscal": 10.5},
            {"programa": "Desarrollo Infantil Integral", "usuarios": 274975, "participacion": 0.4, "unidades": 6759, "monto_fiscal": None}]
    layout = [("Programa", "programa", "text", 36.0), ("Beneficiarios", "usuarios", "int", 14.0),
              ("Participación", "participacion", "pct", 12.0), ("Unidades de atención", "unidades", "int", 12.5),
              ("Presupuesto 2026 (USD)", "monto_fiscal", "usd", 18.0)]
    content = table_xlsx(rows, layout, title="Cuadro de resultados", subtitle="Septiembre 2026", period_label="Septiembre 2026")
    sheet = load_workbook(BytesIO(content))["Detalle"]
    flat = [cell.value for row in sheet.iter_rows() for cell in row if cell.value is not None]
    assert "Cuadro de resultados" in flat and "Beneficiarios" in flat
    assert any(isinstance(v, str) and v.startswith("=SUBTOTAL(109,") for v in flat)
    assert 619073 in flat and "—" in flat
    assert not any(v in ("N/D", "No aplica") for v in flat)
    assert sheet.auto_filter.ref and sheet.freeze_panes
    with zipfile.ZipFile(BytesIO(content)) as archive:
        assert 'name="MTDH Institucional"' in archive.read("xl/theme/theme1.xml").decode()
