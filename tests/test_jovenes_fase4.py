"""Fourth-period monitoring stays reconciled, undated and independent."""
from datetime import datetime

import openpyxl
import pytest

from scripts import preparar_jovenes_fase4 as ingest
from src import jovenes_fase4 as source


def workbook(tmp_path):
    path = tmp_path / "Transferencias 2035.xlsx"  # filename must not create a payment period
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Pagos"
    sheet["A1"] = "CUARTO PERÍODO JÓVENES EN ACCIÓN"
    for cell, value in {"A2": "Entidad Ejecutora", "B2": "Jóvenes seleccionados", "C2": "Jóvenes autorizados",
                        "F2": "Transferencias Acreditadas", "I2": "Transferencias No Acreditadas"}.items():
        sheet[cell] = value
    for line, (entity, selected, authorized, credited, rejected) in enumerate(
            [("A", 100, 90, 80, 10), ("B", 40, 40, 30, 10), ("Total", 140, 130, 110, 20)], 4):
        values = [entity, selected, authorized, authorized / selected, authorized * 400,
                  credited, credited / authorized, credited * 400, rejected, rejected / authorized, rejected * 400]
        for col, value in enumerate(values, 1):
            sheet.cell(line, col, value)
    rejects = book.create_sheet("Rechazos")
    rejects.append([None, "CUARTO PERÍODO JÓVENES EN ACCIÓN"])
    rejects.append(["Código", "Motivo del Rechazo", "A", "B", "Total"])
    rejects.append([3, "CUENTA NO EXISTE", 8, 2, 10])
    rejects.append([16, "NUMERO DE CEDULA O RUC NO CORRESPONDE", 2, 8, 10])
    rejects.append(["Total", None, 10, 10, 20])
    book.properties.created = datetime(2035, 1, 1)
    book.save(path)
    return path, book


def test_import_keeps_source_units_and_never_dates_the_snapshot_from_metadata(tmp_path):
    path, _ = workbook(tmp_path)
    data = ingest.build(path)
    assert data["corte"] is None and data["fecha_disponible"] is False
    assert data["fase"] == 4 and data["periodo"] == "Cuarto período"
    assert data["totales"]["seleccionados"] == 140
    assert data["totales"]["autorizados"] == 130
    assert data["totales"]["acreditadas"] == 110
    assert data["control"]["rechazos_conciliados_por_entidad"] is True
    assert set(data["pagos"][0]) == set(source.PAYMENT_COLUMNS)
    assert set(data["rechazos"][0]) == set(source.REJECTION_COLUMNS)


@pytest.mark.parametrize("sheet,cell,value,match", [
    ("Pagos", "B6", 141, "porcentaje|total"),
    ("Pagos", "H4", 31999, "montos"),
    ("Pagos", "G4", .80, "denominador"),
    ("Pagos", "C4", None, "numérico"),
    ("Pagos", "I4", -1, "no negativo"),
    ("Rechazos", "E3", 11, "motivo"),
])
def test_import_rejects_inconsistent_totals_rates_or_missing_cached_values(tmp_path, sheet, cell, value, match):
    path, book = workbook(tmp_path)
    book[sheet][cell] = value
    book.save(path)
    with pytest.raises(ValueError, match=match):
        ingest.build(path)


def test_rejection_reconciliation_checks_each_entity_even_when_grand_total_matches(tmp_path):
    path, book = workbook(tmp_path)
    book["Rechazos"]["C3"] = 7
    book["Rechazos"]["D3"] = 3
    book.save(path)
    with pytest.raises(ValueError, match="por entidad"):
        ingest.build(path)


def test_a_new_source_date_requires_explicit_review(tmp_path):
    path, book = workbook(tmp_path)
    book["Pagos"]["A8"] = datetime(2026, 10, 1)
    book.save(path)
    with pytest.raises(ValueError):
        ingest.build(path)


def test_shipped_totals_and_rejection_reasons_reconcile_without_a_calendar_period():
    data = source.load()
    result = source.summary()
    assert result["seleccionados"] == 700000
    assert result["autorizados"] == 699888
    assert result["acreditadas"] == 685867
    assert result["no_acreditadas"] == 14021
    assert result["monto_autorizado"] == 279955200
    assert result["monto_acreditado"] == 274346800
    assert result["monto_no_acreditado"] == 5608400
    assert result["sin_autorizacion"] == 112
    assert result["tasa_acreditacion"] == pytest.approx(685867 / 699888 * 100)
    assert source.metadata()["corte"] is None
    assert data["control"]["entidades"] == 7 and data["control"]["motivos_rechazo"] == 10
    assert len(source.rejections()) == 70
    counts = source.rejections().groupby("codigo_rechazo").transferencias.sum().to_dict()
    assert counts[3] == 6780 and counts[16] == 6253 and sum(counts.values()) == 14021
    assert source.rejections().groupby("entidad").transferencias.sum().to_dict() == (
        source.payments().set_index("entidad").no_acreditadas.to_dict())


def test_entity_filter_applies_to_both_tables_and_recalculates_weighted_rates():
    selected = source.select(["MDEP", "SNGR"])
    summary = selected["resumen"]
    assert set(selected["pagos"].entidad) == set(selected["rechazos"].entidad) == {"MDEP", "SNGR"}
    assert summary["acreditadas"] == 138134 + 193249
    assert summary["no_acreditadas"] == 4597 + 3811 == selected["rechazos"].transferencias.sum()
    assert summary["tasa_acreditacion"] == pytest.approx((138134 + 193249) / (142731 + 197060) * 100)
    assert summary["tasa_acreditacion"] != pytest.approx(selected["pagos"].tasa_acreditacion.mean())
    empty = source.select([])
    assert empty["pagos"].empty and empty["rechazos"].empty
    assert empty["resumen"]["tasa_acreditacion"] is None
    assert source.summary("MIT")["acreditadas"] == 176526


def test_loading_or_filtering_does_not_mutate_the_cached_register():
    before = source.load()
    altered = source.load()
    altered["pagos"][0]["acreditadas"] = -1
    table = source.payments()
    table.loc[:, "acreditadas"] = -1
    assert source.load() == before
