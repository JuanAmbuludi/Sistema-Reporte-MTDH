"""Source ingestion must not discard history or manufacture event identities."""
from datetime import datetime
import json
import sys

import openpyxl
import pandas as pd
import pytest

from scripts import build_enos, update_kobo_kits
from src import enos


def alert_source(tmp_path, events, decisions):
    path = tmp_path / 'alertas.xlsx'
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = 'Detalle'
    sheet.append(['Detalle'])
    sheet.append(['Resumen'])
    sheet.append([None])
    columns = ['Código', 'Fecha', 'Tipo', 'Código SIIMDH', 'Unidad de atención',
               'Provincia', 'Cantón', 'Estado', 'Prioridad', 'Afectados']
    sheet.append(columns)
    for event in events:
        sheet.append([event.get(key) for key in columns])
    workbook.save(path)
    review = tmp_path / 'revision.csv'
    pd.DataFrame(decisions, columns=['codigo', 'incluir', 'servicio', 'afectacion_servicio',
                                     'afectacion_confirmada']).to_csv(review, index=False)
    return path, review


def alert(identifier, when, ua=123):
    return {'Código': identifier, 'Fecha': when, 'Tipo': 'Desastre natural', 'Código SIIMDH': ua,
            'Unidad de atención': 'Unidad', 'Provincia': 'GUAYAS', 'Cantón': 'BALAO',
            'Estado': 'Pendiente', 'Prioridad': 'Baja', 'Afectados': 1}


def reviewed(identifier, include='si', confirmed='si'):
    return {'codigo': identifier, 'incluir': include, 'servicio': 'Desarrollo Infantil',
            'afectacion_servicio': 'Acceso afectado', 'afectacion_confirmada': confirmed}


def kobo_source(tmp_path, records):
    path = tmp_path / 'kobo.xlsx'
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = 'Detalle de entregas'
    sheet.append(['Fecha entrega', 'ID Kobo', 'Provincia', 'Cantón', 'Kits', 'Familias'])
    for identifier, when, count in records:
        sheet.append([when, identifier, 'GUAYAS', 'BALAO', count, count])
    total = sum(count for _, _, count in records)
    controls = workbook.create_sheet('Control y notas')
    for row in [('Corte', '30/09/2026 19:00:00'), ('Registros', len(records)),
                ('Kits', total), ('Familias reportadas', total), ('Provincias', 1)]:
        controls.append(row)
    provincial = workbook.create_sheet('Resumen provincial')
    provincial.append(['Provincia', 'Registros', 'Kits', 'Familias', 'Nota'])
    provincial.append(['GUAYAS', len(records), total, total, None])
    provincial.append(['TOTAL', len(records), total, total, None])
    workbook.save(path)
    return path


@pytest.fixture
def open_snapshot(tmp_path, monkeypatch):
    monkeypatch.setattr(enos, 'FOLDER', tmp_path)
    monkeypatch.setattr(enos, 'LEDGER', tmp_path / 'cortes_riesgos.json')
    target = tmp_path / 'seguimiento_enos_2026-09-30.json'
    target.write_text(json.dumps({'desde': '2026-09-28', 'corte': '2026-09-30',
                                  'kits': [], 'cortes': {}, 'fuentes': []}), encoding='utf-8')
    return target


def test_alert_ingestion_retains_early_history_and_complete_final_day(tmp_path):
    path, review = alert_source(tmp_path,
        [alert('A', '2026-08-28'), alert('B', '2026-09-30T23:59:59-05:00'), alert('FUTURE', '2026-10-01')],
        [reviewed('A'), reviewed('B')])
    rows, control = build_enos.alert_rows(path, '2026-09-28', '2026-09-30', review)
    assert [(row['codigo'], row['fecha']) for row in rows] == [('A', '2026-08-28'), ('B', '2026-09-30')]
    assert control['alertas_anteriores_inicio'] == 1
    assert control['alertas_sin_revision'] == 0


def test_unreviewed_early_alert_cannot_silently_disappear(tmp_path):
    path, review = alert_source(tmp_path, [alert('A', '2026-08-28')], [])
    with pytest.raises(ValueError, match='Falta revisión documental.*A'):
        build_enos.alert_rows(path, '2026-09-28', '2026-09-30', review)


def test_explicit_exclusions_are_preserved_separately_from_unconfirmed_impact(tmp_path):
    path, review = alert_source(tmp_path,
        [alert('EXCLUDED', '2026-08-28'), alert('NO-UA', '2026-09-01', ua=None), alert('KEEP', '2026-09-02')],
        [reviewed('EXCLUDED', include='no'), reviewed('NO-UA'), reviewed('KEEP', confirmed='no')])
    rows, control = build_enos.alert_rows(path, '2026-09-28', '2026-09-30', review)
    assert [row['codigo'] for row in rows] == ['KEEP']
    assert rows[0]['afectacion_servicio_confirmada'] is False
    assert control['alertas_excluidas_codigos'] == ['EXCLUDED', 'NO-UA']


@pytest.mark.parametrize('identifiers, message', [(['A', ' A '], 'repetidos'), ([None], 'identificador')])
def test_alert_identity_is_checked_before_date_selection(tmp_path, identifiers, message):
    path, review = alert_source(tmp_path, [alert(identifier, '2026-08-28') for identifier in identifiers], [reviewed('A')])
    with pytest.raises(ValueError, match=message):
        build_enos.alert_rows(path, '2026-09-28', '2026-09-30', review)


def test_kit_csv_preserves_early_rows_and_requires_source_ids(tmp_path):
    path = tmp_path / 'kits.csv'
    pd.DataFrame([{'id_kobo': '123.0', 'fecha': '2026-08-28', 'kits': 5, 'contacto': 'Dato privado'},
                  {'id_kobo': '456', 'fecha': '2026-09-30T23:59:59', 'kits': 2}]).to_csv(path, index=False)
    rows, control = build_enos.kit_rows(path, '2026-09-30')
    assert [(row['id_kobo'], row['fecha']) for row in rows] == [('123', '2026-08-28'), ('456', '2026-09-30')]
    assert control['kits_conciliables_por_id'] is True
    assert all('contacto' not in row for row in rows)
    pd.DataFrame([{'fecha': '2026-08-28', 'kits': 5}]).to_csv(path, index=False)
    with pytest.raises(ValueError, match='requiere id_kobo'):
        build_enos.kit_rows(path, '2026-09-30')
    rows, control = build_enos.kit_rows(path, '2026-09-30', allow_legacy=True)
    assert 'id_kobo' not in rows[0] and control['kits_conciliables_por_id'] is False


def test_kobo_keeps_pre_scope_history_and_normalizes_numeric_ids(tmp_path, open_snapshot):
    source = kobo_source(tmp_path, [('123.0', datetime(2026, 8, 28), 5),
                                    (456, datetime(2026, 9, 30, 18, 30), 2)])
    update_kobo_kits.update(source, open_snapshot)
    data = json.loads(open_snapshot.read_text())
    assert data['desde'] == '2026-09-28'
    assert [(row['id_kobo'], row['fecha']) for row in data['kits']] == [('123', '2026-08-28'), ('456', '2026-09-30')]
    assert data['kits_metadatos']['reemplaza_corte_anterior'] is True
    assert data['kits_metadatos']['desde'] == '2026-08-28'


@pytest.mark.parametrize('identifiers, message', [(['123', '123.0'], 'duplicados'), ([None], 'identificador')])
def test_invalid_kobo_id_does_not_modify_existing_snapshot(tmp_path, open_snapshot, identifiers, message):
    before = open_snapshot.read_bytes()
    source = kobo_source(tmp_path, [(identifier, datetime(2026, 8, 28), 5) for identifier in identifiers])
    with pytest.raises(ValueError, match=message):
        update_kobo_kits.update(source, open_snapshot)
    assert open_snapshot.read_bytes() == before


@pytest.mark.parametrize('script', ['build', 'kobo'])
def test_ingestion_refuses_closed_snapshot_before_processing_source(tmp_path, monkeypatch, script):
    checked = []

    def closed(path):
        checked.append(path)
        raise ValueError('Publicación cerrada')

    monkeypatch.setattr(enos, 'ensure_open', closed)
    with pytest.raises(ValueError, match='cerrada'):
        if script == 'kobo':
            update_kobo_kits.update(tmp_path / 'missing.xlsx', tmp_path / 'seguimiento_enos_2026-09-30.json')
        else:
            monkeypatch.setattr(build_enos, 'ROOT', tmp_path)
            monkeypatch.setattr(sys, 'argv', ['build_enos', '--alerts', str(tmp_path / 'missing.xlsx'),
                                            '--bonos', 'missing.docx', '--service-review', 'missing.csv',
                                            '--cut', '2026-09-30'])
            build_enos.main()
    assert len(checked) == 1
