"""Complete emergency-brigade register, independent of report date selections.

Only administrative identifiers, dates, province and canton leave the source
workbook. The source cut and the actual first/last event dates remain distinct.
Published daily ENOS snapshots are never rewritten by this importer.
"""
from __future__ import annotations

from collections import Counter
from datetime import date, datetime, timezone
import hashlib
import json
from pathlib import Path
import re
import unicodedata

import openpyxl

from src import names


ROOT = Path(__file__).resolve().parents[1]
CANONICAL = ROOT / "data/processed/brigadas_emergencia.json"
_COLUMNS = {"#": "id", "FECHA": "fecha", "PROVINCIA": "provincia",
            "CANTON": "canton", "TIPO DE BRIGADA": "tipo"}


def _key(value) -> str:
    raw = " ".join(str(value or "").split()).upper()
    return "".join(c for c in unicodedata.normalize("NFKD", raw)
                   if not unicodedata.combining(c))


def _date(value, label: str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        for pattern in ("%Y-%m-%d", "%d/%m/%Y"):
            try:
                return datetime.strptime(value.strip(), pattern).date()
            except ValueError:
                pass
    raise ValueError(f"{label}: fecha ausente o inválida.")


def _identifier(value, row_number: int) -> str:
    raw = str(value).strip()
    if not re.fullmatch(r"[0-9]+(?:\.0+)?", raw) or int(raw.split('.')[0]) <= 0:
        raise ValueError(f"Fila {row_number}: identificador administrativo inválido.")
    # Avoid converting a potentially large integer through a floating point value.
    return f"BRI-{int(raw.split('.')[0])}"


def _territory(value, row_number: int, field: str) -> str:
    if _key(value) in ("", "-", "—", "N/A", "NA", "NONE", "NAN"):
        raise ValueError(f"Fila {row_number}: falta {field}.")
    if not isinstance(value, str):
        raise ValueError(f"Fila {row_number}: {field} debe ser un nombre territorial.")
    return names.place(value)


def read_workbook(source, corte, *, sheet: str | None = None) -> dict:
    """Read every EMERGENCIA row; never apply a fixed start/end reporting window.

    Identical administrative records are deduplicated. Conflicting IDs, malformed
    emergency rows and event dates after the declared source cut fail the import.
    ``corte`` is supplied explicitly by the source owner, not inferred from mtime.
    """
    source = Path(source)
    cut = _date(corte, "Corte de la fuente")
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        sheets = [workbook[sheet]] if sheet else workbook.worksheets
        candidates = []
        for worksheet in sheets:
            for number, row in enumerate(worksheet.iter_rows(values_only=True), 1):
                keys = [_key(value) for value in row]
                if set(_COLUMNS).issubset(keys):
                    columns = {_COLUMNS[key]: keys.index(key) for key in _COLUMNS}
                    candidates.append((worksheet, number, columns))
                    break
                if number >= 20:
                    break
        if len(candidates) != 1:
            raise ValueError("Se requiere una única hoja con #, FECHA, PROVINCIA, CANTON y TIPO DE BRIGADA; use --sheet si hay varias.")
        worksheet, header_row, columns = candidates[0]
        records, duplicate_ids, types = {}, [], Counter()
        source_rows = emergency_rows = footer_rows = 0
        for number, row in enumerate(worksheet.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
            values = {key: row[index] for key, index in columns.items()}
            # Spreadsheet totals outside the five published fields are not events.
            if all(value is None or str(value).strip() == "" for value in values.values()):
                footer_rows += 1
                continue
            source_rows += 1
            kind = _key(values["tipo"])
            types[kind or "SIN TIPO"] += 1
            if kind != "EMERGENCIA":
                continue
            emergency_rows += 1
            identifier = _identifier(values["id"], number)
            event_date = _date(values["fecha"], f"Fila {number}")
            if event_date > cut:
                raise ValueError(f"Fila {number}: fecha de brigada posterior al corte de la fuente.")
            record = {"id": identifier, "fecha": event_date.isoformat(),
                      "provincia": _territory(values["provincia"], number, "provincia"),
                      "canton": _territory(values["canton"], number, "cantón"),
                      "tipo": "Emergencia"}
            if identifier in records:
                if records[identifier] != record:
                    raise ValueError(f"Fila {number}: {identifier} tiene registros contradictorios.")
                duplicate_ids.append(identifier)
            else:
                records[identifier] = record
        ordered = sorted(records.values(), key=lambda record: (record["fecha"], record["id"]))
        dates = [record["fecha"] for record in ordered]
        return {
            "schema": 1,
            "corte": cut.isoformat(),
            "desde": min(dates) if dates else None,
            "hasta": max(dates) if dates else None,
            "generado_utc": datetime.now(timezone.utc).isoformat(),
            "criterio": "Todas las filas de tipo EMERGENCIA de la fuente; sin filtro por período del reporte.",
            "brigadas": ordered,
            "fuentes": [{"nombre": source.name, "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
                         "corte": cut.isoformat(), "hoja": worksheet.title}],
            "control": {"filas_fuente": source_rows, "filas_emergencia": emergency_rows,
                        "registros": len(ordered), "duplicados_identicos": len(duplicate_ids),
                        "ids_duplicados": sorted(set(duplicate_ids)), "filas_totales_o_vacias": footer_rows,
                        "tipos_fuente": dict(sorted(types.items()))},
        }
    finally:
        workbook.close()


def load(path=CANONICAL) -> dict:
    """Load the complete operational register; missing optional data yield no rows."""
    path = Path(path)
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else {"brigadas": [], "fuentes": []}
