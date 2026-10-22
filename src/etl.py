"""Validated, reproducible import for the Acumulado VIS-VIE workbook.

The Excel is an aggregated statistical source, not a beneficiary register.
Its cached formulas are checked, never executed as code. Input files are not
modified. All output tables carry row-level provenance.
"""
from __future__ import annotations

import csv
import hashlib
import json
import math
import re
import unicodedata
from datetime import date, datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import load_workbook

from .data import summarize

SCHEMA_VERSION = "1.0.0"
CENT_TOLERANCE = 0.01
MONEY_TYPE = pa.decimal128(24, 6)
MONTHS = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"]

PROGRAMS = {
    "DESARROLLO INFANTIL INTEGRAL": ("DII", "Desarrollo Infantil Integral", "proteccion_integral"),
    "PERSONAS ADULTOS MAYORES - PEJ": ("PAM_SERV", "Personas adultas mayores", "proteccion_integral"),
    "PERSONAS CON DISCAPACIDAD": ("PCD_SERV", "Personas con discapacidad", "proteccion_integral"),
    "PROTECCION ESPECIAL": ("PE", "Protección Especial", "proteccion_integral"),
    "BONO MIL DIAS": ("BMD", "Bono 1.000 Días", "bonos_pensiones"),
    "BDH": ("BDH", "Bono de Desarrollo Humano", "bonos_pensiones"),
    "BDHV": ("BDHV", "Bono de Desarrollo Humano Variable", "bonos_pensiones"),
    "BJGL": ("BJGL", "Bono Joaquín Gallegos Lara", "bonos_pensiones"),
    "PENSION MMA": ("MMA", "Pensión Mis Mejores Años", "bonos_pensiones"),
    "PENSION TUV": ("PTV", "Pensión Toda Una Vida", "bonos_pensiones"),
    "PENSION PAM": ("PAM", "Pensión para Adultos Mayores", "bonos_pensiones"),
    "PENSION PCD": ("PCD", "Pensión para Personas con Discapacidad", "bonos_pensiones"),
}
PROVINCES = {
    "AZUAY": ("01", "Azuay"), "BOLIVAR": ("02", "Bolívar"), "CANAR": ("03", "Cañar"),
    "CARCHI": ("04", "Carchi"), "COTOPAXI": ("05", "Cotopaxi"), "CHIMBORAZO": ("06", "Chimborazo"),
    "EL ORO": ("07", "El Oro"), "ESMERALDAS": ("08", "Esmeraldas"), "GUAYAS": ("09", "Guayas"),
    "IMBABURA": ("10", "Imbabura"), "LOJA": ("11", "Loja"), "LOS RIOS": ("12", "Los Ríos"),
    "MANABI": ("13", "Manabí"), "MORONA SANTIAGO": ("14", "Morona Santiago"), "NAPO": ("15", "Napo"),
    "PASTAZA": ("16", "Pastaza"), "PICHINCHA": ("17", "Pichincha"), "TUNGURAHUA": ("18", "Tungurahua"),
    "ZAMORA CHINCHIPE": ("19", "Zamora Chinchipe"), "GALAPAGOS": ("20", "Galápagos"),
    "SUCUMBIOS": ("21", "Sucumbíos"), "ORELLANA": ("22", "Orellana"),
    "SANTO DOMINGO DE LOS TSACHILAS": ("23", "Santo Domingo de los Tsáchilas"),
    "SANTA ELENA": ("24", "Santa Elena"), "ZONA EN ESTUDIO": ("90", "Zona en estudio"),
}

SCHEMA = pa.schema([
    pa.field("periodo", pa.date32()), pa.field("anio", pa.int16()),
    pa.field("seccion", pa.string(), nullable=False), pa.field("programa_id", pa.string(), nullable=False),
    pa.field("programa", pa.string(), nullable=False), pa.field("programa_origen", pa.string(), nullable=False),
    pa.field("provincia_codigo", pa.string(), nullable=False), pa.field("provincia", pa.string(), nullable=False),
    pa.field("usuarios", pa.int64(), nullable=False), pa.field("unidades", pa.int64()),
    pa.field("unidades_origen", pa.int64(), nullable=False),
    pa.field("monto_fiscal", MONEY_TYPE, nullable=False), pa.field("monto_gobierno", MONEY_TYPE, nullable=False),
    pa.field("monto_anterior", MONEY_TYPE, nullable=False), pa.field("presupuesto_ejecutado", MONEY_TYPE),
    pa.field("fuente_hoja", pa.string(), nullable=False), pa.field("fuente_fila", pa.int32(), nullable=False),
], metadata={b"mtdh.schema_version": SCHEMA_VERSION.encode(), b"mtdh.grain": b"periodo/programa_id/provincia_codigo"})


def norm(value) -> str:
    text = str(value or "").strip()
    text = "".join(c for c in unicodedata.normalize("NFKD", text) if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", text).upper()


def number(value, context: str, *, count: bool = False, allow_none: bool = False):
    if value is None or value == "":
        if allow_none:
            return None
        raise ValueError(f"Valor faltante en {context}; vuelva a guardar el Excel con sus fórmulas calculadas.")
    if isinstance(value, bool):
        raise ValueError(f"Valor booleano inesperado en {context}.")
    try:
        result = float(value)
    except (ValueError, TypeError):
        raise ValueError(f"Se esperaba un número en {context}: {value!r}") from None
    if not math.isfinite(result) or result < 0:
        raise ValueError(f"Valor no finito o negativo en {context}: {value!r}")
    if count:
        if not result.is_integer():
            raise ValueError(f"El conteo debe ser entero en {context}: {value!r}")
        return int(result)
    return Decimal(str(value)).quantize(Decimal("0.000001"), rounding=ROUND_HALF_UP)


def parse_period(value: str) -> date:
    if not re.fullmatch(r"\d{4}-\d{2}", value):
        raise ValueError("El corte debe tener formato AAAA-MM, por ejemplo 2026-09.")
    try:
        return date.fromisoformat(value + "-01")
    except ValueError:
        raise ValueError("El corte contiene un año o mes inválido.") from None


def _check_headers(sheet, expected: list[str]) -> None:
    row = next(sheet.iter_rows(min_row=1, max_row=1, max_col=len(expected), values_only=True))
    actual = [norm(v) for v in row]
    if actual != [norm(v) for v in expected]:
        raise ValueError(f"Encabezados inesperados en '{sheet.title}'. Esperados: {expected}; recibidos: {list(row)}")


def _identity(service, province, sheet: str, row: int) -> dict:
    if norm(service) not in PROGRAMS:
        raise ValueError(f"Programa sin homologación en {sheet}!A{row}: {service!r}. Actualice PROGRAMS antes de importar.")
    if norm(province) not in PROVINCES:
        raise ValueError(f"Provincia sin homologación en {sheet}!B{row}: {province!r}.")
    pid, program, section = PROGRAMS[norm(service)]
    code, province_name = PROVINCES[norm(province)]
    return dict(seccion=section, programa_id=pid, programa=program, programa_origen=str(service).strip(),
                provincia_codigo=code, provincia=province_name, fuente_hoja=sheet, fuente_fila=row)


def _read_snapshot(sheet, year: int, period: date | None, layout: str) -> list[dict]:
    expected = {
        "current": ["servicio", "Provincia", f"acumulado 2023-{year-1}", "Usuarios", "unidades de atención", f"Presupuesto {year}", "acumulado"],
        "2025": ["servicio", "Provincia", "solo dic", "Presupuesto 2024", "Usuarios", "unidades de atención", "Presupuesto 2025", "acumulado 2023-2024"],
        "2024": ["servicio", "Provincia", "solo dic", "Usuarios", "unidades de atención", "Presupuesto 2024", "acumulado 2023-2024"],
    }[layout]
    _check_headers(sheet, expected)
    result, seen = [], set()
    for rownum, row in enumerate(sheet.iter_rows(min_row=2, max_col=len(expected), values_only=True), 2):
        if all(v is None for v in row):
            continue
        rec = _identity(row[0], row[1], sheet.title, rownum)
        key = (rec["programa_id"], rec["provincia_codigo"])
        if key in seen:
            raise ValueError(f"Clave duplicada {key} en '{sheet.title}', fila {rownum}; no se elimina automáticamente.")
        seen.add(key)
        context = f"{sheet.title}!fila {rownum}"
        if layout == "2025":
            users, units, fiscal, government = row[4:8]
            previous = number(row[2], context) + number(row[3], context)
        else:
            users, units, fiscal, government = row[3:7]
            previous = number(row[2], context)
        users = number(users, context + " Usuarios", count=True)
        units = number(units, context + " unidades", count=True)
        fiscal, government = number(fiscal, context + " Presupuesto"), number(government, context + " acumulado")
        if abs(float(previous + fiscal - government)) > CENT_TOLERANCE:
            raise ValueError(f"Acumulado no concilia en {context}: anterior + fiscal != acumulado.")
        rec.update(periodo=period, anio=year, usuarios=users, unidades=units if rec["seccion"] == "proteccion_integral" else None,
                   unidades_origen=units, monto_fiscal=fiscal, monto_gobierno=government, monto_anterior=previous,
                   presupuesto_ejecutado=None)
        if layout != "current":
            rec.update(anio_referencia=year, corte_documentado=False, etiqueta_corte=f"{year} · mes de corte no informado",
                       monto_solo_diciembre_origen=number(row[2], context),
                       presupuesto_2024_origen=number(row[3] if layout == "2025" else row[5], context),
                       encabezado_acumulado_origen=expected[-1])
        result.append(rec)
    if not result:
        raise ValueError(f"La hoja '{sheet.title}' no contiene registros válidos.")
    return result


def _read_auxiliary(sheet) -> list[dict]:
    _check_headers(sheet, ["servicio", "Provincia", "Canton", "Usuarios", "%", "Presupuesto prov", "Presupuesto prop", "val %", "val presupuesto"])
    records = []
    for rownum, row in enumerate(sheet.iter_rows(min_row=2, max_col=9, values_only=True), 2):
        if all(v is None for v in row):
            continue
        rec = _identity(row[0], row[1], sheet.title, rownum)
        rec.update(canton=str(row[2]).strip(), usuarios=number(row[3], f"{sheet.title}!D{rownum}", count=True),
                   proporcion=number(row[4], f"{sheet.title}!E{rownum}", allow_none=True),
                   presupuesto_provincial=number(row[5], f"{sheet.title}!F{rownum}", allow_none=True),
                   presupuesto_prorrateado=number(row[6], f"{sheet.title}!G{rownum}", allow_none=True),
                   validacion_proporcion=number(row[7], f"{sheet.title}!H{rownum}", allow_none=True),
                   validacion_presupuesto=number(row[8], f"{sheet.title}!I{rownum}", allow_none=True),
                   periodo=None, utilizable_corte_actual=False)
        records.append(rec)
    return records


def join_execution(records: list[dict], csv_path: str | Path, period: date) -> dict:
    """Attach explicitly supplied execution; never infer it from Presupuesto."""
    lookup = {(r["programa_id"], r["provincia_codigo"]): r for r in records}
    seen, count = set(), 0
    with Path(csv_path).open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        required = {"periodo", "programa_id", "provincia_codigo", "presupuesto_ejecutado"}
        if not required.issubset(reader.fieldnames or []):
            raise ValueError("El CSV de ejecución requiere: " + ", ".join(sorted(required)))
        for lineno, row in enumerate(reader, 2):
            incoming = str(row["periodo"]).strip()
            if incoming not in {period.isoformat(), period.strftime("%Y-%m")}:
                raise ValueError(f"El período de ejecución en fila {lineno} no coincide con el corte {period:%Y-%m}.")
            key = (row["programa_id"].strip(), row["provincia_codigo"].strip().zfill(2))
            if key in seen:
                raise ValueError(f"Clave duplicada en ejecución, fila {lineno}: {key}")
            if key not in lookup:
                raise ValueError(f"Clave huérfana en ejecución, fila {lineno}: {key}")
            seen.add(key)
            lookup[key]["presupuesto_ejecutado"] = number(row["presupuesto_ejecutado"], f"ejecución fila {lineno}")
            count += 1
    if not count:
        raise ValueError("El CSV de ejecución solo tiene encabezados; complete al menos un registro o retire --ejecucion.")
    return {"archivo": Path(csv_path).name, "sha256": _sha256(csv_path), "filas": count, "cobertura_pct": 100 * count / len(records)}


def _sha256(path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _write_table(records: list[dict], path: Path, schema: pa.Schema | None = None) -> None:
    table = pa.Table.from_pylist(records, schema=schema)
    tmp = path.with_suffix(".parquet.tmp")
    pq.write_table(table, tmp, compression="zstd", use_dictionary=True)
    tmp.replace(path)


def build_dataset(excel_path: str | Path, output_dir: str | Path, period: str, *,
                  period_source: str = "Corte indicado por el operador; no consta en el Excel.",
                  execution_csv: str | Path | None = None, current_sheet: str = "provincial") -> dict:
    """Validate everything before writing a new current dataset."""
    excel_path, out = Path(excel_path), Path(output_dir)
    cut = parse_period(period)
    if not excel_path.is_file():
        raise FileNotFoundError(f"No se encuentra el Excel: {excel_path}")
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    try:
        if current_sheet not in workbook.sheetnames:
            raise ValueError(f"No existe la hoja vigente '{current_sheet}'. Hojas: {workbook.sheetnames}")
        records = _read_snapshot(workbook[current_sheet], cut.year, cut, "current")
        history = []
        for name, year in [("provincial (2025)", 2025), ("Hoja1 (2)", 2024)]:
            if name in workbook.sheetnames and name != current_sheet:
                history.extend(_read_snapshot(workbook[name], year, None, str(year)))
        auxiliary = _read_auxiliary(workbook["calculos dic"]) if "calculos dic" in workbook.sheetnames else []
        sheet_inventory = [{"hoja": s.title, "filas_excel": s.max_row, "columnas_declaradas": s.max_column} for s in workbook]
    finally:
        workbook.close()
    execution = join_execution(records, execution_csv, cut) if execution_csv else None
    frame = pd.DataFrame(records)
    for col in ["monto_fiscal", "monto_gobierno", "monto_anterior", "presupuesto_ejecutado"]:
        frame[col] = pd.to_numeric(frame[col], errors="raise").astype(float)
    frame["unidades"] = frame["unidades"].astype("Int64")
    totals = summarize(frame)
    known_2025 = {(r["programa_id"], r["provincia_codigo"]): r["monto_gobierno"] for r in history if r["anio"] == 2025}
    linkage_mismatches = sum(abs(float(r["monto_anterior"] - known_2025[(r["programa_id"], r["provincia_codigo"])])) > CENT_TOLERANCE
                              for r in records if cut.year == 2026 and (r["programa_id"], r["provincia_codigo"]) in known_2025)
    warnings = [
        "Los conteos son agregados reportados por programa y provincia. Sin identificadores nominales no se pueden deduplicar personas entre programas, servicios o cortes.",
        "El mes de corte no consta en el Excel; se registra por parámetro explícito con su evidencia. No existe una serie mensual en este archivo.",
        "Presupuesto fiscal es el campo Presupuesto del año. La fuente no certifica si es devengado, pagado, codificado u otra fase; no se interpreta automáticamente como presupuesto ejecutado ni monto mensual transferido.",
        "El monto acumulado de gobierno combina el acumulado anterior y el presupuesto fiscal; no se suma nuevamente con esos componentes.",
        "Las referencias de 2024 y 2025 no informan mes de corte. Sus cifras no son comparables como variaciones interanuales ni como ejecución fiscal de igual período.",
        "Las unidades de bonos y pensiones se marcan como no aplicables (nulo); el cero original queda en unidades_origen.",
        "Zona en estudio se conserva con código interno 90; no se cuenta como una provincia adicional.",
        "La hoja calculos dic contiene un prorrateo cantonal de referencia, con período no documentado y datos incompletos. Se conserva separada y no se suma al corte provincial vigente.",
    ]
    if not execution:
        warnings.append("Presupuesto ejecutado no disponible: use un CSV complementario verificado para completar ese indicador.")
    elif execution["cobertura_pct"] < 100:
        warnings.append("La ejecución complementaria es parcial; los totales solo se muestran cuando todos los registros seleccionados cuentan con ejecución.")
    if linkage_mismatches:
        warnings.append(f"Hay {linkage_mismatches} diferencias entre acumulado anterior vigente y la referencia 2025; revisar antes de compararlos.")
    definitions = {
        "periodo": "Primer día del mes de corte indicado por el operador; no representa fecha de transacción.",
        "usuarios": "Conteo agregado reportado en Usuarios; en bonos se presenta como beneficiarios reportados. No equivale a usuarios únicos verificados.",
        "unidades": "Suma de unidades de atención reportadas en servicios; no aplica a bonos y pensiones.",
        "monto_fiscal": f"Valor de Presupuesto {cut.year}, en USD, sin fase presupuestaria certificada en el insumo.",
        "monto_anterior": f"Acumulado 2023–{cut.year-1} reportado, en USD.",
        "monto_gobierno": "Acumulado reportado de gobierno: monto_anterior + monto_fiscal, en USD.",
        "presupuesto_ejecutado": "Ejecución verificada, solo desde CSV complementario. Nulo significa no disponible, no cero.",
        "provincia_codigo": "Código provincial ecuatoriano de dos dígitos; 90 es una clave interna para Zona en estudio.",
        "fuente_hoja / fuente_fila": "Procedencia exacta del registro dentro del Excel original.",
    }
    manifest = {
        "schema_version": SCHEMA_VERSION, "generado_utc": datetime.now(timezone.utc).isoformat(),
        "periodo": cut.isoformat(), "periodo_label": f"{MONTHS[cut.month-1].capitalize()} {cut.year}",
        "periodo_origen": period_source, "fuente": {"archivo": excel_path.name, "sha256": _sha256(excel_path)},
        "granularidad": "Una fila por corte, programa y provincia; fuente agregada.",
        "moneda": "USD", "filas": len(records), "totales": totals,
        "advertencias": warnings, "definiciones": definitions, "ejecucion_complementaria": execution,
        "control_calidad": {"claves_duplicadas": 0, "valores_negativos": 0, "conteos_fraccionarios": 0,
                            "acumulados_no_conciliados": 0, "diferencias_acumulado_2025": linkage_mismatches,
                            "ejecucion_disponible_filas": execution["filas"] if execution else 0,
                            "precision_monetaria_decimales": 6, "tolerancia_conciliacion_usd": CENT_TOLERANCE},
        "historicos": [{"anio_referencia": year, "filas": sum(r["anio"] == year for r in history), "mes_corte": None}
                       for year in sorted(set(r["anio"] for r in history))],
        "hojas_origen": sheet_inventory,
        "archivos": {},
    }
    out.mkdir(parents=True, exist_ok=True)
    _write_table(records, out / "indicadores.parquet", SCHEMA)
    if history:
        history_schema = SCHEMA.append(pa.field("anio_referencia", pa.int16())).append(pa.field("corte_documentado", pa.bool_())).append(pa.field("etiqueta_corte", pa.string())).append(pa.field("monto_solo_diciembre_origen", MONEY_TYPE)).append(pa.field("presupuesto_2024_origen", MONEY_TYPE)).append(pa.field("encabezado_acumulado_origen", pa.string()))
        _write_table(history, out / "historico_referencia.parquet", history_schema)
    else:
        (out / "historico_referencia.parquet").unlink(missing_ok=True)
    if auxiliary:
        auxiliary_schema = pa.schema([
            pa.field("seccion", pa.string()), pa.field("programa_id", pa.string()), pa.field("programa", pa.string()),
            pa.field("programa_origen", pa.string()), pa.field("provincia_codigo", pa.string()), pa.field("provincia", pa.string()),
            pa.field("fuente_hoja", pa.string()), pa.field("fuente_fila", pa.int32()), pa.field("canton", pa.string()),
            pa.field("usuarios", pa.int64()), pa.field("proporcion", pa.decimal128(18, 6)),
            pa.field("presupuesto_provincial", MONEY_TYPE), pa.field("presupuesto_prorrateado", MONEY_TYPE),
            pa.field("validacion_proporcion", pa.decimal128(18, 6)), pa.field("validacion_presupuesto", MONEY_TYPE),
            pa.field("periodo", pa.date32()), pa.field("utilizable_corte_actual", pa.bool_()),
        ], metadata={b"mtdh.schema_version": SCHEMA_VERSION.encode(), b"mtdh.role": b"auxiliary_not_current"})
        _write_table(auxiliary, out / "auxiliar_cantonal.parquet", auxiliary_schema)
    else:
        (out / "auxiliar_cantonal.parquet").unlink(missing_ok=True)
    for filename in ["indicadores.parquet", "historico_referencia.parquet", "auxiliar_cantonal.parquet"]:
        path = out / filename
        if path.exists():
            manifest["archivos"][filename] = {"sha256": _sha256(path), "bytes": path.stat().st_size,
                                             "filas": pq.ParquetFile(path).metadata.num_rows}
    (out / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    return manifest
