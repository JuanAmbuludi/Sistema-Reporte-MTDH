"""Editable Excel sources for current ENOS queries; issued JSON reports stay immutable.

Only this adapter knows the workbook layout. Its output keeps the existing
alerts, kits, continuity and bonus contracts used by the dashboard and exports.
"""
from __future__ import annotations

from collections import Counter
from copy import deepcopy
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from functools import lru_cache
from pathlib import Path

import openpyxl

FILES = {name: f"{name}.xlsx" for name in ("alertas", "kits", "unidades_suspendidas", "contingencias")}
FORMAT = "ENOS_EXCEL_1"
ALERT_FIELDS = ["codigo", "fecha", "provincia", "canton", "estado", "prioridad", "tipo_alerta", "codigo_ua",
                "unidad", "servicio", "afectados_fuente", "afectacion_servicio_confirmada", "afectacion_servicio", "incluir"]
KIT_FIELDS = ["id_kobo", "fecha", "provincia", "canton", "kits", "familias", "incluir"]
UNIT_FIELDS = ["id_unidad", "corte_fuente", "provincia", "canton", "parroquia", "unidad", "codigo_ua", "servicio",
               "modalidad", "programa_id", "estado", "cobertura"]
BONUS_FIELDS = ["corte", "provincia", "evento", "nucleos_habilitados", "monto_habilitado", "nucleos_cobrados", "monto_cobrado"]
FIELDS = dict(alertas=ALERT_FIELDS, kits=KIT_FIELDS, unidades_suspendidas=UNIT_FIELDS, contingencias=BONUS_FIELDS)
SERVICES = {"CNH": "DII", "CDI": "DII", "AF": "AF", "PAM": "PAM_SERV", "PCD": "PCD_SERV", "PE": "PE"}
STATES = {"Pendiente", "En seguimiento", "Resuelto"}
PRIORITIES = {"Crítica", "Alta", "Media", "Baja", "Sin prioridad", "No reportada"}


def directory(processed: Path) -> Path:
    return Path(processed).parent / "enos_excel"


def paths(processed: Path) -> list[Path]:
    folder = directory(processed)
    present = [folder / filename for filename in FILES.values() if (folder / filename).exists()]
    if present and len(present) != len(FILES):
        missing = [name for name in FILES.values() if not (folder / name).exists()]
        raise ValueError('La base Excel está incompleta. Faltan: ' + ', '.join(missing))
    return present


def _error(kind, row, message):
    raise ValueError(f"{FILES[kind]} · Datos, fila {row}: {message}")


def _text(value):
    return str(value).strip() if value is not None else ""


def _identifier(value):
    if isinstance(value, (int, float)) and not isinstance(value, bool) and float(value).is_integer():
        return str(int(value))
    return _text(value)


def _date(value, label):
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds") if value.time().isoformat() != "00:00:00" else value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = _text(value)
    try:
        parsed = datetime.fromisoformat(text)
    except ValueError as exc:
        raise ValueError(f"{label}: use una fecha de Excel o AAAA-MM-DD.") from exc
    return _date(parsed, label)


def _number(value, label, *, integer=False, optional=False):
    if value is None or value == "":
        if optional:
            return None
        raise ValueError(f"{label}: falta el valor; un blanco no equivale a cero.")
    try:
        n = Decimal(str(value))
    except InvalidOperation as exc:
        raise ValueError(f"{label}: se requiere un número.") from exc
    if isinstance(value, bool) or not n.is_finite() or n < 0 or (integer and n != n.to_integral_value()):
        raise ValueError(f"{label}: se requiere un {'entero ' if integer else ''}no negativo.")
    if not integer and n != n.quantize(Decimal('0.01')):
        raise ValueError(f"{label}: máximo dos decimales.")
    return int(n) if integer else float(n)


def _bool(value, label, *, optional=False):
    if optional and value in (None, "", "Sin confirmar"):
        return None
    if value is True or _text(value).casefold() in {"sí", "si"}:
        return True
    if value is False or _text(value).casefold() == "no":
        return False
    raise ValueError(f"{label}: seleccione Sí o No" + (" o deje en blanco si no está confirmado." if optional else "."))


@lru_cache(maxsize=24)
def _read(path: str, modified: int, size: int):
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        if not {"Datos", "Control"}.issubset(book.sheetnames):
            raise ValueError(f"{Path(path).name}: faltan las hojas Datos o Control.")
        control = {}
        for row in book["Control"].iter_rows(min_row=2, max_row=8, max_col=2, values_only=True):
            if row[0]:
                control[_text(row[0])] = row[1]
        kind = control.get("tipo")
        if kind not in FILES or control.get("formato") != FORMAT:
            raise ValueError(f"{Path(path).name}: formato o tipo no reconocido.")
        cut = _date(control.get("corte"), f"{Path(path).name} · Control · corte")
        table = book["Datos"]
        headers = [_text(v) for v in next(table.iter_rows(min_row=1, max_row=1, values_only=True))]
        while headers and not headers[-1]:
            headers.pop()
        if headers != FIELDS[kind]:
            raise ValueError(f"{Path(path).name}: conserve los encabezados y su orden: {', '.join(FIELDS[kind])}.")
        rows = []
        for number, cells in enumerate(table.iter_rows(min_row=2), 2):
            if all(cell.value is None for cell in cells):
                continue
            if any(cell.data_type == "f" for cell in cells):
                _error(kind, number, "no use fórmulas en Datos; ingrese valores.")
            if any(cell.value is not None for cell in cells[len(headers):]):
                _error(kind, number, "hay valores fuera de las columnas establecidas.")
            rows.append((number, dict(zip(headers, [c.value for c in cells]))))
        return kind, cut, control, rows
    finally:
        book.close()


def read(processed: Path, kind: str):
    path = directory(processed) / FILES[kind]
    if not path.exists():
        return None
    st = path.stat()
    actual, cut, control, rows = _read(str(path.resolve()), st.st_mtime_ns, st.st_size)
    if actual != kind:
        raise ValueError(f"{path.name}: el tipo de la hoja Control debe ser {kind}.")
    return cut, deepcopy(control), deepcopy(rows)


def _unique(seen, identifier, kind, row):
    if not identifier or identifier in seen:
        _error(kind, row, f"identificador vacío o duplicado: {identifier!r}. Actualice la fila existente.")
    seen.add(identifier)


def _check_date(value, cut, kind, row, field):
    parsed = _date(value, f"{FILES[kind]} · fila {row} · {field}")
    if parsed[:10] > cut[:10]:
        _error(kind, row, f"{field} supera el corte de Control; actualice el corte o corrija la fecha.")
    return parsed


def _missing(previous, seen, kind):
    missing = set(previous) - seen
    if missing:
        raise ValueError(f"{FILES[kind]}: faltan {len(missing)} registros del histórico (p. ej. {sorted(missing)[0]}). "
                         "No borre filas anteriores; para excluir un registro use incluir=No.")


def overlay(data: dict, processed: Path) -> dict:
    """Excel owns current registers, while base records retain reviewed optional metadata."""
    paths(processed)
    result = deepcopy(data)
    for kind, id_field in (("alertas", "codigo"), ("kits", "id_kobo")):
        workbook = read(processed, kind)
        if workbook is None:
            continue
        cut, control, rows = workbook
        old = {r[id_field]: r for r in data[kind]}
        current, seen = [], set()
        for line, row in rows:
            identifier = _identifier(row[id_field])
            _unique(seen, identifier, kind, line)
            record = deepcopy(old.get(identifier, {}))
            before = deepcopy(record)
            record[id_field] = identifier
            record['fecha'] = _check_date(row['fecha'], cut, kind, line, 'fecha')[:10]
            for name in ('provincia', 'canton'):
                if not _text(row[name]):
                    _error(kind, line, f"falta {name}.")
                record[name] = _text(row[name])
            if kind == 'kits':
                for field in ('kits', 'familias'):
                    record[field] = _number(row[field], f"fila {line} · {field}", integer=True)
            else:
                for field in ('estado', 'prioridad', 'tipo_alerta', 'unidad', 'servicio', 'afectacion_servicio'):
                    record[field] = _text(row[field])
                record['codigo_ua'] = _identifier(row['codigo_ua'])
                record['afectados_fuente'] = _number(row['afectados_fuente'], f"fila {line} · afectados_fuente", integer=True, optional=True)
                record['afectacion_servicio_confirmada'] = _bool(row['afectacion_servicio_confirmada'], 'afectacion_servicio_confirmada', optional=True)
                if record['estado'] not in STATES or record['prioridad'] not in PRIORITIES:
                    _error(kind, line, 'estado o prioridad fuera del catálogo de la plantilla.')
                if not record['tipo_alerta'] or not record['servicio']:
                    _error(kind, line, 'tipo_alerta y servicio son obligatorios; use Sin servicio identificado cuando corresponda.')
                if not before:
                    record.update(origen='aplicativo', fuente=FILES[kind], fecha_texto=record['fecha'], mes_reporte=record['fecha'][:7])
                elif record['fecha'] != before.get('fecha'):
                    record.update(fecha_texto=record['fecha'], mes_reporte=record['fecha'][:7])
            if _bool(row['incluir'], 'incluir'):
                current.append(record)
        _missing(old, seen, kind)
        result[kind] = current
        result['cortes'][kind] = cut
        if kind == 'alertas' and current != data[kind]:
            result.setdefault('control', {})['alertas_total'] = len(current)
        if kind == 'kits' and current != data[kind]:
            result.setdefault('kits_metadatos', {}).update(registros=len(current), kits=sum(r['kits'] for r in current),
                familias_reportadas=sum(r['familias'] for r in current), reemplaza_corte_anterior=True,
                desde=min((r['fecha'] for r in current), default=None), hasta=max((r['fecha'] for r in current), default=None))
    workbook = read(processed, 'contingencias')
    if workbook:
        cut, control, rows = workbook
        groups, seen = {}, set()
        for line, row in rows:
            when = _check_date(row['corte'], cut, 'contingencias', line, 'corte')[:10]
            province, event = _text(row['provincia']), _text(row['evento'])
            if not province or not event:
                _error('contingencias', line, 'provincia y evento son obligatorios.')
            _unique(seen, (when, province, event), 'contingencias', line)
            values = {field: _number(row[field], f"fila {line} · {field}", integer=field.startswith('nucleos')) for field in BONUS_FIELDS[3:]}
            if values['nucleos_cobrados'] > values['nucleos_habilitados'] or values['monto_cobrado'] > values['monto_habilitado']:
                _error('contingencias', line, 'lo cobrado supera lo habilitado.')
            groups.setdefault(when, []).append((province, event, values))
        if not groups:
            raise ValueError('contingencias.xlsx: falta el saldo provincial. No se interpreta una hoja vacía como cero.')
        latest = max(groups)
        previous_provinces = {r['provincia'] for r in data['bonos'].get('articulacion', [])}
        for when, balances in sorted(groups.items()):
            provinces = {p for p, _, _ in balances}
            if previous_provinces - provinces:
                raise ValueError(f'contingencias.xlsx: corte {when} incompleto. Conserve todas las provincias, incluso con saldo cero.')
            previous_provinces = provinces
        values = groups[latest]
        for province in {p for p, _, _ in values}:
            events = [e for p, e, _ in values if p == province]
            if 'Todos los eventos' in events and len(events) > 1:
                raise ValueError(f'contingencias.xlsx: {province} mezcla un total con sus eventos; duplicaría el saldo.')
        bonus = {}
        for stage, count_field, amount_field in [('articulacion','nucleos_habilitados','monto_habilitado'), ('gestionados','nucleos_cobrados','monto_cobrado')]:
            bonus[stage] = [dict(provincia=p, evento=e, personas=v[count_field], monto=v[amount_field]) for p,e,v in values]
        changed = any(result['bonos'][stage] != records for stage, records in bonus.items())
        result['bonos'].update(bonus)
        result['cortes'].update(bonos=latest, bonos_gestionados=latest)
        if changed:
            result['bonos_periodos_confirmados'] = {}
            result['bonos_referencia_confirmada'] = {'articulacion': {'confirmacion': 'Saldo provincial confirmado en contingencias.xlsx.'}}
            result.setdefault('bonos_presentacion', {}).update(unidad='Familias habilitadas', estado='Habilitadas al pago', cobro_confirmado=True)
            result['notas']['bono'] = 'Bono de contingencia: saldo habilitado y cobrado del último corte de contingencias.xlsx. Los cortes no se suman.'
            result.setdefault('control', {})['bonos_revision'] = result['notas']['bono']
    return result


def observations(processed: Path, base: dict | None = None) -> list[dict]:
    """One full stock per observed cut. Reopened units do not remain suspended."""
    paths(processed)
    workbook = read(processed, 'unidades_suspendidas')
    if workbook is None:
        return []
    cut, control, rows = workbook
    base = base or {}
    previous = {(r['provincia'],r['canton'],r['servicio'],r['unidad']): r for r in base.get('detalle', [])}
    groups, seen, units_seen, ids_by_cut = {}, set(), set(), {}
    for line, row in rows:
        when = _check_date(row['corte_fuente'], cut, 'unidades_suspendidas', line, 'corte_fuente')
        identifier = _identifier(row['id_unidad'])
        if not identifier:
            _error('unidades_suspendidas', line, 'id_unidad es obligatorio y debe conservarse al cambiar de corte.')
        _unique(seen, (when, identifier), 'unidades_suspendidas', line)
        ids_by_cut.setdefault(when, set()).add(identifier)
        key = tuple(_text(row[field]) for field in ('provincia','canton','servicio','unidad'))
        if not all(key):
            _error('unidades_suspendidas', line, 'provincia, cantón, servicio y unidad son obligatorios.')
        _unique(units_seen, (when, *key), 'unidades_suspendidas', line)
        service, program = _text(row['servicio']), _text(row['programa_id'])
        if SERVICES.get(service) != program:
            _error('unidades_suspendidas', line, 'servicio y programa_id no concuerdan con el catálogo.')
        state = _text(row['estado'])
        if state not in ('Suspendida', 'Suspendida presencial', 'Reanudada'):
            _error('unidades_suspendidas', line, 'estado debe ser Suspendida, Suspendida presencial o Reanudada.')
        group = groups.setdefault(when, [])
        coverage = _number(row['cobertura'], f'fila {line} · cobertura', integer=True)
        if state == 'Reanudada':
            continue
        record = deepcopy(previous.get(key, {}))
        record.update({f: _text(row[f]) for f in UNIT_FIELDS if f not in ('id_unidad','corte_fuente','cobertura')})
        record.update(codigo_ua=_identifier(row['codigo_ua']), corte_fuente=when, cobertura=coverage, unidades=1)
        group.append(record)
    if not groups:
        raise ValueError('unidades_suspendidas.xlsx: falta el corte. Para informar cero, marque las unidades como Reanudada.')
    result = []
    previous_ids = set()
    for when, details in sorted(groups.items()):
        if previous_ids - ids_by_cut[when]:
            raise ValueError(f'unidades_suspendidas.xlsx: corte {when} incompleto. Copie el inventario anterior y marque Reanudada cuando corresponda.')
        previous_ids = ids_by_cut[when]
        counts = Counter((r['provincia'],r['servicio'],r['programa_id'],r['estado']) for r in details)
        observation = deepcopy(base)
        observation.update(corte=when, detalle=details,
            registros=[dict(provincia=p,servicio=s,programa_id=g,estado=e,unidades=n,corte_fuente=when) for (p,s,g,e),n in sorted(counts.items())],
            fuentes_por_provincia={r['provincia']:when for r in details})
        if details != base.get('detalle'):
            observation.setdefault('validacion', {}).update(registros_detalle=len(details), suspendidas_total=len(details),
                cobertura_total=sum(r['cobertura'] for r in details), criterio='Inventario completo del corte en unidades_suspendidas.xlsx; no se suman cortes.')
        result.append(observation)
    return result
