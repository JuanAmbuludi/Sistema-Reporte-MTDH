"""Editable Word bulletin with the institutional identity, native tables and maps.

Fonts are the Office families also used by the Excel workbook (Century Gothic
for headings, Calibri for text), so the document renders the same on any
computer with Microsoft Office.  Charts and maps are high-resolution images
built from the same report context.
"""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from src import theme
from .charts import bar_chart, map_png
from .context import money, number

ROOT = Path(__file__).resolve().parents[2]
LOGO = ROOT / "assets" / "brand" / "mtdh_logo.png"
DISPLAY = theme.OFFICE_DISPLAY
TEXT = theme.OFFICE_TEXT
PURPLE = theme.MORADO.lstrip("#")
NIGHT = theme.MORADO_900.lstrip("#")
GOLD = theme.ORO.lstrip("#")
INK = theme.TINTA.lstrip("#")
MUTED = "6E6880"
PALE = theme.NIEBLA.lstrip("#")
LAVENDER = theme.LAVANDA.lstrip("#")
SOFT = theme.LILA_50.lstrip("#")
LINE = "E7E2EF"
CONTENT_WIDTH = 7.1


def _shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    element = props.find(qn("w:shd"))
    if element is None:
        element = OxmlElement("w:shd")
        props.append(element)
    element.set(qn("w:val"), "clear")
    element.set(qn("w:color"), "auto")
    element.set(qn("w:fill"), fill)


def _borders(cell, **edges):
    """edges: top/bottom/left/right = (size_eighths, color) or None to remove."""
    props = cell._tc.get_or_add_tcPr()
    borders = props.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        props.append(borders)
    for name in ["top", "left", "bottom", "right"]:
        spec = edges.get(name, "none")
        edge = borders.find(qn("w:" + name))
        if edge is None:
            edge = OxmlElement("w:" + name)
            borders.append(edge)
        if spec == "none" or spec is None:
            edge.set(qn("w:val"), "nil")
        else:
            size, color = spec
            edge.set(qn("w:val"), "single")
            edge.set(qn("w:sz"), str(size))
            edge.set(qn("w:color"), color)


def _padding(cell, top=80, bottom=80, left=110, right=110):
    props = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for side, value in {"top": top, "bottom": bottom, "left": left, "right": right}.items():
        item = OxmlElement("w:" + side)
        item.set(qn("w:w"), str(value))
        item.set(qn("w:type"), "dxa")
        margins.append(item)
    props.append(margins)


def _font(run, size=10.5, bold=False, color=INK, family=TEXT, italic=False):
    run.font.name = family
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:eastAsia"), family)
    rpr.rFonts.set(qn("w:hAnsi"), family)
    rpr.rFonts.set(qn("w:ascii"), family)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def _para(container, text, size=10.5, color=INK, bold=False, after=6, before=0, family=TEXT, align=None, italic=False, style=None):
    p = container.add_paragraph(style=style) if style else container.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing = 1.12
    if align:
        p.alignment = align
    _font(p.add_run(str(text)), size=size, bold=bold, color=color, family=family, italic=italic)
    return p


def _cell_text(cell, text, size=9.5, bold=False, color=INK, family=TEXT, align=WD_ALIGN_PARAGRAPH.LEFT, after=0):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.0
    _font(p.add_run(str(text)), size=size, bold=bold, color=color, family=family)
    return p


def _heading(doc, text, size=13, before=12, after=6, color=NIGHT):
    p = _para(doc, text, size=size, color=color, bold=True, family=DISPLAY, before=before, after=after)
    p.paragraph_format.keep_with_next = True
    return p


def _table(doc, headers, rows, widths, total=None, align=None):
    align = align or ["left"] + ["right"] * (len(headers) - 1)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    for i, value in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width = Inches(widths[i])
        _shade(cell, PURPLE)
        _cell_text(cell, value, size=8.5, bold=True, color="FFFFFF", family=DISPLAY,
                   align=WD_ALIGN_PARAGRAPH.RIGHT if align[i] == "right" else WD_ALIGN_PARAGRAPH.LEFT)
    header = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(header)
    body = list(rows) + ([total] if total else [])
    for index, values in enumerate(body):
        cells = table.add_row().cells
        is_total = total is not None and index == len(body) - 1
        for i, value in enumerate(values):
            cells[i].width = Inches(widths[i])
            _shade(cells[i], LAVENDER if is_total else ("FFFFFF" if index % 2 == 0 else PALE))
            _cell_text(cells[i], value, size=9.5, bold=is_total or i == 0, color=NIGHT if (i == 0 or is_total) else INK,
                       align=WD_ALIGN_PARAGRAPH.RIGHT if align[i] == "right" else WD_ALIGN_PARAGRAPH.LEFT)
    for r, row in enumerate(table.rows):
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        last = total is not None and r == len(table.rows) - 1
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _borders(cell, top=(12, PURPLE) if last else None, bottom=(4, LINE))
            _padding(cell, 70, 70, 100, 100)
    return table


def _note(doc, text, label=""):
    box = doc.add_table(rows=1, cols=1)
    box.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = box.rows[0].cells[0]
    cell.width = Inches(CONTENT_WIDTH)
    _shade(cell, SOFT)
    _borders(cell, left=(24, PURPLE))
    _padding(cell, 110, 110, 160, 160)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    if label:
        _font(p.add_run(label + " "), size=9.5, bold=True, color=theme.MORADO_800.lstrip("#"), family=DISPLAY)
    _font(p.add_run(text), size=9.5, color="4C3A6A")
    _para(doc, "", after=2)


def _cover(doc, context):
    """A purple plate with title, period and the map, followed by KPI tiles."""
    plate = doc.add_table(rows=1, cols=2)
    plate.alignment = WD_TABLE_ALIGNMENT.CENTER
    plate.autofit = False
    widths = [4.1, 3.0]
    for column, width in zip(plate.columns, widths):
        column.width = Inches(width)
    left, right = plate.rows[0].cells
    for cell, width, fill in [(left, widths[0], NIGHT), (right, widths[1], PURPLE)]:
        cell.width = Inches(width)
        _shade(cell, fill)
        _borders(cell)
        _padding(cell, 200, 200, 220, 160)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    left.text = ""
    p = left.paragraphs[0]
    p.paragraph_format.space_after = Pt(16)
    if LOGO.exists():
        p.add_run().add_picture(str(LOGO), width=Inches(3.0))
    _para(left, "Boletín estadístico institucional", size=9.5, color="D9CCF0", bold=True, after=4)
    _para(left, context["title"], size=24, color="FFFFFF", bold=True, family=DISPLAY, after=6)
    _para(left, context["period_label"], size=14, color=GOLD, bold=True, family=DISPLAY, after=4)
    _para(left, "Alcance: " + context["scope_label"], size=9.5, color="E9E2F6", after=0)
    territory = context.get("territory", [])
    if territory:
        values = {t["code"]: t["people"] for t in territory}
        few = {t["code"] for t in territory if t["code"] != "90"}
        image = map_png(values, mode="night", selected=few if len(few) <= 3 else set(), annotate=0, legend=False,
                        scale="quantile", width=5.0, height=4.4)
        right.text = ""
        right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        right.paragraphs[0].add_run().add_picture(BytesIO(image), width=Inches(2.7))
    rule = doc.add_table(rows=1, cols=1)
    rule.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = rule.rows[0].cells[0]
    cell.width = Inches(CONTENT_WIDTH)
    _shade(cell, GOLD)
    _borders(cell)
    _padding(cell, 0, 0, 0, 0)
    cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    _font(cell.paragraphs[0].add_run(" "), size=2)
    row = rule.rows[0]
    trpr = row._tr.get_or_add_trPr()
    height = OxmlElement("w:trHeight")
    height.set(qn("w:val"), "70")
    height.set(qn("w:hRule"), "exact")
    trpr.append(height)


def _kpis(doc, context):
    tiles = []
    for section in context["sections"]:
        tone = theme.FUCSIA if section["id"] == "proteccion_integral" else theme.MORADO
        noun = "servicios" if section["id"] == "proteccion_integral" else "programas"
        tiles.append((section["people_label"], number(section["people"]), f"{len(section['rows'])} {noun} en la selección", tone))
    if context["units_total"] is not None:
        tiles.append(("Unidades de atención", number(context["units_total"]), "Servicios de protección integral", theme.VIOLETA_500))
    tiles.append((f"Monto fiscal {context['year']}", f"USD {money(context['fiscal_total'], True)} M",
                  f"Acumulado de gobierno: USD {money(context['government_total'], True)} M", theme.ORO))
    if len(tiles) % 2:
        executed = context["executed_total"]
        tiles.append(("Presupuesto ejecutado", "N/D" if executed is None else f"USD {money(executed, True)} M",
                      "La fuente no certifica la fase de ejecución" if executed is None else "Ejecución verificada", theme.LILA_300))
    grid = doc.add_table(rows=(len(tiles) + 1) // 2, cols=2)
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    grid.autofit = False
    for index, (label, value, note, tone) in enumerate(tiles):
        cell = grid.rows[index // 2].cells[index % 2]
        cell.width = Inches(CONTENT_WIDTH / 2)
        _shade(cell, PALE)
        _borders(cell, top=(24, tone.lstrip("#")), left=(24, "FFFFFF"), right=(24, "FFFFFF"), bottom=(48, "FFFFFF"))
        _padding(cell, 140, 140, 180, 140)
        _cell_text(cell, label, size=9, bold=True, color=theme.PIZARRA.lstrip("#"), after=4)
        _para(cell, value, size=20, bold=True, color=NIGHT, family=DISPLAY, after=2)
        _para(cell, note, size=8.5, color=MUTED, after=0)


def _configure(doc, context):
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin, section.bottom_margin = Inches(1.15), Inches(.7)
    section.left_margin, section.right_margin = Inches(.7), Inches(.7)
    section.header_distance, section.footer_distance = Inches(.35), Inches(.35)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = TEXT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), TEXT)
    for name in ["Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"]:
        style = styles[name]
        style.font.name = DISPLAY
        style.font.color.rgb = RGBColor.from_string(NIGHT)
        props = style.element.find(qn("w:pPr"))
        if props is not None:
            for border in list(props.findall(qn("w:pBdr"))):
                props.remove(border)
    # The cover plate carries the logo on page one; inner pages repeat it in the header.
    section.different_first_page_header_footer = True
    first_footer = section.first_page_footer.paragraphs[0]
    _font(first_footer.add_run("Ministerio de Trabajo y Desarrollo Humano · InfoReportes"), size=8, color=MUTED)
    header = section.header
    hp = header.paragraphs[0]
    hp.paragraph_format.space_after = Pt(0)
    table = header.add_table(rows=1, cols=2, width=Inches(CONTENT_WIDTH))
    table.autofit = False
    table.columns[0].width, table.columns[1].width = Inches(4.3), Inches(2.8)
    for cell in table.rows[0].cells:
        _shade(cell, PURPLE)
        _borders(cell, bottom=(18, GOLD))
        _padding(cell, 50, 50, 140, 140)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = table.cell(0, 0).paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    if LOGO.exists():
        p.add_run().add_picture(str(LOGO), width=Inches(2.6))
    p = table.cell(0, 1).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    _font(p.add_run("InfoReportes\n"), size=9, bold=True, color="FFFFFF", family=DISPLAY)
    _font(p.add_run(context["period_label"]), size=9, color="E4DBF3")
    footer = section.footer.paragraphs[0]
    footer.paragraph_format.space_after = Pt(0)
    _font(footer.add_run("Ministerio de Trabajo y Desarrollo Humano · InfoReportes"), size=8, color=MUTED)
    footer.add_run("\t")
    footer.paragraph_format.tab_stops.clear_all()
    footer.paragraph_format.tab_stops.add_tab_stop(Inches(CONTENT_WIDTH), WD_TAB_ALIGNMENT.RIGHT)
    _font(footer.add_run("Página "), size=8, color=MUTED)
    run = footer.add_run()
    _font(run, size=8, color=MUTED)
    begin, instr, end = OxmlElement("w:fldChar"), OxmlElement("w:instrText"), OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end.set(qn("w:fldCharType"), "end")
    for item in [begin, instr, end]:
        run._r.append(item)
    doc.core_properties.title = context["title"]
    doc.core_properties.author = context["author"]
    doc.core_properties.subject = context["period_label"] + ". Estadísticas agregadas de protección social."
    doc.core_properties.comments = "Selección de datos: " + context["selection_hash"]
    doc.core_properties.category = "Estadística institucional"


def _page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def generate_docx(context: dict) -> bytes:
    if not context.get("sections"):
        raise ValueError("El reporte no contiene secciones.")
    doc = Document()
    _configure(doc, context)
    _cover(doc, context)
    _heading(doc, "Indicadores principales", before=14)
    _kpis(doc, context)
    _heading(doc, "Lectura del corte")
    for section in context["sections"]:
        for text in section["insights"][1:3]:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.left_indent = Inches(.22)
            p.paragraph_format.first_line_indent = Inches(-.18)
            _font(p.add_run("●  "), size=8, color=GOLD)
            _font(p.add_run(text), size=10.5)
    executed = context["executed_total"]
    _note(doc, ("N/D. El Excel de origen presenta montos presupuestarios acumulados, pero no acredita por sí solo su fase de ejecución. "
                "Este indicador se habilita con una fuente adicional validada.") if executed is None else
          f"USD {money(executed)}. Ejecución verificada para toda la selección.", "Presupuesto ejecutado:")
    _heading(doc, "Criterios de lectura", size=11.5)
    _para(doc, context["notes"][0], size=9, color=MUTED, after=4)
    _para(doc, context["notes"][1] + f" Fuente: {context['source_label']}; corte {context['period_label'].lower()}.", size=9, color=MUTED, after=0)
    for section in context["sections"]:
        _page_break(doc)
        service = section["id"] == "proteccion_integral"
        tone = theme.FUCSIA if service else theme.SERIE_BONOS
        _heading(doc, section["title"], size=20, before=0, after=2)
        _para(doc, f"{number(section['people'])} {section['people_label'].lower()}" +
              (f" · {number(section['units'])} unidades de atención" if section["units"] is not None else ""),
              size=12, bold=True, color=PURPLE, after=6)
        _para(doc, section["insights"][1], after=8)
        headers = ["Servicio" if service else "Programa", "Usuarios" if service else "Beneficiarios"]
        if service:
            headers.append("Unidades")
        headers += [f"Fiscal {context['year']}\n(USD M)", "Gobierno\n(USD M)"]
        widths = [2.45, 1.05, .85, 1.35, 1.4] if service else [3.0, 1.3, 1.4, 1.4]
        rows = []
        for row in section["rows"]:
            values = [row["label"], number(row["people"])] + ([number(row["units"])] if service else [])
            rows.append(values + [money(row["fiscal"], True), money(row["government"], True)])
        total = ["Total de la selección", number(section["people"])] + ([number(section["units"])] if service else [])
        total += [money(section["fiscal"], True), money(section["government"], True)]
        _table(doc, headers, rows, widths, total)
        _para(doc, "Montos acumulados en millones de USD. El acumulado de gobierno ya incluye el monto fiscal; no se suman entre sí.",
              size=8.5, color=MUTED, before=4, after=6)
        if len(section["rows"]) > 1:
            _heading(doc, "Cobertura por " + ("servicio" if service else "programa"), size=12)
            tall = 2.05 if len(section["rows"]) > 5 else 1.7
            doc.add_paragraph().add_run().add_picture(BytesIO(bar_chart(section["rows"], label=section["people_label"], limit=8,
                                                                        height=tall, color=tone)), width=Inches(CONTENT_WIDTH))
        if context.get("include_territory", True) and len(section.get("territories", [])) > 1:
            _heading(doc, "Provincias con mayor cobertura", size=12)
            picture = bar_chart(section["territories"], label=section["people_label"], limit=6, height=1.7, color=tone)
            doc.add_paragraph().add_run().add_picture(BytesIO(picture), width=Inches(CONTENT_WIDTH))
        closing = (section["insights"][2] + " " if len(section["insights"]) > 2 and context.get("include_territory", True) else "")
        _para(doc, closing + "Los conteos agregados no permiten verificar personas únicas. Presupuesto ejecutado: " +
              ("N/D." if section["executed"] is None else "USD " + money(section["executed"]) + "."), size=9, color=MUTED)
    territory = context.get("territory", [])
    if context.get("include_territory", True) and territory:
        _page_break(doc)
        _heading(doc, "Cobertura en el territorio", size=20, before=0, after=2)
        _para(doc, "Beneficiarios y usuarios reportados por provincia en la selección", size=12, bold=True, color=PURPLE, after=6)
        values = {t["code"]: t["people"] for t in territory}
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(BytesIO(map_png(values, label="Cobertura reportada", annotate=6, width=6.6, height=5.6)), width=Inches(5.2))
        top = [t for t in territory if t["code"] != "90"][:8]
        total_people = sum(t["people"] or 0 for t in territory)
        ranking = [[t["label"], number(t["people"]), f"{number((t['people'] or 0) / total_people * 100 if total_people else 0, 1)} %", t.get("region", "")]
                   for t in top]
        _heading(doc, "Provincias con mayor cobertura", size=12)
        _table(doc, ["Provincia", "Cobertura reportada", "Participación", "Región natural"], ranking, [2.9, 1.5, 1.3, 1.4],
               align=["left", "right", "right", "left"])
        _para(doc, "Galápagos se muestra en un recuadro fuera de su posición real. La zona en estudio se incluye en los totales sin polígono "
                   "en el mapa. La cobertura suma registros de distintos programas.", size=8.5, color=MUTED, before=4)
    out = BytesIO()
    doc.save(out)
    return out.getvalue()
