"""Risk exports using the shared institutional PDF, PowerPoint and Excel templates."""
from __future__ import annotations

from src import names, riesgos, theme, ui
from .areas import _base, _ranked, kpis_from_cards
from .context import number

PALETTE = [theme.LILA_300, theme.MORADO, theme.FUCSIA_700]
LEVEL_KEYS = ["media", "alta", "muy_alta"]
SOURCES = ("SIIMTDH · DGID / SGDE / MTDH. Capas de inundaciones y sequía: MAGAP 2015; "
           "movimientos en masa: SNGR 2011; incendios forestales: SNGR 2015.")
CAUTION = ("Una unidad puede estar expuesta a varias amenazas: los totales entre amenazas no se suman. "
           "«Sin categoría reportada» no acredita ausencia de riesgo. Hay unidades sin coordenadas en la fuente.")


def build_context(view, *, title: str) -> dict:
    from src.views.riesgos import METHODOLOGY
    if view.selected.empty:
        raise ValueError("No hay unidades expuestas con los filtros seleccionados.")
    if not view.matrix_available:
        raise ValueError("La matriz fuente de multiamenaza no está disponible para este corte.")
    units = view.detail_selected.copy()
    total = riesgos.count_units(view.selected)
    if len(units) != total or units.codigo_ua.duplicated().any():
        raise ValueError("El detalle de unidades no coincide con el resumen de este corte.")
    units["provincia"] = units.provincia.map(names.place)
    units = units.sort_values(["provincia", "servicio", "nombre", "codigo_ua"])
    services = [{"name": str(name), "units": int(value)} for name, value in
                view.selected.groupby("servicio").unidades.sum().sort_values(ascending=False).items()]
    multi_provinces = view.multihazard.groupby("provincia_codigo").registros.sum().to_dict()
    provinces = [{"code": str(code), "name": names.place(str(name)), "units": int(value),
                  "multihazard": int(multi_provinces.get(code, 0)), "share": value / total}
                 for (code, name), value in view.selected.groupby(["provincia_codigo", "provincia"]).unidades.sum()
                 .sort_values(ascending=False).items()]
    counts = riesgos.threat_counts(view.selected, view.threat, view.levels)
    threats = []
    for name, rows in counts.groupby("amenaza", sort=False):
        values = rows.set_index("nivel").unidades.to_dict()
        row = {"name": name, **{key: int(values.get(level, 0)) for key, level in zip(LEVEL_KEYS, riesgos.LEVELS)}}
        row["units"] = sum(row[key] for key in LEVEL_KEYS)
        threats.append(row)
    ctx = _base(view, "riesgos", title)
    if len(ctx["territory_label"]) > 48:
        ctx["territory_label"] = f"{len(view.provinces)} provincias seleccionadas"
    if len(ctx["programs_label"]) > 48:
        ctx["programs_label"] = f"{len(view.programs)} servicios seleccionados"
    ctx.update(source="Fuente: SIIMTDH · DGID / SGDE / MTDH", kpis=kpis_from_cards(view.kpi_cards()),
               services=services, provinces=provinces, threats=threats, units=units.to_dict("records"),
               scope=view.scope_items(), methodology=METHODOLOGY, sources=SOURCES, caution=CAUTION, multi_note=riesgos.MULTI_NOTE,
               geoportal=ui.GEOPORTAL_URL,
               selection=f"{riesgos.THREATS.get(view.threat, 'Todas las amenazas')} · {', '.join(view.levels)}",
               insights=[{"icon": "building", "text": f"{number(total)} unidades únicas expuestas en la selección."},
                         {"icon": "storm", "text": "La exposición territorial no confirma una afectación efectiva."}])
    return ctx


def render_pdf(c, ctx: dict, pages: int) -> None:
    from xml.sax.saxutils import escape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Paragraph
    from .charts import emblem_png
    from .pdf import H, M, W, INK, C, _band, _bars, _heading, _kpis, _next_page, _slim_band, _table
    from .pdf_areas import _legend, _two_findings
    width = W - 2 * M

    def paragraph(text, top, size=8.5):
        p = Paragraph(escape(text), ParagraphStyle("risk", fontName="Text", fontSize=size,
                                                  leading=size * 1.4, textColor=INK))
        _, height = p.wrap(width, 1000)
        p.drawOn(c, M, top - height)
        return top - height - 10

    y = _band(c, ctx, H, emblem_png())
    y = _kpis(c, ctx, y - 12)
    y = paragraph(ctx["selection"], y - 8)
    _heading(c, "Unidades expuestas por servicio", M, y - 6, width)
    y -= 30
    _bars(c, _ranked(ctx["services"], "units"), M, y, width, row_h=30)
    y -= len(ctx["services"]) * 30 + 14
    _heading(c, "Amenazas y susceptibilidad", M, y, width)
    _legend(c, list(zip(PALETTE, riesgos.LEVELS)), M, y - 22)
    y -= 45
    peak = max((r["units"] for r in ctx["threats"]), default=1) or 1
    for row in ctx["threats"]:
        c.setFillColor(INK)
        c.setFont("Text-Semi", 8.5)
        c.drawString(M, y, row["name"])
        c.drawRightString(W - M, y, number(row["units"]))
        x = M
        for key, color in zip(LEVEL_KEYS, PALETTE):
            length = width * row[key] / peak
            c.setFillColor(C(color))
            c.rect(x, y - 15, length, 8, stroke=0, fill=1)
            x += length
        y -= 34
    y = paragraph("Una unidad puede estar expuesta a varias amenazas; sus totales no se suman.", y, 8)
    _two_findings(c, ctx["insights"], y - 6, rows=1)
    _next_page(c, ctx, 1, pages)
    y = _slim_band(c, ctx, "Unidades expuestas por provincia", ctx["selection"])
    y = paragraph(" · ".join(f"{key}: {value}" for key, value in ctx["scope"] if key in ("Territorio", "Servicios")), y - 12, 8)
    y = _table(c, M, y - 16, width,
               [("Provincia", "name", None, "text"), ("UA expuestas", "units", 90, "int"),
                ("Multiamenaza*", "multihazard", 100, "int")], ctx["provinces"], row_h=14.2)
    _heading(c, "Nota metodológica", M, y - 23, width)
    y = paragraph(ctx["methodology"], y - 42)
    y = paragraph("* " + ctx["multi_note"], y, 8)
    y = paragraph(ctx["caution"], y, 8)
    y = paragraph(ctx["sources"], y, 7.5)
    c.setFillColor(C(theme.MORADO))
    c.setFont("Text-Semi", 8.5)
    c.drawString(M, y - 3, "Consultar el geoportal del Ministerio")
    c.linkURL(ctx["geoportal"], (M, y - 6, M + 200, y + 8), relative=0)


def render_pptx(prs, ctx: dict) -> None:
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
    from pptx.util import Inches, Pt
    from .charts import emblem_png
    from .pptx import TEXT, _bar_chart, _cover, _footer, _header, _rgb, _text
    from .pptx_areas import _table, _title
    _cover(prs, ctx, emblem_png(), 4, share_word="")

    def slide(title):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        _header(s, title, ctx["selection"])
        _footer(s, ctx, len(prs.slides), 4)
        s.notes_slide.notes_text_frame.text = "\n".join(f"{k}: {v}" for k, v in ctx["scope"]) + "\n" + ctx["sources"]
        return s

    s = slide("Exposición por servicio y amenaza")
    _title(s, "Unidades por servicio", .55, 1.36, 5.9)
    _title(s, "Amenazas y susceptibilidad", 6.95, 1.36, 5.8)
    _bar_chart(s, [r["name"] for r in reversed(ctx["services"])], [r["units"] for r in reversed(ctx["services"])],
               "Unidades expuestas", x=.45, y=1.95, w=6.1, h=4.45)
    data = CategoryChartData()
    data.categories = [r["name"] for r in reversed(ctx["threats"])]
    for key, level in zip(LEVEL_KEYS, riesgos.LEVELS):
        data.add_series(level, [r[key] for r in reversed(ctx["threats"])])
    chart = s.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED, Inches(6.7), Inches(1.95), Inches(6.1), Inches(4.45), data).chart
    chart.font.name = TEXT
    chart.font.size = Pt(11)
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.value_axis.minimum_scale = 0
    chart.value_axis.has_major_gridlines = False
    chart.category_axis.format.line.fill.background()
    chart.plots[0].gap_width = 65
    for series, color in zip(chart.series, PALETTE):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = _rgb(color)
    chart.plots[0].has_data_labels = True
    chart.plots[0].data_labels.position = XL_LABEL_POSITION.CENTER
    chart.plots[0].data_labels.number_format = '#,##0;-#,##0;""'
    chart.plots[0].data_labels.font.size = Pt(10)
    chart.plots[0].data_labels.font.color.rgb = _rgb("#FFFFFF")
    _text(s, .55, 6.58, 12.2, .32, "Unidades únicas por servicio. Una unidad puede aparecer en varias amenazas; sus totales no se suman.",
          size=11, color=theme.PIZARRA)
    s = slide("Unidades expuestas por provincia")
    columns = [("Provincia", "name", None, "text"), ("UA expuestas", "units", 1.15, "int"),
               ("Multiamenaza*", "multihazard", 1.3, "int")]
    split = (len(ctx["provinces"]) + 1) // 2
    for x, rows in [(0.55, ctx["provinces"][:split]), (6.95, ctx["provinces"][split:])]:
        if rows:
            _table(s, x, 1.5, 5.85, columns, rows, row_h=.36, size=11)
    _text(s, .55, 6.4, 12.2, .5, "* " + ctx["multi_note"], size=11, color=theme.PIZARRA)
    s = slide("Nota metodológica")
    _text(s, .65, 1.52, 12, 2.55, ctx["methodology"], size=18, spacing=1.15)
    _text(s, .65, 4.25, 12, .85, ctx["caution"], size=14, color=theme.PIZARRA)
    _text(s, .65, 5.32, 12, .78, ctx["sources"], size=13, color=theme.PIZARRA)
    link = _text(s, .65, 6.35, 12, .35, "Consultar el geoportal del Ministerio", size=15, color=theme.MORADO, bold=True)
    link.text_frame.paragraphs[0].runs[0].hyperlink.address = ctx["geoportal"]
    # Full scope travels with every slide, including the shared institutional cover.
    prs.slides[0].notes_slide.notes_text_frame.text += "\n" + "\n".join(f"{k}: {v}" for k, v in ctx["scope"])


def render_workbook(book) -> None:
    from .excel import TEXT, DISPLAY
    from .excel_areas import _block, _finish_summary, _sheet, _summary
    ctx = book.ctx
    ws, row = _summary(book, "")
    row = book.heading(ws, row, "Unidades expuestas por servicio", ctx["selection"], last=12)
    first, end, row = book.table(ws, row, [("Servicio", "name", "text", 1, 8), ("Unidades expuestas", "units", "int", 9, 4)],
                                ctx["services"], total={"units": len(ctx["units"])}, height=24)
    chart = book.bar_chart("Resumen", first, end, 1, 9, "Unidades expuestas por servicio", [theme.MORADO], width=850, height=310)
    ws.insert_chart(row + 1, 1, chart)
    _finish_summary(book, ws, row + 19, ctx["insights"])

    ws, row = _sheet(book, "Amenazas", theme.FUCSIA_700, "Amenazas y susceptibilidad", [2.4, 34, 15, 15, 15, 19, 2.4])
    first, end, row = _block(book, ws, row + 1, "Unidades por amenaza", [("Amenaza", "name", "text"),
                  ("Media", "media", "int"), ("Alta", "alta", "int"), ("Muy alta", "muy_alta", "int"),
                  ("UA expuestas", "units", "int")], ctx["threats"], note=ctx["selection"], last=5, height=25)
    chart = book.wb.add_chart({"type": "bar", "subtype": "stacked"})
    for col, level, color in zip(range(2, 5), riesgos.LEVELS, PALETTE):
        chart.add_series({"name": level, "categories": ["Amenazas", first, 1, end, 1],
                          "values": ["Amenazas", first, col, end, col], "fill": {"color": color}, "border": {"none": True}})
    chart.set_title({"name": "Unidades por amenaza y susceptibilidad", "name_font": {"name": DISPLAY, "size": 12, "color": theme.MORADO_900}})
    chart.set_legend({"position": "bottom", "font": {"name": TEXT, "size": 10}})
    chart.set_y_axis({"reverse": True, "num_font": {"name": TEXT, "size": 10}, "line": {"none": True}})
    chart.set_x_axis({"min": 0, "num_format": "#,##0", "num_font": {"name": TEXT, "size": 9}, "major_gridlines": {"visible": False}})
    chart.set_chartarea({"border": {"none": True}})
    chart.set_size({"width": 760, "height": 320})
    ws.insert_chart(row + 1, 1, chart)
    book.findings(ws, row + 20, [{"icon": "storm", "text": "Una unidad puede aparecer en varias amenazas: sus totales no se suman."}], last=5)
    ws.freeze_panes(first, 2)

    ws, row = _sheet(book, "Provincias", theme.VIOLETA_500, "Unidades expuestas por provincia", [2.4, 34, 20, 22, 20, 2.4])
    first, end, row = _block(book, ws, row + 1, "Distribución territorial", [("Provincia", "name", "text"),
        ("UA expuestas", "units", "int"), ("Multiamenaza*", "multihazard", "int"), ("Participación", "share", "pct")],
        ctx["provinces"], last=4, note=ctx["selection"], total={"units": len(ctx["units"]),
            "multihazard": sum(r["multihazard"] for r in ctx["provinces"]), "share": 1}, autofilter=True, height=23)
    ws.merge_range(row + 1, 1, row + 1, 4, "* " + ctx["multi_note"], book.f(font_size=10, text_wrap=True, font_color=theme.PIZARRA))
    ws.set_row(row + 1, 55)
    ws.freeze_panes(first, 2)

    ws, row = _sheet(book, "Unidades", theme.MORADO, "Unidades de atención expuestas",
                    [2.4, 17, 58, 36, 28, 25, 25, 25, 25, 2.4])
    columns = [("Código UA", "codigo_ua", "code"), ("Unidad de atención", "nombre", "text"),
               ("Servicio", "servicio", "text"), ("Provincia", "provincia", "text")]
    columns += [(label, key, "text") for key, label in riesgos.THREATS.items()]
    first, end, row = _block(book, ws, row + 1, "Detalle de la selección", columns, ctx["units"], last=8,
                             note=ctx["selection"], height=36, autofilter=True)
    for offset, unit in enumerate(ctx["units"]):
        ws.set_row(first + offset, max(36, 14 * ((len(unit["nombre"]) + 46) // 47)))
    ws.freeze_panes(first, 3)
    ws.repeat_rows(first - 1)
    ws.set_zoom(75)

    facts = [("Institución", ctx["institution"]), *ctx["scope"], *book.update_facts(),
             ("Generado", ctx["generated_label"]), ("Unidad de medida", "UA únicas; multiamenaza conserva el conteo de registros de la matriz fuente.")]
    book.facts_sheet(facts,
        [("UA expuestas", "Unidades únicas con susceptibilidad media, alta o muy alta en las amenazas seleccionadas."),
         ("Multiamenaza", ctx["multi_note"]),
         ("Detalle", "Una fila por código UA; las cuatro amenazas conservan la categoría reportada en la fuente.")],
        [("Resumen", "Cifras clave y servicios"), ("Amenazas", "Categorías por amenaza"), ("Provincias", "Distribución territorial"),
         ("Unidades", "Código, nombre, servicio, provincia y susceptibilidad")],
        sources=[("Fuentes", ctx["sources"], ""), ("Geoportal", "Visor geográfico del Ministerio", ctx["geoportal"])])
    ws = book.wb.get_worksheet_by_name("Ficha técnica")
    # The shared fact template uses compact rows; expand its long scope/source entries.
    strings = {index: text for text, index in book.wb.str_table.string_table.items()}
    for r, cells in ws.table.items():
        cell = cells.get(2)
        if cell is not None and hasattr(cell, "string"):
            text = strings.get(cell.string, "")
            if len(text) > 95:
                ws.set_row(r, max(38, 15 * ((len(text) + 84) // 85)))
    # Long methodological text gets its own wrapped block, beyond the compact definition rows.
    row = book.heading(ws, ws.dim_rowmax + 2, "Nota metodológica", last=2)
    ws.merge_range(row, 1, row, 2, ctx["methodology"], book.f(font_size=11, text_wrap=True, valign="top", font_color=theme.TINTA))
    ws.set_row(row, 130)
    ws.merge_range(row + 2, 1, row + 2, 2, ctx["caution"], book.f(font_size=10, text_wrap=True, font_color=theme.PIZARRA))
    ws.set_row(row + 2, 50)
