"""Presentations of the areas with their own structure, built with the native shapes, charts and tables of python-pptx.

Contingencies and compensation: 1. cover; 2. results by program (table) and findings; 3. each program month by month;
4. territory (map and leading places); 5. budget (accrued by month and execution) or, without budget, the types.
Jóvenes en Acción: 1. cover; 2. the phases; 3. registration and payment of the current phase; 4. payments by month.
Labour management: 1. cover; 2. territory; 3. profile and economic activity; 4. contracts registered in the year (or the
types of contract); 5. retirement compensation.
Social budget: 1. cover; 2. execution by program; 3. the year month by month; 4. detail by program and the pace the year
needs.
Every figure carries its unit, and every chart and table stays editable.
"""
from __future__ import annotations

from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from src import theme
from .charts import emblem_png, map_png, nice_bounds
from .context import millions, number
from .pptx import (DISPLAY, TEXT, _badge, _bar_chart, _box, _cover, _finding, _findings_column, _footer, _header, _line_chart,
                   _picture, _ranked, _rgb, _text)

LEVEL_LABELS = {"provincia": "Provincia", "canton": "Cantón", "parroquia": "Parroquia"}
LEVEL_PLURALS = {"provincia": "Provincias", "canton": "Cantones", "parroquia": "Parroquias"}
PHASE_COLORS = {1: theme.LILA_300, 2: theme.VIOLETA_500, 3: theme.ORO}
SHORT_UNITS = {"beneficiarios": "Beneficiarios", "personas habilitadas": "Habilitadas", "personas con pago acreditado": "Pago acreditado"}


# ----------------------------------------------------------------------------------------------- helpers
def _usd(value) -> str:
    if value is None or value != value:
        return "—"
    return f"USD {number(value / 1e6, 1)} M" if abs(value) >= 1_000_000 else f"USD {number(value)}"


def _cell(value, kind: str) -> str:
    if value is None or value != value:
        return "—"
    if kind == "int":
        return number(value)
    if kind in ("pct", "rate"):
        return f"{number(value, 1)} %"
    if kind == "spct":
        sign = "+" if value > 0.05 else ("−" if value < -0.05 else "")
        return f"{sign}{number(abs(value), 1)} %"
    return str(value)


def _table(slide, x: float, y: float, w: float, columns: list[tuple], rows: list[dict], *, row_h: float = 0.36, header_h: float = 0.5,
           size: float = 11):
    """Native table: ``columns`` are (header, key, width or None, kind); the header in purple, rows banded, ``selected`` in gold."""
    fixed = sum(c[2] for c in columns if c[2])
    widths = [c[2] if c[2] else w - fixed for c in columns]
    shape = slide.shapes.add_table(len(rows) + 1, len(columns), Inches(x), Inches(y), Inches(w), Inches(header_h + row_h * len(rows)))
    table = shape.table
    table.first_row = True
    table.horz_banding = False
    for index, width in enumerate(widths):
        table.columns[index].width = Inches(width)
    table.rows[0].height = Inches(header_h)
    for r in range(1, len(rows) + 1):
        table.rows[r].height = Inches(row_h)

    def fill(cell, text, *, color, bold=False, fg=theme.TINTA, align=PP_ALIGN.LEFT, font_size=size):
        cell.fill.solid()
        cell.fill.fore_color.rgb = _rgb(color)
        cell.margin_left = cell.margin_right = Inches(0.08)
        cell.margin_top = cell.margin_bottom = Inches(0.03)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        frame = cell.text_frame
        frame.word_wrap = True
        paragraph = frame.paragraphs[0]
        paragraph.alignment = align
        run = paragraph.add_run()
        run.text = text
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.name = TEXT
        run.font.color.rgb = _rgb(fg)

    for index, (header, _key, _w, kind) in enumerate(columns):
        fill(table.cell(0, index), header, color=theme.MORADO, bold=True, fg="#FFFFFF",
             align=PP_ALIGN.LEFT if kind == "text" else PP_ALIGN.RIGHT, font_size=size - 1)
    for r, row in enumerate(rows, start=1):
        band = theme.ORO_100 if row.get("selected") else (theme.NIEBLA if r % 2 == 0 else "#FFFFFF")
        for index, (_header, key, _w, kind) in enumerate(columns):
            value = row.get(key)
            fill(table.cell(r, index), str(value) if kind == "text" else _cell(value, kind), color=band,
                 bold=bool(row.get("selected")) or index == 0 and kind == "text",
                 fg=theme.MORADO_900 if index == 0 else theme.TINTA, align=PP_ALIGN.LEFT if kind == "text" else PP_ALIGN.RIGHT)
    return shape


def _columns_chart(slide, categories: list[str], values: list[float], name: str, *, x, y, w, h, money: bool = False,
                   highlight: int | None = None, colors: list[str] | None = None, number_format: str | None = None):
    """Native columns with their figures (millions of USD for money), the month of the report or each phase in its colour."""
    scaled = [(v or 0) / 1e6 if money else (v or 0) for v in values]
    # Small amounts keep two decimals, so no column of a chart reads «0,0 M».
    small = money and any(0 < abs(v or 0) < 50_000 for v in values)
    fmt = number_format or (('#,##0.00" M"' if small else '#,##0.0" M"') if money else "#,##0")
    chart = _bar_chart(slide, categories, scaled, name, x=x, y=y, w=w, h=h, number_format=fmt, chart_type=XL_CHART_TYPE.COLUMN_CLUSTERED)
    chart.plots[0].gap_width = 60
    chart.font.size = Pt(11)
    chart.plots[0].data_labels.font.size = Pt(11)
    series = chart.plots[0].series[0]
    for index in range(len(values)):
        color = (colors[index] if colors else (theme.ORO if index == highlight else theme.SERIE_BONOS))
        point = series.points[index]
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = _rgb(color)
    return chart


def _skip_labels(chart, every: int) -> None:
    """Show one category label out of ``every`` (the schema places tickLblSkip before noMultiLvlLbl)."""
    from lxml import etree
    from pptx.oxml.ns import qn
    axis = chart.category_axis._element
    skip = axis.find(qn("c:tickLblSkip"))
    if skip is None:
        skip = etree.Element(qn("c:tickLblSkip"))
        anchor = axis.find(qn("c:noMultiLvlLbl"))
        if anchor is not None:
            anchor.addprevious(skip)
        else:
            axis.append(skip)
    skip.set("val", str(every))


def _title(slide, text: str, x: float, y: float, w: float, note: str = ""):
    _text(slide, x, y, w, 0.4, text, size=18, color=theme.MORADO_900, font=DISPLAY, bold=True)
    if note:
        _text(slide, x, y + 0.42, w, 0.3, note, size=11.5, color=theme.PIZARRA)


def _stacked(slide, title: str, parts: list[dict], x: float, y: float, w: float, palette: list[str]) -> float:
    """A 100 % bar in native shapes, its parts named under it."""
    _text(slide, x, y, w, 0.3, title, size=12.5, color=theme.PIZARRA, bold=True)
    total = sum(p["share"] or 0 for p in parts) or 100
    cx = x
    for index, part in enumerate(parts):
        length = w * (part["share"] or 0) / total
        if length > 0.01:
            _box(slide, cx, y + 0.36, max(0.01, length - 0.02), 0.26, palette[index % len(palette)])
        cx += length
    labels = [f"{p['name']}: {number(p['share'], 1)} %" for p in parts]
    _text(slide, x, y + 0.68, w, 0.3, "   ·   ".join(labels), size=11, color=theme.TINTA)
    return y + 1.08


def _closing(slide, ctx: dict, page: int, total: int, findings: list[dict]) -> None:
    _footer(slide, ctx, page, total)
    slide.notes_slide.notes_text_frame.text = " ".join(f["text"] for f in findings)


# ------------------------------------------------------------------------------------------ contingencies
def contingency(prs, ctx: dict) -> None:
    levels = ctx["levels"]
    has_map = bool(levels and ctx["territories"].get(levels[0]))
    budget = ctx.get("budget")
    total = 3 + (1 if has_map else 0) + (1 if budget or ctx["types"] else 0)
    _cover(prs, ctx, emblem_png(width=4.2, height=3.7, dpi=220), total, share_word="")
    page = 1

    # Results by program: the criterion of each source, the month and the year.
    page += 1
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _header(slide, "Resultados por programa", f"{ctx['when'].capitalize()} · cada programa con el criterio de su fuente")
    columns = [("Programa", "name", None, "text"), ("Criterio", "criterion", 1.75, "text"), ("Personas", "people", 1.05, "int"),
               ("Monto del mes (USD)", "amount", 1.45, "int")]
    if budget:
        columns += [("Devengado acumulado (USD)", "accrued", 1.75, "int"), ("Ejecución", "execution", 1.05, "rate")]
    rows = [{**p, "criterion": SHORT_UNITS.get(p["unit"], "")} for p in ctx["programs"]]
    width = 12.25
    row_h = min(0.55, 3.3 / max(1, len(rows)))
    _table(slide, 0.55, 1.5, width, columns, rows, row_h=row_h, size=12.5)
    findings = ctx["insights"][:2]
    top = 1.5 + 0.5 + row_h * len(rows) + 0.35
    if len(rows) <= 2:
        # Few programs: the year month by month next to the findings.
        program = ctx["programs"][0]
        months = ctx["months"]
        year_columns = [("Mes", "label", None, "text"), (SHORT_UNITS.get(program["unit"], "Personas"), "people", 1.35, "int"),
                        ("Monto (USD)", "amount", 1.45, "int")]
        if budget:
            year_columns += [("Devengado (USD)", "month_accrued", 1.45, "int"), ("Ejecución", "execution", 1.05, "rate")]
        year_rows = [{"label": m["label"], "people": m["people"].get(program["id"]), "amount": m["amount"], "month_accrued": m["month_accrued"],
                      "execution": m["execution"], "selected": m["selected"]} for m in months]
        year_h = min(0.34, (6.85 - top - 0.45) / max(1, len(year_rows)))
        _text(slide, 0.55, top - 0.05, 7.3, 0.35, f"{program['name']} mes a mes", size=15, color=theme.MORADO_900, font=DISPLAY, bold=True)
        _table(slide, 0.55, top + 0.35, 7.3, year_columns, year_rows, row_h=year_h, header_h=0.42, size=10.5)
        _findings_column(slide, findings, x=8.15, y=top + 0.35, w=4.65, h=1.45, step=1.6)
    else:
        half = (width - 0.25) / 2
        for index, finding in enumerate(findings):
            _finding(slide, finding, 0.55 + index * (half + 0.25), top, half, min(1.5, 6.9 - top))
    _closing(slide, ctx, page, total, findings)

    # Each program month by month: native line charts.
    page += 1
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _header(slide, "Evolución mensual", "Personas de cada programa en los últimos doce meses · el mes de la consulta en dorado")
    history = ctx["history"][:6]
    per_row = 3 if len(history) > 2 else max(1, len(history))
    rows_count = (len(history) + per_row - 1) // per_row or 1
    cell_w = (12.25 - 0.3 * (per_row - 1)) / per_row
    cell_h = (5.45 - 0.2 * (rows_count - 1)) / rows_count
    for index, item in enumerate(history):
        x = 0.55 + (index % per_row) * (cell_w + 0.3)
        y = 1.45 + (index // per_row) * (cell_h + 0.2)
        _text(slide, x, y, cell_w, 0.32, item["name"], size=14, color=theme.MORADO_900, font=DISPLAY, bold=True)
        _text(slide, x, y + 0.32, cell_w, 0.28, f"{item['short_criterion']} · {item['latest_label']}: {number(item['latest'])}", size=11,
              color=theme.PIZARRA)
        present = [v for v in item["values"] if v is not None] or [0, 1]
        low, high, step = nice_bounds(min(present), max(present), steps=3)
        chart = _line_chart(slide, [label for label in item["labels"]], [(item["short_criterion"], item["values"], theme.SERIE_BONOS)],
                            x=x - 0.05, y=y + 0.6, w=cell_w + 0.1, h=cell_h - 0.65, highlight=item.get("highlight"), minimum=low,
                            maximum=high, step=step)
        if per_row > 1:
            _skip_labels(chart, 3)
    _closing(slide, ctx, page, total, [])

    if has_map:
        page += 1
        level = levels[0]
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _header(slide, f"Bonos habilitados por {LEVEL_LABELS[level].lower()}", f"{ctx['map_program_label']} · {ctx['span']}")
        picture = map_png(ctx["level_maps"][level], layer=level, mode="light", label=f"Bonos habilitados de {ctx['span']}", width=6.0,
                          height=5.4, dpi=220)
        _picture(slide, picture, 0.45, 1.42, h=5.55)
        rows = ctx["territories"][level][:8]
        peak = max([r["records"] or 0 for r in rows] or [1]) or 1
        ranked = [{"name": r["name"], "text": f"{number(r['records'])} · {number(r['share'], 1)} %", "fraction": (r["records"] or 0) / peak}
                  for r in rows]
        _title(slide, f"{LEVEL_PLURALS[level]} con más bonos habilitados", 7.05, 1.45, 5.7)
        _ranked(slide, ranked, 7.05, 2.0, 5.7, step=0.47)
        finding = next((i for i in ctx["insights"] if i["icon"] in ("map", "storm")), ctx["insights"][0])
        _finding(slide, finding, 7.05, 5.9, 5.7, 1.0)
        _closing(slide, ctx, page, total, [finding])

    if budget:
        page += 1
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _header(slide, "Presupuesto", f"Devengado de {ctx['span']} sobre el presupuesto aprobado · {ctx['programs_label']}")
        months = [m for m in ctx["months"] if m.get("month_accrued") is not None]
        selected = next((i for i, m in enumerate(months) if m["selected"]), None)
        _title(slide, "Presupuesto devengado por mes", 0.55, 1.45, 6.5, "Millones de USD · el mes de la consulta en dorado")
        _columns_chart(slide, [m["short"] for m in months], [m["month_accrued"] for m in months], "Devengado (millones de USD)", x=0.45,
                       y=2.2, w=6.6, h=4.65, money=True, highlight=selected)
        funded = sorted([p for p in ctx["programs"] if p.get("approved")], key=lambda p: -(p["approved"] or 0))
        if len(funded) >= 3:
            _title(slide, "Ejecución por programa", 7.4, 1.45, 5.4, f"Total: {number(budget['execution'], 1)} % de {_usd(budget['approved'])}")
            rows = [{"name": p["name"], "text": f"{number(p['execution'], 1)} %" if p["execution"] is not None else "—",
                     "fraction": (p["execution"] or 0) / 100} for p in funded]
            _ranked(slide, rows, 7.4, 2.2, 5.4, step=min(0.62, 3.4 / max(1, len(rows))))
        else:
            _title(slide, "Ejecución acumulada", 7.4, 1.45, 5.4, f"{number(budget['execution'], 1)} % de {_usd(budget['approved'])} aprobados")
            rates = [(m["execution"] or 0) / 100 if m["execution"] is not None else None for m in months]
            _line_chart(slide, [m["short"] for m in months], [("Ejecución acumulada", rates, theme.SERIE_BONOS),
                                                            ("Ritmo uniforme del año", [m["month"] / 12 for m in months], theme.LILA_300)],
                        x=7.3, y=2.2, w=5.6, h=3.5, number_format="0%", highlight=selected, minimum=0, maximum=1.0, legend=True,
                        label_format="0.0%", step=0.25)
        finding = next((i for i in ctx["insights"] if i["icon"] == "target"), ctx["insights"][-1])
        _finding(slide, finding, 7.4, 5.85, 5.4, 1.05)
        _closing(slide, ctx, page, total, [finding])
    elif ctx["types"]:
        page += 1
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _header(slide, "Tipos y montos", f"Bonos habilitados por tipo y monto de cada mes · {ctx['span']}")
        months = [m for m in ctx["months"] if m.get("amount")]
        selected = next((i for i, m in enumerate(months) if m["selected"]), None)
        _title(slide, "Monto habilitado o acreditado por mes", 0.55, 1.45, 6.5, "Millones de USD · el mes de la consulta en dorado")
        _columns_chart(slide, [m["short"] for m in months], [m["amount"] for m in months], "Monto (millones de USD)", x=0.45, y=2.2,
                       w=6.6, h=4.65, money=True, highlight=selected)
        types = ctx["types"][:7]
        peak = max([t["records"] or 0 for t in types] or [1]) or 1
        _title(slide, "Bonos habilitados por tipo", 7.4, 1.45, 5.4, "Bonos habilitados · participación")
        _ranked(slide, [{"name": t["name"], "text": f"{number(t['records'])} · {number(t['share'], 1)} %",
                         "fraction": (t["records"] or 0) / peak} for t in types], 7.4, 2.2, 5.4, step=0.6)
        _closing(slide, ctx, page, total, [])


# --------------------------------------------------------------------------------------- Jóvenes en Acción
def youth(prs, ctx: dict) -> None:
    info = ctx["info"]
    total = 4
    _cover(prs, ctx, emblem_png(width=4.2, height=3.7, dpi=220), total)
    phases = ctx["phases"]
    phase = info["phase_name"].lower()

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _header(slide, "Fases del programa", "Jóvenes participantes, presupuesto, ejecución y pagos acreditados de cada fase")
    _title(slide, "Jóvenes participantes por fase", 0.55, 1.45, 5.2, "La fase en curso en dorado")
    _columns_chart(slide, [f"{p['name']} ({p['period_short']})" for p in phases], [p["participants"] for p in phases], "Jóvenes",
                   x=0.45, y=2.2, w=5.4, h=4.65, highlight=len(phases) - 1)
    rows = [{"name": p["name"], "budget": p["budget"], "executed": p["execution_amount"], "execution": p["execution"],
             "accredited": p["accredited"], "selected": p["current"]} for p in phases]
    _title(slide, "Presupuesto y ejecución", 6.2, 1.45, 6.6, "Montos en USD")
    _table(slide, 6.2, 2.2, 6.6, [("Fase", "name", 0.85, "text"), ("Presupuesto", "budget", 1.45, "int"),
                                  ("Ejecutado", "executed", 1.45, "int"), ("Ejecución", "execution", 1.1, "rate"),
                                  ("Pagos acreditados", "accredited", None, "int")], rows, row_h=0.45)
    finding = ctx["insights"][1]
    _finding(slide, finding, 6.2, 5.0, 6.6, 1.2)
    _closing(slide, ctx, 2, total, [finding])

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _header(slide, f"Registro y pagos · {phase}", f"{ctx['phases'][-1]['period'].capitalize()} · corte {ctx['period_label'].lower()}")
    _title(slide, "Del registro al pago", 0.55, 1.45, 6.0)
    steps = ctx["funnel"]
    first = steps[0][1] or 1
    tones = [theme.LILA_300, theme.SERIE_BONOS, theme.ORO]
    for index, (label, value) in enumerate(steps):
        y = 2.1 + index * 1.25
        _text(slide, 0.55, y, 3.6, 0.35, label, size=15, color=theme.TINTA, bold=True)
        _text(slide, 3.6, y - 0.05, 2.95, 0.42, f"{number(value)} jóvenes", size=20, color=theme.MORADO_900, font=DISPLAY, bold=True,
              align=PP_ALIGN.RIGHT)
        _box(slide, 0.55, y + 0.45, 6.0, 0.2, "#EFECF5", shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
        _box(slide, 0.55, y + 0.45, max(0.1, 6.0 * (value or 0) / first), 0.2, tones[index % 3], shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
        if index:
            previous = steps[index - 1][1]
            _text(slide, 0.55, y + 0.72, 6.0, 0.3, f"{number((value or 0) / previous * 100 if previous else 0, 1)} % de los "
                                                   f"{steps[index - 1][0].lower()}", size=12, color=theme.MORADO, bold=True)
    months = ctx["budget_months"]
    selected = next((i for i, m in enumerate(months) if m["selected"]), None)
    _title(slide, f"Presupuesto devengado de la {phase}", 7.1, 1.45, 5.7, "Millones de USD por mes · el mes de la consulta en dorado")
    shown = [m for m in months if m["month_accrued"]] or months
    _columns_chart(slide, [m["label"] for m in shown], [m["month_accrued"] for m in shown], "Devengado (millones de USD)", x=7.0, y=2.2,
                   w=5.85, h=3.4, money=True, highlight=next((i for i, m in enumerate(shown) if m["selected"]), selected))
    finding = next((i for i in ctx["insights"] if i["icon"] == "banknote"), ctx["insights"][0])
    _finding(slide, finding, 7.1, 5.75, 5.7, 1.1)
    _closing(slide, ctx, 3, total, [finding])

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _header(slide, "Pagos acreditados por mes", "Millones de USD acreditados cada mes, por fase del programa")
    payments = ctx["payments"]
    _columns_chart(slide, [p["short"] for p in payments], [p["amount"] for p in payments], "Monto acreditado (millones de USD)", x=0.45,
                   y=1.5, w=7.6, h=5.0, money=True, colors=[PHASE_COLORS.get(p["phase"], theme.SERIE_BONOS) for p in payments],
                   )
    x = 0.65
    for phase_id in sorted({p["phase"] for p in payments}):
        _box(slide, x, 6.62, 0.18, 0.18, PHASE_COLORS[phase_id], shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.2)
        _text(slide, x + 0.25, 6.57, 1.2, 0.3, f"Fase {phase_id}", size=11.5, color=theme.TINTA)
        x += 1.3
    _title(slide, "Aspectos destacados", 8.35, 1.45, 4.45)
    findings = [i for i in ctx["insights"] if i["icon"] in ("send", "layers", "graduation")][:3]
    _findings_column(slide, findings, x=8.35, y=2.0, w=4.45, h=1.4, step=1.55)
    _closing(slide, ctx, 4, total, findings)


# --------------------------------------------------------------------------------------- labour management
def labour(prs, ctx: dict) -> None:
    selected = set(ctx["selected"])
    total = 5
    _cover(prs, ctx, emblem_png(selected, width=4.2, height=3.7, dpi=220), total)
    level = ctx["map_level"]

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _header(slide, f"Contratos vigentes por {LEVEL_LABELS[level].lower()}", f"{ctx['territory_label']} · {ctx['programs_label']}")
    focus = selected if level == "canton" else None
    picture = map_png(ctx["map_values"], layer=level, focus=focus, selected=selected if level == "provincia" else None, mode="light",
                      label=f"Contratos vigentes por {LEVEL_LABELS[level].lower()}", width=6.0, height=5.4, dpi=220)
    _picture(slide, picture, 0.45, 1.42, h=5.55)
    rank_level = "canton" if level == "canton" or len(selected) == 1 else "provincia"
    rows = ctx["territories"][rank_level][:8]
    peak = max([r["contracts"] or 0 for r in rows] or [1]) or 1
    _title(slide, f"{LEVEL_PLURALS[rank_level]} con más contratos vigentes", 7.05, 1.45, 5.7)
    _ranked(slide, [{"name": r["name"], "text": f"{number(r['contracts'])} · {number(r['share'], 1)} %", "fraction": (r["contracts"] or 0) / peak}
                    for r in rows], 7.05, 2.0, 5.7, step=0.47)
    finding = ctx["insights"][1] if len(ctx["insights"]) > 1 else ctx["insights"][0]
    _finding(slide, finding, 7.05, 5.9, 5.7, 1.0)
    _closing(slide, ctx, 2, total, [finding])

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _header(slide, "Perfil y actividad económica", f"Contratos vigentes · {ctx['territory_label']}")
    palette = [theme.SERIE_BONOS, theme.LILA_300, theme.ORO, theme.FUCSIA]
    y = 1.5
    for title, parts in ctx["profile"].items():
        y = _stacked(slide, title, parts, 0.55, y, 6.0, palette)
    profile_finding = next((i for i in ctx["insights"] if i["icon"] == "users"), None)
    if profile_finding:
        _finding(slide, profile_finding, 0.55, max(y + 0.2, 5.0), 6.0, 1.3)
    branches = ctx["branches"][:8]
    peak = max([b["share"] or 0 for b in branches] or [1]) or 1
    _title(slide, "Por rama de actividad", 7.05, 1.45, 5.7, "Contratos vigentes · participación")
    _ranked(slide, [{"name": b["name"], "text": f"{number(b['contracts'])} · {number(b['share'], 1)} %", "fraction": (b["share"] or 0) / peak}
                    for b in branches], 7.05, 2.25, 5.7, step=0.58)
    _closing(slide, ctx, 3, total, [profile_finding] if profile_finding else [])

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    months = ctx["months"]
    if months:
        registrations = ctx["registrations"]
        _header(slide, "Contratación del año", f"Contratos registrados de {registrations['span']} de {registrations['year']} que siguen "
                                               f"vigentes · total nacional")
        _columns_chart(slide, [m["short"] for m in months], [m["value"] for m in months], "Contratos registrados", x=0.45, y=1.5, w=8.1,
                       h=5.4)
        _title(slide, "Contratos del año", 8.95, 1.45, 3.9)
        _text(slide, 8.95, 1.95, 3.9, 0.6, number(registrations["total"]), size=34, color=theme.MORADO_900, font=DISPLAY, bold=True)
        _text(slide, 8.95, 2.62, 3.9, 0.3, f"Promedio: {number(registrations['average'])} por mes", size=12.5, color=theme.PIZARRA)
        finding = next((i for i in ctx["insights"] if i["icon"] == "file-check"), ctx["insights"][0])
        _finding(slide, finding, 8.95, 3.3, 3.9, 1.6)
        kinds = ctx["kinds"][:3]
        lines = [f"{k['name']}: {number(k['share'], 1)} %" for k in kinds]
        _text(slide, 8.95, 5.1, 3.9, 0.3, "Principales tipos de contrato vigentes", size=12.5, color=theme.PIZARRA, bold=True)
        _text(slide, 8.95, 5.45, 3.9, 1.3, lines, size=12, color=theme.TINTA)
    else:
        _header(slide, "Tipos de contrato y autoidentificación", f"Contratos vigentes · {ctx['territory_label']}")
        kinds = ctx["kinds"][:8]
        peak = max([k["share"] or 0 for k in kinds] or [1]) or 1
        _title(slide, "Por tipo de contrato", 0.55, 1.45, 6.0, "Contratos vigentes · participación")
        _ranked(slide, [{"name": k["name"], "text": f"{number(k['contracts'])} · {number(k['share'], 1)} %", "fraction": (k["share"] or 0) / peak}
                        for k in kinds], 0.55, 2.25, 6.0, step=0.58)
        groups = ctx["ethnicity"][:8]
        peak = max([g["share"] or 0 for g in groups] or [1]) or 1
        _title(slide, "Por autoidentificación étnica", 7.05, 1.45, 5.7, "Contratos vigentes · participación")
        _ranked(slide, [{"name": g["name"], "text": f"{number(g['contracts'])} · {number(g['share'], 1)} %", "fraction": (g["share"] or 0) / peak}
                        for g in groups], 7.05, 2.25, 5.7, step=0.58)
        finding = None
    _closing(slide, ctx, 4, total, [finding] if months and finding else [])

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    cut = ctx["cut"].get("retirement_label", "")
    _header(slide, "Compensación jubilar", f"Exservidores y extrabajadores públicos · {ctx['territory_label']} · corte: {cut}")
    retirement = ctx["retirement"]
    rows = [{"name": r["label"], "files": r["files"], "amount": r["amount"], "average": r["average"], "selected": r is retirement[-1]}
            for r in retirement]
    _title(slide, "Expedientes y montos por año", 0.55, 1.45, 6.6, "Montos en USD")
    _table(slide, 0.55, 2.2, 6.6, [("Año", "name", None, "text"), ("Expedientes", "files", 1.3, "int"), ("Monto pagado", "amount", 1.75, "int"),
                                   ("Promedio por expediente", "average", 1.65, "int")], rows, row_h=0.45)
    finding = next((i for i in ctx["insights"] if i["icon"] == "user-check"), ctx["insights"][-1])
    _finding(slide, finding, 0.55, 2.2 + 0.5 + 0.45 * len(rows) + 0.35, 6.6, 1.2)
    parts = ctx["retirement_parts"]
    side = parts["regimes"] or parts["provinces"]
    if len(side) > 1:
        title = (f"Expedientes {ctx['retirement_year']} por régimen laboral" if parts["regimes"] else
                 f"Expedientes {ctx['retirement_year']} por provincia")
        peak = max([r["files"] or 0 for r in side[:6]] or [1]) or 1
        _title(slide, title, 7.55, 1.45, 5.25)
        _ranked(slide, [{"name": r["name"], "text": f"{number(r['files'])} · {number(r['share'], 1)} %", "fraction": (r["files"] or 0) / peak}
                        for r in side[:6]], 7.55, 2.2, 5.25, step=0.62)
    else:
        _title(slide, "Compensaciones pagadas por año", 7.55, 1.45, 5.25, "Millones de USD · el año en curso en dorado")
        _columns_chart(slide, [r["label"] for r in retirement], [r["amount"] for r in retirement], "Compensaciones (millones de USD)",
                       x=7.45, y=2.2, w=5.45, h=4.65, money=True, highlight=len(retirement) - 1)
    sexes = parts["sexes"]
    if sexes:
        _stacked(slide, "Expedientes por sexo", sexes and [{"name": s["name"], "share": s["share"]} for s in sexes], 7.55, 5.95, 5.25,
                 [theme.SERIE_BONOS, theme.LILA_300])
    _closing(slide, ctx, 5, total, [finding])


# ------------------------------------------------------------------------------------------- social budget
def social_budget(prs, ctx: dict) -> None:
    total = 4
    _cover(prs, ctx, emblem_png(width=4.2, height=3.7, dpi=220), total, share_word="")
    stats = ctx["stats"]
    programs = ctx["programs"]

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _header(slide, "Ejecución por programa", f"Devengado de {ctx['span']} sobre el presupuesto aprobado · {ctx['programs_label']}")
    if len(programs) > 4:
        rows = list(reversed(programs))
        chart = _bar_chart(slide, [r["name"] for r in rows], [(r["execution"] or 0) / 100 for r in rows], "Ejecución presupuestaria",
                           number_format="0.0%", maximum=1.0, x=0.45, y=1.45, w=7.0, h=5.45)
        chart.plots[0].gap_width = 60
        chart.category_axis.tick_labels.font.size = Pt(11 if len(rows) > 8 else 12)
    else:
        # Few programs: their execution, then each one month by month.
        _ranked(slide, [{"name": p["name"], "text": f"{number(p['execution'], 1)} % de {_usd(p['approved'])}",
                         "fraction": (p["execution"] or 0) / 100} for p in programs], 0.55, 1.6, 6.8, step=0.6)
        top = 1.6 + 0.6 * len(programs) + 0.2
        _title(slide, "Devengado de cada programa por mes", 0.55, top, 6.8, "Millones de USD")
        matrix = ctx["matrix"]
        data = CategoryChartData(number_format='#,##0.0" M"')
        data.categories = matrix["months"]
        for row in matrix["rows"]:
            data.add_series(row["name"], [(v or 0) / 1e6 for v in row["values"]])
        frame = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.45), Inches(top + 0.75), Inches(7.0), Inches(6.9 - top - 0.8),
                                       data)
        chart = frame.chart
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(10.5)
        chart.font.name = TEXT
        chart.font.size = Pt(10.5)
        chart.value_axis.visible = False
        chart.value_axis.has_major_gridlines = False
        chart.plots[0].gap_width = 70
        chart.plots[0].has_data_labels = True
        chart.plots[0].data_labels.number_format = '#,##0.0" M"'
        chart.plots[0].data_labels.number_format_is_linked = False
        chart.plots[0].data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
        chart.plots[0].data_labels.font.size = Pt(9)
        for series, color in zip(chart.plots[0].series, [theme.SERIE_BONOS, theme.LILA_300, theme.ORO, theme.FUCSIA]):
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = _rgb(color)
    _title(slide, "Aspectos destacados", 7.75, 1.45, 5.0)
    findings = [i for i in ctx["insights"] if i["icon"] in ("target", "award", "trending-up")][:3]
    _findings_column(slide, findings)
    _closing(slide, ctx, 2, total, findings)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    months = ctx["months"]
    selected = next((i for i, m in enumerate(months) if m["selected"]), None)
    _header(slide, "Evolución mensual", f"Presupuesto devengado de cada mes y ejecución acumulada · {ctx['programs_label']}")
    _title(slide, "Presupuesto devengado por mes", 0.55, 1.45, 6.5, "Millones de USD · el mes de la consulta en dorado")
    _columns_chart(slide, [m["short"] for m in months], [m["month_accrued"] for m in months], "Devengado (millones de USD)", x=0.45, y=2.2,
                   w=6.6, h=4.7, money=True, highlight=selected)
    _title(slide, "Ejecución acumulada", 7.35, 1.45, 5.5, "Frente al ritmo uniforme del año")
    rates = [(m["execution"] or 0) / 100 if m["execution"] is not None else None for m in months]
    pace = [m["month"] / 12 for m in months]
    _line_chart(slide, [m["short"] for m in months], [("Ejecución acumulada", rates, theme.SERIE_BONOS), ("Ritmo uniforme del año", pace, theme.LILA_300)],
                x=7.25, y=2.2, w=5.65, h=3.4, number_format="0%", highlight=selected, minimum=0, maximum=1.0, legend=True,
                label_format="0.0%", step=0.25)
    finding = next((i for i in ctx["insights"] if i["icon"] == "calendar"), ctx["insights"][0])
    _finding(slide, finding, 7.35, 5.8, 5.5, 1.1)
    _closing(slide, ctx, 3, total, [finding])

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _header(slide, "Detalle por programa", f"Presupuesto aprobado, devengado y saldo · montos en USD · {ctx['span']}")
    month_word = next((m["label"] for m in months if m["selected"]), "mes")
    closing = ctx.get("closing")
    width = 9.0 if closing else 12.25
    columns = [("Programa presupuestario", "name", None, "text"), ("Aprobado", "approved", 1.25, "int"),
               (f"Devengado {month_word.lower()}", "month_accrued", 1.2, "int"), ("Devengado acumulado", "accrued", 1.3, "int"),
               ("Ejecución", "execution", 0.85, "rate"), ("Saldo por devengar", "balance", 1.2, "int")]
    rows = [dict(p) for p in programs]
    if len(rows) > 1:
        rows.append({"name": "Total de la selección", "approved": stats["approved"], "month_accrued": stats["month"], "accrued": stats["accrued"],
                     "execution": stats["execution"], "balance": stats["balance"], "selected": True})
    row_h = min(0.42, 4.75 / max(1, len(rows)))
    _table(slide, 0.55, 1.45, width, columns, rows, row_h=row_h, header_h=0.5, size=10 if len(rows) > 10 else 11)
    if closing:
        x, w = 9.85, 2.95
        _box(slide, x, 1.45, w, 2.75, theme.ORO_100, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
        _box(slide, x, 1.7, 0.07, 2.25, theme.ORO)
        _badge(slide, "hourglass", x + 0.55, 1.95, 0.3, fill=theme.ORO)
        _text(slide, x + 0.3, 2.4, w - 0.45, 0.6, f"Ritmo necesario de {closing['months']}", size=12.5, color=theme.PIZARRA, bold=True)
        _text(slide, x + 0.3, 2.95, w - 0.45, 0.45, _usd(closing["required"]), size=22, color=theme.MORADO_900, font=DISPLAY, bold=True)
        _text(slide, x + 0.3, 3.4, w - 0.45, 0.7, f"por mes. Promedio de enero a {month_word.lower()}: {_usd(closing['average'])} por mes",
              size=11, color=theme.TINTA)
        finding = next((i for i in ctx["insights"] if i["icon"] == "layers"), None)
        if finding:
            _finding(slide, finding, x, 4.45, w, 2.4)
    table_bottom = 1.45 + 0.5 + row_h * len(rows)
    pending = sorted([p for p in programs if p.get("required") is not None and (p["balance"] or 0) > 0], key=lambda p: -(p["balance"] or 0))
    if closing and pending and 6.85 - table_bottom > 1.6:
        # Room left: the pace each program needs to accrue its approved budget.
        top = table_bottom + 0.35
        _text(slide, 0.55, top, width, 0.35, f"Ritmo necesario por programa de {closing['months']}", size=15, color=theme.MORADO_900,
              font=DISPLAY, bold=True)
        limit = max(1, min(len(pending), int((6.85 - top - 0.9) / 0.36)))
        _table(slide, 0.55, top + 0.42, width, [("Programa presupuestario", "name", None, "text"),
                                               (f"Promedio mensual (enero a {month_word.lower()})", "average", 2.6, "int"),
                                               ("Necesario por mes", "required", 1.9, "int"), ("Diferencia", "difference", 1.3, "spct")],
               [{"name": p["name"], "average": p["average"], "required": p["required"],
                 "difference": (p["required"] / p["average"] - 1) * 100 if p["average"] else None} for p in pending[:limit]],
               row_h=0.36, header_h=0.45, size=11)
    _closing(slide, ctx, 4, total, [])


from .risks import render_pptx as risk_presentation
from .enos import render_pptx as enos_presentation

RENDERERS = {"contingencias": contingency, "jovenes": youth, "laboral": labour, "presupuesto_social": social_budget,
             "riesgos": risk_presentation, "seguimiento_enos": enos_presentation}
