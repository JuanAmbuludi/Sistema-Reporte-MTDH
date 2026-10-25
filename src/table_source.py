"""Read Excel inputs or their compact, cached-value Parquet equivalents."""
from pathlib import Path
import json

import pyarrow.parquet as pq
from openpyxl import load_workbook as open_excel


class ParquetSheet:
    def __init__(self, path):
        table = pq.read_table(path)
        self.title = (table.schema.metadata or {}).get(b"source.sheet", b"Sheet1").decode()
        self.headers = table.column_names
        self.rows = [tuple(r.values()) for r in table.to_pylist()]
        self.max_row, self.max_column = len(self.rows) + 1, len(self.headers)

    def iter_rows(self, min_row=1, max_row=None, max_col=None, values_only=True, **kwargs):
        if not values_only:
            raise ValueError("El lector Parquet devuelve valores, no celdas Excel.")
        rows = [tuple(self.headers), *self.rows]
        for row in rows[min_row-1:max_row]:
            yield tuple(row[:max_col])


class ParquetBook:
    def __init__(self, path):
        path = Path(path)
        paths = [path] if path.is_file() else sorted(path.glob("*.parquet"))
        if not paths:
            raise ValueError(f"No hay hojas Parquet en {path}.")
        self.sheets = [ParquetSheet(p) for p in paths]
        self.sheetnames = [s.title for s in self.sheets]
        self.active = self.sheets[0]

    def __getitem__(self, title):
        return next(s for s in self.sheets if s.title == title)

    def __iter__(self):
        return iter(self.sheets)

    def close(self):
        pass


def load_workbook(path, **kwargs):
    path = Path(path)
    return ParquetBook(path) if path.is_dir() or path.suffix == ".parquet" else open_excel(path, **kwargs)
